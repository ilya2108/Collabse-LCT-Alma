"""Распознавание и парсинг файлов импорта: XLSX / XLS / CSV / JSON (api-contract.md §6.2).

Решения кейсодержателя (Р-18): формат определяется по magic bytes, не по расширению;
XLSX — openpyxl (read-only), XLS (BIFF) — xlrd 2.x; CSV — charset-normalizer с цепочкой
`utf-8-sig → utf-8 → cp1251 → koi8-r → cp866` и `csv.Sniffer` для разделителя;
JSON — массив объектов, ключи становятся «колонками». Объединённые ячейки Excel
(вуз на несколько строк лицензий в таблице КАМа) раскрываются: значение левой верхней
ячейки получает весь диапазон. Модуль чистый (без БД) — юнит-тестируется на байтовых
фикстурах.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from charset_normalizer import from_bytes

from app.core.errors import ApiError

# порядок перебора кодировок зафиксирован контрактом (§6.2)
ENCODING_CHAIN: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1251", "koi8-r", "cp866")
MIN_ENCODING_CONFIDENCE = 0.5

# Защита от decompression-бомб: лимит 20 МБ (files.service.MAX_IMPORT_BYTES) проверяет
# только сжатый размер, а XLSX — это zip с кратностью распаковки в сотни раз.
# Материализация листа обрывается по этим порогам ДО исчерпания памяти.
MAX_IMPORT_ROWS = 100_000
MAX_IMPORT_CELLS = 2_000_000
# суммарный распакованный размер членов zip-архива XLSX (sharedStrings-бомбы и пр.)
MAX_XLSX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024

# <mergeCell ref="A2:A5"/> в XML листа (префикс пространства имён допускается)
_MERGE_CELL_RE = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)"')
_MERGE_SCAN_CHUNK = 1 << 20
_MERGE_SCAN_OVERLAP = 128

_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# символы, ожидаемые в русскоязычных таблицах: кириллица + типографика
_EXPECTED_NON_ASCII = set(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюяАБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ«»—–№·°"
)
_CYRILLIC_LOWER = set("абвгдеёжзийклмнопрстуфхцчшщъыьэюя")
_CYRILLIC_UPPER = set("АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ")

# однобайтовые кодировки часто дают одинаковый _non_ascii_score (можибака
# koi8-r↔cp1251 — тоже «чистая» кириллица, но с перевёрнутым регистром);
# при разнице меньше ε выбор делает вторичный скоринг по доле строчных букв,
# а уверенность занижается, чтобы UI предупредил пользователя (§6.2)
_SCORE_EPSILON = 0.05
AMBIGUOUS_ENCODING_CONFIDENCE = 0.75


def import_bad_format(message: str = "Файл импорта не распознан") -> ApiError:
    return ApiError(400, "import_bad_format", f"{message}. Поддерживаются XLSX, XLS, CSV, JSON")


def import_too_big(detail: str) -> ApiError:
    return import_bad_format(f"Файл импорта слишком велик ({detail}) — разбейте его на части")


def import_encoding_error(tried: tuple[str, ...] = ENCODING_CHAIN) -> ApiError:
    return ApiError(
        400,
        "import_encoding_error",
        "Не удалось определить кодировку файла. Укажите кодировку вручную в опциях маппинга",
        details=[
            {
                "field": "options.encoding",
                "code": "encoding_not_detected",
                "message": f"Пробовали кодировки: {', '.join(tried)}",
            }
        ],
    )


@dataclass
class ParsedTable:
    """Результат парсинга: заголовки + строки данных (все значения — str | None)."""

    file_format: str  # xlsx | xls | csv | json
    columns: list[str]
    rows: list[list[str | None]]
    encoding: str | None = None
    encoding_confidence: float | None = None
    sheets: list[dict[str, Any]] = field(default_factory=list)
    active_sheet: int = 0
    header_row: int | None = 0  # None — в файле нет строки заголовков
    # ошибки уровня парсинга (например, null-элементы JSON): не отказ файла,
    # а построчные ошибки {row, code, message} — сервис кладёт их в import_row_error
    row_issues: list[dict[str, Any]] = field(default_factory=list)
    # исходные 1-based номера строк для rows (JSON с пропусками); None — сплошная нумерация
    row_numbers: list[int] | None = None


# --------------------------------------------------------------------------------------
# Определение формата и кодировки
# --------------------------------------------------------------------------------------


def detect_format(data: bytes) -> str:
    """Формат по magic bytes; текст без JSON-скобки считается CSV."""
    if not data:
        raise import_bad_format("Файл пуст")
    if data[:4] == _XLSX_MAGIC:
        return "xlsx"
    if data[:8] == _XLS_MAGIC:
        return "xls"
    head = data.lstrip(b"\xef\xbb\xbf \t\r\n")
    if head[:1] in (b"[", b"{"):
        return "json"
    if b"\x00" in data[:1024]:
        # бинарный мусор, не текст
        raise import_bad_format("Бинарный формат не распознан")
    return "csv"


def _non_ascii_score(text: str) -> float:
    """Доля «ожидаемых» символов среди подозрительных — отличает cp1251 от koi8-r/cp866.

    Подозрительные — всё не-ASCII плюс управляющие байты (кроме \\t\\r\\n): мусор
    вроде UTF-16, прочитанного как однобайтовая кодировка, роняет оценку ниже порога.
    """
    considered = [ch for ch in text if ord(ch) > 127 or (ord(ch) < 32 and ch not in "\t\r\n")]
    if not considered:
        return 1.0
    expected = sum(1 for ch in considered if ch in _EXPECTED_NON_ASCII)
    return expected / len(considered)


def _lowercase_ratio(text: str) -> float:
    """Доля строчных среди кириллических букв — вторичный скоринг кодировок.

    В реальном русском тексте строчные преобладают; можибака koi8-r↔cp1251
    инвертирует регистр («Значение» → «ъОБЮЕОЙЕ»), поэтому у неверной кодировки
    доля строчных резко падает. Без кириллицы возвращает 0.5 (нейтрально).
    """
    lower = sum(1 for ch in text if ch in _CYRILLIC_LOWER)
    upper = sum(1 for ch in text if ch in _CYRILLIC_UPPER)
    total = lower + upper
    if not total:
        return 0.5
    return lower / total


def detect_encoding(data: bytes) -> tuple[str, float]:
    """(кодировка, уверенность 0..1) перебором цепочки §6.2; ниже 0.5 — ошибка.

    utf-8 проверяется строгим декодированием (BOM — приоритетно); однобайтовые
    русские кодировки различаются скорингом по доле кириллицы среди не-ASCII —
    charset-normalizer используется как подсказка, финальное слово за скорингом
    (детерминированность важна для тестов и повторяемости диалога импорта).
    """
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", 1.0
    try:
        data.decode("utf-8", errors="strict")
        return "utf-8", 0.99
    except UnicodeDecodeError:
        pass

    candidates: list[tuple[str, float, float]] = []
    for encoding in ("cp1251", "koi8-r", "cp866"):
        try:
            text = data.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
        candidates.append((encoding, _non_ascii_score(text), _lowercase_ratio(text)))
    if not candidates:
        raise import_encoding_error()

    try:
        match = from_bytes(bytes(data), cp_isolation=["cp1251", "koi8_r", "cp866"]).best()
    except Exception:  # noqa: BLE001 — сторонняя эвристика не должна ронять импорт
        match = None
    hint = match.encoding.replace("_", "-") if match is not None else None

    top_score = max(score for _enc, score, _low in candidates)
    near = [c for c in candidates if top_score - c[1] < _SCORE_EPSILON]
    if len(near) > 1:
        # неоднозначность (например, koi8-r против cp1251): решают доля строчных
        # букв и совпадение с вердиктом charset-normalizer; уверенность занижена,
        # чтобы фронт предложил проверить сэмплы и options.encoding
        best_encoding = max(
            near,
            key=lambda item: (item[2] + (0.1 if item[0] == hint else 0.0), item[1]),
        )[0]
        best_score = min(top_score, AMBIGUOUS_ENCODING_CONFIDENCE)
        if best_score < MIN_ENCODING_CONFIDENCE:
            raise import_encoding_error()
        return best_encoding, round(best_score, 2)

    best_encoding, best_score, _low = max(candidates, key=lambda item: item[1])
    # подсказка charset-normalizer повышает уверенность при совпадении вердиктов
    if hint == best_encoding and match is not None:
        best_score = max(best_score, 1.0 - match.chaos)
    if best_score < MIN_ENCODING_CONFIDENCE:
        raise import_encoding_error()
    return best_encoding, round(best_score, 2)


def decode_text(data: bytes, forced_encoding: str | None = None) -> tuple[str, str, float]:
    """(текст, кодировка, уверенность); `forced_encoding` — переопределение из options."""
    if forced_encoding:
        try:
            return data.decode(forced_encoding), forced_encoding, 1.0
        except (UnicodeDecodeError, LookupError) as exc:
            raise import_encoding_error((forced_encoding,)) from exc
    encoding, confidence = detect_encoding(data)
    return data.decode(encoding), encoding, confidence


# --------------------------------------------------------------------------------------
# Нормализация значений ячеек
# --------------------------------------------------------------------------------------


def _cell_to_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def _normalize_grid(
    grid: list[list[Any]], header_row: int | None
) -> tuple[list[str], list[list[str | None]]]:
    if header_row is None:
        # файл без заголовков: все строки — данные, колонки «Колонка N» по самой широкой строке
        if not grid:
            raise import_bad_format("В файле нет данных")
        width = max(len(raw) for raw in grid)
        columns = [f"Колонка {idx + 1}" for idx in range(width)]
        data_rows = grid
    else:
        if header_row >= len(grid):
            raise import_bad_format("В файле нет строки заголовков")
        raw_header = grid[header_row]
        columns = [
            (_cell_to_str(cell) or f"Колонка {idx + 1}") for idx, cell in enumerate(raw_header)
        ]
        width = len(columns)
        data_rows = grid[header_row + 1 :]
    rows: list[list[str | None]] = []
    for raw in data_rows:
        row = [_cell_to_str(cell) for cell in raw[:width]]
        row += [None] * (width - len(row))
        rows.append(row)
    return columns, rows


# --------------------------------------------------------------------------------------
# Парсеры форматов
# --------------------------------------------------------------------------------------


def _column_index(letters: bytes) -> int:
    """`b"A"` → 0, `b"AB"` → 27."""
    index = 0
    for char in letters:
        index = index * 26 + (char - 64)
    return index - 1


def _xlsx_merged_ranges(workbook: Any, worksheet: Any) -> list[tuple[int, int, int, int]]:
    """Объединённые диапазоны листа (0-based, границы включительно).

    Read-only режим openpyxl не разбирает `<mergeCells>`, поэтому XML листа читается
    потоком из уже открытого zip и сканируется регуляркой: память — один чанк.
    """
    path = getattr(worksheet, "_worksheet_path", None)
    archive = getattr(workbook, "_archive", None)
    if not path or archive is None:
        return []
    found: set[tuple[int, int, int, int]] = set()
    tail = b""
    try:
        with archive.open(path) as stream:
            while chunk := stream.read(_MERGE_SCAN_CHUNK):
                window = tail + chunk
                for match in _MERGE_CELL_RE.finditer(window):
                    first_col, first_row, last_col, last_row = match.groups()
                    found.add(
                        (
                            int(first_row) - 1,
                            _column_index(first_col),
                            int(last_row) - 1,
                            _column_index(last_col),
                        )
                    )
                tail = window[-_MERGE_SCAN_OVERLAP:]
    except (KeyError, OSError, zipfile.BadZipFile):
        return []
    return sorted(found)


def _fill_merged(grid: list[list[Any]], ranges: list[tuple[int, int, int, int]]) -> None:
    """Значение левой верхней ячейки диапазона — во все его пустые ячейки.

    Диапазон обрезается по фактическим размерам листа: «объединение до XFD1048576»
    не раздувает таблицу сверх уже проверенных лимитов строк и ячеек.
    """
    if not grid:
        return
    width = max(len(row) for row in grid)
    for first_row, first_col, last_row, last_col in ranges:
        if first_row >= len(grid) or first_col >= width:
            continue
        top_row = grid[first_row]
        value = top_row[first_col] if first_col < len(top_row) else None
        if value is None:
            continue
        for row_index in range(first_row, min(last_row, len(grid) - 1) + 1):
            row = grid[row_index]
            last = min(last_col, width - 1)
            if len(row) <= last:
                row.extend([None] * (last + 1 - len(row)))
            for col_index in range(first_col, last + 1):
                if row[col_index] is None:
                    row[col_index] = value


def _check_zip_bomb(data: bytes) -> None:
    """Суммарный распакованный размер членов zip — до передачи openpyxl."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise import_bad_format("Файл XLSX повреждён или не является книгой Excel") from exc
    if unpacked > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise import_too_big(
            f"более {MAX_XLSX_UNCOMPRESSED_BYTES // (1024 * 1024)} МБ после распаковки"
        )


def _parse_xlsx(data: bytes, sheet: int | str, header_row: int | None) -> ParsedTable:
    import openpyxl

    _check_zip_bomb(data)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — битый zip/xml приводим к ошибке контракта
        raise import_bad_format("Файл XLSX повреждён или не является книгой Excel") from exc
    try:
        sheets = [
            {"index": idx, "name": ws.title, "rows": ws.max_row or 0}
            for idx, ws in enumerate(workbook.worksheets)
        ]
        active = _resolve_sheet_index(sheets, sheet)
        worksheet = workbook.worksheets[active]
        grid: list[list[Any]] = []
        cells = 0
        for row in worksheet.iter_rows(values_only=True):
            if len(grid) >= MAX_IMPORT_ROWS:
                raise import_too_big(f"более {MAX_IMPORT_ROWS} строк")
            cells += len(row)
            if cells > MAX_IMPORT_CELLS:
                raise import_too_big(f"более {MAX_IMPORT_CELLS} ячеек")
            grid.append(list(row))
        merged = _xlsx_merged_ranges(workbook, worksheet)
    finally:
        workbook.close()
    _fill_merged(grid, merged)
    columns, rows = _normalize_grid(grid, header_row)
    return ParsedTable(
        file_format="xlsx",
        columns=columns,
        rows=rows,
        sheets=sheets,
        active_sheet=active,
        header_row=header_row,
    )


def _parse_xls(data: bytes, sheet: int | str, header_row: int | None) -> ParsedTable:
    import xlrd

    try:
        # formatting_info — ради merged_cells (объединённые ячейки); старые книги без
        # записей форматирования читаются вторым заходом без него
        workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except Exception:  # noqa: BLE001
        try:
            workbook = xlrd.open_workbook(file_contents=data)
        except Exception as exc:  # noqa: BLE001
            raise import_bad_format("Файл XLS повреждён или не является книгой Excel") from exc
    sheets = [
        {"index": idx, "name": ws.name, "rows": ws.nrows}
        for idx, ws in enumerate(workbook.sheets())
    ]
    active = _resolve_sheet_index(sheets, sheet)
    worksheet = workbook.sheet_by_index(active)
    if worksheet.nrows > MAX_IMPORT_ROWS:
        raise import_too_big(f"более {MAX_IMPORT_ROWS} строк")
    if worksheet.nrows * max(worksheet.ncols, 1) > MAX_IMPORT_CELLS:
        raise import_too_big(f"более {MAX_IMPORT_CELLS} ячеек")
    grid: list[list[Any]] = []
    for row_idx in range(worksheet.nrows):
        row: list[Any] = []
        for cell in worksheet.row(row_idx):
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                row.append(bool(cell.value))
            elif cell.ctype == xlrd.XL_CELL_EMPTY:
                row.append(None)
            else:
                row.append(cell.value)
        grid.append(row)
    # merged_cells xlrd: (rlo, rhi, clo, chi), верхние границы — не включительно
    merged = [
        (rlo, clo, rhi - 1, chi - 1)
        for rlo, rhi, clo, chi in getattr(worksheet, "merged_cells", None) or []
        if rhi > rlo and chi > clo
    ]
    _fill_merged(grid, merged)
    columns, rows = _normalize_grid(grid, header_row)
    return ParsedTable(
        file_format="xls",
        columns=columns,
        rows=rows,
        sheets=sheets,
        active_sheet=active,
        header_row=header_row,
    )


def _parse_csv(data: bytes, header_row: int | None, forced_encoding: str | None) -> ParsedTable:
    text, encoding, confidence = decode_text(data, forced_encoding)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid: list[list[str]] = []
    cells = 0
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        # те же пороги, что у XLSX/XLS: байтовый лимит файла не спасает от
        # миллионов коротких строк, обрываем материализацию заранее
        if len(grid) >= MAX_IMPORT_ROWS:
            raise import_too_big(f"более {MAX_IMPORT_ROWS} строк")
        cells += len(row)
        if cells > MAX_IMPORT_CELLS:
            raise import_too_big(f"более {MAX_IMPORT_CELLS} ячеек")
        grid.append(row)
    if not grid:
        raise import_bad_format("CSV-файл не содержит данных")
    columns, rows = _normalize_grid(grid, header_row)
    table = ParsedTable(
        file_format="csv",
        columns=columns,
        rows=rows,
        encoding=encoding,
        encoding_confidence=confidence,
        header_row=header_row,
    )
    table.sheets = [{"index": 0, "name": "CSV", "rows": len(grid)}]
    return table


def _parse_json(data: bytes) -> ParsedTable:
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise import_bad_format("JSON повреждён или не в UTF-8") from exc
    if isinstance(payload, dict) and isinstance(payload.get("items"), list):
        payload = payload["items"]
    if not isinstance(payload, list):
        raise import_bad_format("Ожидается JSON-массив объектов")
    if not payload:
        raise import_bad_format("JSON-массив пуст")
    if len(payload) > MAX_IMPORT_ROWS:
        raise import_too_big(f"более {MAX_IMPORT_ROWS} строк")
    # толерантность к «грязным» выгрузкам: null и не-объектные элементы пропускаются
    # и становятся построчными ошибками; отказ целиком — только когда объектов нет вовсе
    items: list[tuple[int, dict[str, Any]]] = []
    row_issues: list[dict[str, Any]] = []
    for index, item in enumerate(payload):
        if isinstance(item, dict):
            items.append((index, item))
            continue
        kind = "null" if item is None else type(item).__name__
        row_issues.append(
            {
                # 1-based номер + псевдострока заголовка — та же схема, что header_offset
                "row": index + 2,
                "code": "row_not_object",
                "message": f"Элемент №{index + 1} не является JSON-объектом ({kind}) — пропущен",
            }
        )
    if not items:
        raise import_bad_format("Ожидается JSON-массив объектов")
    columns: list[str] = []
    for _, item in items:
        for key in item:
            if key not in columns:
                columns.append(str(key))
    if len(items) * max(len(columns), 1) > MAX_IMPORT_CELLS:
        raise import_too_big(f"более {MAX_IMPORT_CELLS} ячеек")
    rows = [[_cell_to_str(item.get(column)) for column in columns] for _, item in items]
    return ParsedTable(
        file_format="json",
        columns=columns,
        rows=rows,
        encoding="utf-8",
        encoding_confidence=1.0,
        sheets=[{"index": 0, "name": "JSON", "rows": len(payload) + 1}],
        row_issues=row_issues,
        row_numbers=[index + 2 for index, _ in items],
    )


def _resolve_sheet_index(sheets: list[dict[str, Any]], sheet: int | str) -> int:
    if isinstance(sheet, int):
        if 0 <= sheet < len(sheets):
            return sheet
        raise import_bad_format(f"Лист №{sheet} отсутствует в книге")
    for info in sheets:
        if info["name"] == sheet:
            return int(info["index"])
    raise import_bad_format(f"Лист «{sheet}» отсутствует в книге")


def parse_table(
    data: bytes,
    file_format: str | None = None,
    *,
    sheet: int | str = 0,
    header_row: int | None = 0,
    encoding: str | None = None,
) -> ParsedTable:
    """Единая точка парсинга; формат определяется по magic bytes, если не передан.

    `header_row=None` — файл без строки заголовков: все строки — данные.
    """
    fmt = file_format or detect_format(data)
    if fmt == "xlsx":
        return _parse_xlsx(data, sheet, header_row)
    if fmt == "xls":
        return _parse_xls(data, sheet, header_row)
    if fmt == "json":
        return _parse_json(data)
    if fmt == "csv":
        return _parse_csv(data, header_row, encoding)
    raise import_bad_format(f"Неизвестный формат «{fmt}»")
