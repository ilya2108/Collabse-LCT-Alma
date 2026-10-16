"""Выгрузка отчёта «Активность КАМов»: многолистовой XLSX, PDF для заказчика, CSV/JSON.

Вызывается механизмом выгрузок `POST /export/jobs {entity_type: "report:kam-activity"}`
(MinIO, аудит `export.created`, журнал выгрузок, та же зона видимости, что у API).
- XLSX (openpyxl): «Сводка / КАМы / Вузы / Заявки / Динамика», жирная шапка с заливкой,
  закреплённая первая строка, автофильтр, настоящие даты, проценты `0.0`, нативная
  диаграмма «Взаимодействия по КАМам»; защита CWE-1236 — строки «=…» пишутся текстом.
- PDF (fpdf2): стр. 1 — вид Б для заказчика: 4 KPI (охват, вузы с договором, обучаются,
  выиграно), выводы, «Вузы B2B по этапам на конец периода» с правилом отнесения, риски;
  стр. 2 — показатели активности КАМов, конверсия по этапам, разрез по КАМам, «где
  застревают», студенты на программах; дальше — детализация по КАМам.
- CSV/JSON — плоский лист «Заявки» общими писателями выгрузки.

Определения метрик — те же, что в UI («Как считается») и в `kam_activity.py`.
"""

from __future__ import annotations

import asyncio
import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.models.user import AppUser
from app.modules.reporting.kam_activity import (
    INTERACTION_LABELS,
    REPORT_TZ,
    WORKFLOW_LABELS,
    KamActivityParams,
    build_kam_activity,
    build_kam_activity_summary,
)
from app.modules.reporting.pdf import (
    BRAND,
    MUTED,
    Column,
    KpiItem,
    make_report_pdf,
)

REPORT_TITLE = "Отчёт об активности КАМов"
SHEET_NAMES = ("Сводка", "КАМы", "Вузы", "Заявки", "Динамика")
EMPTY = "—"

METRIC_DEFINITIONS: tuple[str, ...] = (
    "Взаимодействие — собственное действие ответственного КАМа в периоде: ручной переход "
    "по этапам, комментарий, файл заявки, мероприятие с вузом. Системные записи (перевод "
    "схемы, LMS, импорт) и действия других пользователей не считаются.",
    "Этап остановки — этап заявки на конец периода, восстановленный по истории переходов; "
    "дней на этапе — от входа в этот этап до конца периода (или до сейчас). Перевод схемы "
    "процесса и события LMS срок на этапе не обнуляют — как на доске.",
    "Зависла — этап не финальный и заявка на нём дольше порога (этап → воронка → общий).",
    "Выиграно / отказ — заявка вошла в финальный этап в периоде; конверсия — "
    "выиграно / (выиграно + отказ). Охват — закреплённые вузы с контактом / все закреплённые.",
    "Вузы B2B по этапам — каждый вуз один раз, по самому продвинутому этапу его B2B-заявок "
    "на конец периода; «Отказ» — только если других заявок нет. Вузы с договором — договор "
    "действует на конец периода (статус «Действует» или завершён, но срок не истёк).",
    "Сравнение — с периодом той же длины непосредственно перед текущим. Границы периода — "
    "по московскому времени.",
)


def _now() -> datetime:
    return datetime.now(timezone.utc)


@dataclass(frozen=True)
class ExportMeta:
    generated_by: str
    generated_at: datetime
    filters_text: str


# --------------------------------------------------------------------------------------
# Форматирование (общие для XLSX и PDF)
# --------------------------------------------------------------------------------------


def _dt(value: str | None) -> datetime | None:
    return datetime.fromisoformat(value) if value else None


def _msk_date(value: str | None) -> date | None:
    moment = _dt(value)
    return moment.astimezone(REPORT_TZ).date() if moment else None


def _d(value: str | date | None) -> str:
    if value is None:
        return EMPTY
    day = value if isinstance(value, date) else date.fromisoformat(value)
    return day.strftime("%d.%m.%Y")


def _moment(value: str | None) -> str:
    day = _msk_date(value)
    return day.strftime("%d.%m.%Y") if day else EMPTY


def _num(value: Any) -> str:
    if value is None:
        return EMPTY
    if isinstance(value, float) and not value.is_integer():
        return f"{value:.1f}".replace(".", ",")
    try:
        return f"{int(value):,}".replace(",", " ")
    except (TypeError, ValueError):
        return str(value)


def _pct(value: Any) -> str:
    return EMPTY if value is None else f"{_num(value)}%"


def _money(value: Any) -> str:
    if value in (None, ""):
        return EMPTY
    try:
        amount = Decimal(str(value))
    except InvalidOperation:
        return str(value)
    return f"{int(amount.quantize(Decimal(1))):,}".replace(",", " ") + " ₽"


def _money_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(Decimal(str(value)))
    except InvalidOperation:
        return None


def _signed(value: Any, unit: str = "") -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        try:
            value = Decimal(value)
        except InvalidOperation:
            return value
    if value == 0:
        return f"±0{unit}"
    sign = "+" if value > 0 else "−"
    if isinstance(value, Decimal):
        body = _money(abs(value)).removesuffix(" ₽")
    else:
        body = _num(abs(value))
    return f"{sign}{body}{unit}"


def period_text(summary_or_detail: Mapping[str, Any]) -> str:
    period = summary_or_detail["period"]
    parts = [f"Период: {_d(period['from'])} — {_d(period['to'])}"]
    previous = summary_or_detail.get("previous_period")
    if previous:
        parts.append(f"сравнение с {_d(previous['from'])} — {_d(previous['to'])}")
    parts.append(f"данные на {_moment(period['as_of'])}")
    return " · ".join(parts)


def describe_filters(detail: Mapping[str, Any]) -> str:
    """Фильтры текстом для шапки файлов: имена берутся из самого отчёта, без запросов."""
    params = detail.get("params") or {}
    parts = ["Процесс: " + WORKFLOW_LABELS.get(params.get("workflow_type") or "", "B2B и B2C")]
    scope = detail.get("scope") or {}
    kam_ids = params.get("kam_id") or []
    if kam_ids:
        names = {row["kam_id"]: row["kam"] for row in detail.get("rows", [])}
        listed = [names.get(kam_id, kam_id) for kam_id in kam_ids]
        label = "Только свои заявки" if scope.get("own_only") else "КАМ"
        parts.append(f"{label}: {', '.join(listed)}")
    else:
        parts.append("КАМ: все")
    university_ids = params.get("university_id") or []
    if university_ids:
        names = {
            node["university_id"]: node.get("short_name") or node["name"]
            for row in detail.get("rows", [])
            for node in row["universities"]
            if node["university_id"]
        }
        parts.append("Вузы: " + ", ".join(names.get(uid, uid) for uid in university_ids))
    return " · ".join(parts)


def _interactions_short(interactions: Mapping[str, int]) -> str:
    return (
        f"П {interactions['transitions']} · К {interactions['comments']} · "
        f"Ф {interactions['files']} · М {interactions['activities']}"
    )


def _stages_text(stages: Sequence[Mapping[str, Any]]) -> str:
    return "; ".join(f"{stage['name']} ×{stage['count']}" for stage in stages) or EMPTY


def _outcome_text(outcome: str | None) -> str:
    return {"won": "Выиграна", "lost": "Отказ"}.get(outcome or "", EMPTY)


def kam_activity_flat_rows(
    detail: Mapping[str, Any],
) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    """Лист «Заявки» плоской таблицей (CSV/JSON и XLSX)."""
    columns = [
        ("kam", "КАМ"),
        ("university", "Вуз / клиент"),
        ("title", "Заявка"),
        ("workflow_type", "Тип"),
        ("stage", "Этап на конец периода"),
        ("stage_since", "На этапе с"),
        ("days_on_stage", "Дней на этапе"),
        ("threshold_days", "Порог, дней"),
        ("is_stuck", "Зависла"),
        ("closed_in_period", "Итог в периоде"),
        ("created_in_period", "Создана в периоде"),
        ("transitions", "Переходы"),
        ("returns", "Возвраты"),
        ("comments", "Комментарии"),
        ("files", "Файлы"),
        ("last_contact_at", "Последний контакт"),
        ("amount", "Сумма, ₽"),
        ("request_id", "ID заявки"),
    ]
    rows: list[dict[str, Any]] = []
    for kam in detail.get("rows", []):
        for node in kam["universities"]:
            place = node.get("short_name") or node["name"]
            for request in node["requests"]:
                interactions = request["interactions"]
                rows.append(
                    {
                        "kam": kam["kam"],
                        "university": request["client"] if node["university_id"] is None else place,
                        "title": request["title"],
                        "workflow_type": WORKFLOW_LABELS.get(
                            request["workflow_type"], request["workflow_type"]
                        ),
                        "stage": request["stage"]["name"],
                        "stage_since": _msk_date(request["stage_since"]),
                        "days_on_stage": request["days_on_stage"],
                        "threshold_days": request["threshold_days"],
                        "is_stuck": "Да" if request["is_stuck"] else "Нет",
                        "closed_in_period": _outcome_text(request["closed_in_period"]),
                        "created_in_period": "Да" if request["created_in_period"] else "Нет",
                        "transitions": interactions["transitions"],
                        "returns": interactions["returns"],
                        "comments": interactions["comments"],
                        "files": interactions["files"],
                        "last_contact_at": _msk_date(request["last_contact_at"]),
                        "amount": _money_number(request["amount"]),
                        "request_id": request["request_id"],
                    }
                )
    return columns, rows


def universities_stage_rows(by_stage: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Строки главной диаграммы вида Б: этапы + «Без заявок в работе» (если такие вузы были)."""
    rows = [dict(row) for row in by_stage.get("rows", [])]
    without = by_stage.get("without_requests") or {}
    if without.get("value") or without.get("prev"):
        rows.append(
            {
                "id": None,
                "code": "without_requests",
                "name": without.get("label") or "Без заявок в работе",
                "color": MUTED,
                "value": without.get("value", 0),
                "prev": without.get("prev", 0),
                "universities": list(without.get("universities", [])),
            }
        )
    return rows


def risk_note(risk: Mapping[str, Any]) -> str:
    parts = [part for part in (risk.get("note"), "; ".join(risk.get("items") or [])) if part]
    return " — ".join(parts)


def _json_ready(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {key: value.isoformat() if isinstance(value, date) else value for key, value in row.items()}
        for row in rows
    ]


# --------------------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------------------

_DATE_FORMAT = "DD.MM.YYYY"
_PCT_FORMAT = "0.0"
_MONEY_FORMAT = "#,##0"


class _Sheet:
    """Тонкая обёртка над листом openpyxl: строки, стили, защита CWE-1236."""

    def __init__(self, sheet: Any) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        self.sheet = sheet
        self.row = 1
        self._header_font = Font(bold=True, color="FFFFFF")
        self._header_fill = PatternFill("solid", fgColor=BRAND.lstrip("#"))
        self._bold = Font(bold=True)
        self._title = Font(bold=True, size=14)
        self._muted = Font(color=MUTED.lstrip("#"), size=9)
        self._wrap = Alignment(wrap_text=True, vertical="top")

    def put(self, column: int, value: Any, number_format: str | None = None) -> Any:
        cell = self.sheet.cell(row=self.row, column=column, value=value)
        # openpyxl делает строку с ведущим «=» ЖИВОЙ формулой (data_type='f');
        # принудительно сохраняем как текст — значение не меняется (CWE-1236)
        if cell.data_type == "f":
            cell.data_type = "s"
        if number_format and value is not None:
            cell.number_format = number_format
        return cell

    def line(self, values: Sequence[Any], formats: Sequence[str | None] | None = None) -> None:
        for index, value in enumerate(values, start=1):
            fmt = formats[index - 1] if formats and index - 1 < len(formats) else None
            self.put(index, value, fmt)
        self.row += 1

    def header(self, labels: Sequence[str]) -> int:
        for index, label in enumerate(labels, start=1):
            cell = self.put(index, label)
            cell.font = self._header_font
            cell.fill = self._header_fill
            cell.alignment = self._wrap
        header_row = self.row
        self.row += 1
        return header_row

    def title(self, value: str) -> None:
        self.put(1, value).font = self._title
        self.row += 1

    def subtitle(self, value: str, *, bold: bool = False) -> None:
        self.put(1, value).font = self._bold if bold else self._muted
        self.row += 1

    def skip(self, rows: int = 1) -> None:
        self.row += rows

    def widths(self, widths: Sequence[float]) -> None:
        from openpyxl.utils import get_column_letter

        for index, width in enumerate(widths, start=1):
            self.sheet.column_dimensions[get_column_letter(index)].width = width

    def finish_table(self, header_row: int, columns: int) -> None:
        from openpyxl.utils import get_column_letter

        self.sheet.freeze_panes = f"A{header_row + 1}"
        last = max(self.row - 1, header_row)
        self.sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(columns)}{last}"


_KPI_LABELS: tuple[tuple[str, str, str | None], ...] = (
    ("coverage_pct", "Охват закреплённых вузов, %", _PCT_FORMAT),
    ("universities_with_contract", "Вузы с действующим договором", None),
    ("students_studying", "Обучаются (студенты вузов отчёта)", None),
    ("assigned_touched", "Закреплённых вузов с контактом", None),
    ("universities_assigned", "Закреплённых вузов всего", None),
    ("universities_touched", "Вузов с взаимодействиями", None),
    ("requests_count", "Заявок в работе за период", None),
    ("new_requests", "Новые заявки", None),
    ("new_amount", "Сумма новых заявок, ₽", _MONEY_FORMAT),
    ("won", "Выиграно", None),
    ("lost", "Отказы", None),
    ("conversion", "Конверсия, %", _PCT_FORMAT),
    ("interactions", "Взаимодействий", None),
    ("stuck", "Зависших на конец периода", None),
    ("talent_pool_added", "Добавлено в пул талантов", None),
)


def _kpi_number(value: Any) -> Any:
    if isinstance(value, str):
        return _money_number(value)
    return value


def _write_summary_sheet(
    sheet: _Sheet, detail: Mapping[str, Any], summary: Mapping[str, Any], meta: ExportMeta
) -> None:
    from openpyxl.chart import BarChart, Reference

    sheet.title(REPORT_TITLE)
    sheet.subtitle(period_text(summary))
    sheet.subtitle(f"Фильтры: {meta.filters_text}")
    sheet.subtitle(
        f"Сформировал: {meta.generated_by}, "
        f"{meta.generated_at.astimezone(REPORT_TZ):%d.%m.%Y %H:%M} МСК"
    )
    sheet.skip()

    sheet.subtitle("Показатели", bold=True)
    sheet.header(["Показатель", "Текущий", "Предыдущий", "Изменение"])
    for key, label, fmt in _KPI_LABELS:
        item = summary["kpi"].get(key)
        if item is None:
            continue
        sheet.line(
            [
                label,
                _kpi_number(item["value"]),
                _kpi_number(item["prev"]),
                _kpi_number(item["delta"]),
            ],
            [None, fmt, fmt, fmt],
        )
    sheet.skip()

    sheet.subtitle("Ключевые выводы", bold=True)
    for text in summary.get("highlights", []):
        sheet.line([f"• {text}"])
    sheet.skip()

    by_stage = summary.get("universities_by_stage")
    if by_stage:
        sheet.subtitle("Вузы B2B по этапам на конец периода", bold=True)
        sheet.header(["Этап", "Вузов", "Прошлый период", "Вузы"])
        for row in universities_stage_rows(by_stage):
            sheet.line([row["name"], row["value"], row["prev"], ", ".join(row["universities"])])
        sheet.line([f"Правило: {by_stage['rule']}"])
        sheet.skip()

    sheet.subtitle("Риски", bold=True)
    risks = summary.get("risks") or []
    if risks:
        sheet.header(["Риск", "Значение", "Комментарий"])
        for risk in risks:
            sheet.line([risk["label"], risk["value"], risk_note(risk)])
    else:
        sheet.line(["Рисков за период не выявлено"])
    sheet.skip()

    for workflow_type, stages in summary.get("stages", {}).items():
        sheet.subtitle(
            f"Конверсия по этапам — {WORKFLOW_LABELS.get(workflow_type, workflow_type)}",
            bold=True,
        )
        sheet.header(
            [
                "Этап",
                "Дошли",
                "Доля от входа, %",
                "Конверсия шага, %",
                "Сейчас на этапе",
                "Зависли",
                "Ср. дней на этапе",
            ]
        )
        for stage in stages:
            sheet.line(
                [
                    stage["name"],
                    stage["reached"],
                    stage["share_of_first_pct"],
                    stage["step_conversion_pct"],
                    stage["on_stage_at_end"],
                    stage["stuck"],
                    stage["avg_days"],
                ],
                [None, None, _PCT_FORMAT, _PCT_FORMAT, None, None, _PCT_FORMAT],
            )
        sheet.skip()

    sheet.subtitle("По КАМам", bold=True)
    kam_header = sheet.header(
        [
            "КАМ",
            "Взаимодействий",
            "Пред. период",
            "Вузов с контактом",
            "Закреплено",
            "Новых заявок",
            "Выиграно",
            "Отказ",
            "Конверсия, %",
            "Зависших",
            "Последний контакт",
        ]
    )
    for row in summary.get("by_kam", []):
        sheet.line(
            [
                row["kam"],
                row["interactions"],
                row["interactions_prev"],
                row["universities_touched"],
                row["universities_assigned"],
                row["new_requests"],
                row["won"],
                row["lost"],
                row["conversion"],
                row["stuck"],
                _msk_date(row["last_contact_at"]),
            ],
            [None] * 8 + [_PCT_FORMAT, None, _DATE_FORMAT],
        )
    kam_last = sheet.row - 1
    if kam_last > kam_header:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Взаимодействия по КАМам: текущий и прошлый период"
        chart.y_axis.title = "Взаимодействий"
        chart.y_axis.scaling.min = 0  # ось от нуля
        chart.height = max(7.5, 0.55 * (kam_last - kam_header) + 3)
        chart.width = 18
        data = Reference(sheet.sheet, min_col=2, max_col=3, min_row=kam_header, max_row=kam_last)
        categories = Reference(sheet.sheet, min_col=1, min_row=kam_header + 1, max_row=kam_last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        sheet.sheet.add_chart(chart, "M6")
    sheet.skip()

    talent = summary.get("talent") or {}
    sheet.subtitle("Студенты на программах (пул талантов)", bold=True)
    sheet.header(["Статус", "На конец периода", "Перешли в статус за период"])
    for item in talent.get("by_status", []):
        sheet.line([item["label"], item["value"], item["entered"]])
    sheet.skip()
    if talent.get("programs"):
        sheet.header(
            ["Программа", "Вуз", "Кандидаты", "Обучаются", "Выпускники", "Пул талантов", "Всего"]
        )
        for program in talent["programs"]:
            sheet.line(
                [
                    program["name"],
                    program["university"],
                    program["candidate"],
                    program["studying"],
                    program["graduate"],
                    program["talent_pool"],
                    program["total"],
                ]
            )
        sheet.skip()

    sheet.subtitle("Как считается", bold=True)
    for text in METRIC_DEFINITIONS:
        sheet.line([text])
    sheet.widths([46, 16, 16, 16, 16, 14, 14, 12, 14, 12, 18])


def _write_kams_sheet(sheet: _Sheet, detail: Mapping[str, Any]) -> None:
    header = sheet.header(
        [
            "КАМ",
            "Вузов",
            "С контактом",
            "Закреплено",
            "Закреплено без контакта",
            "Заявок",
            "Новых",
            "Выиграно",
            "Отказ",
            "Конверсия, %",
            "Переходы",
            "Возвраты",
            "Комментарии",
            "Файлы",
            "Мероприятия",
            "Всего взаимодействий",
            "Зависших",
            "Последний контакт",
        ]
    )
    for row in detail.get("rows", []):
        interactions = row["interactions"]
        sheet.line(
            [
                row["kam"] + ("" if row["is_active"] else " (неактивен)"),
                row["universities_count"],
                row["universities_touched"],
                row["universities_assigned"],
                row["assigned_untouched"],
                row["requests_count"],
                row["new_requests"],
                row["won"],
                row["lost"],
                row["conversion"],
                interactions["transitions"],
                interactions["returns"],
                interactions["comments"],
                interactions["files"],
                interactions["activities"],
                interactions["total"],
                row["stuck"],
                _msk_date(row["last_contact_at"]),
            ],
            [None] * 9 + [_PCT_FORMAT] + [None] * 7 + [_DATE_FORMAT],
        )
    sheet.finish_table(header, 18)
    sheet.widths([34] + [12] * 16 + [16])


def _write_universities_sheet(sheet: _Sheet, detail: Mapping[str, Any]) -> None:
    header = sheet.header(
        [
            "КАМ",
            "Вуз",
            "Регион",
            "Закреплён",
            "Заявок",
            "Этапы остановки",
            "Переходы",
            "Комментарии",
            "Файлы",
            "Мероприятия",
            "Всего",
            "Зависших",
            "Последний контакт",
        ]
    )
    for row in detail.get("rows", []):
        for node in row["universities"]:
            interactions = node["interactions"]
            sheet.line(
                [
                    row["kam"],
                    node["name"],
                    node.get("region"),
                    "Да" if node["assigned"] else "Нет",
                    node["requests_count"],
                    _stages_text(node["stages"]),
                    interactions["transitions"],
                    interactions["comments"],
                    interactions["files"],
                    interactions["activities"],
                    interactions["total"],
                    node["stuck"],
                    _msk_date(node["last_contact_at"]),
                ],
                [None] * 12 + [_DATE_FORMAT],
            )
    sheet.finish_table(header, 13)
    sheet.widths([30, 44, 24, 11, 9, 40, 11, 13, 9, 13, 9, 10, 16])


def _write_requests_sheet(sheet: _Sheet, detail: Mapping[str, Any]) -> int:
    columns, rows = kam_activity_flat_rows(detail)
    header = sheet.header([label for _key, label in columns])
    formats: list[str | None] = []
    for key, _label in columns:
        if key in ("stage_since", "last_contact_at"):
            formats.append(_DATE_FORMAT)
        elif key == "amount":
            formats.append(_MONEY_FORMAT)
        else:
            formats.append(None)
    for row in rows:
        sheet.line([row[key] for key, _label in columns], formats)
    sheet.finish_table(header, len(columns))
    sheet.widths([28, 22, 44, 7, 26, 13, 10, 10, 9, 13, 11, 10, 10, 12, 8, 16, 14, 38])
    return len(rows)


def _write_dynamics_sheet(sheet: _Sheet, summary: Mapping[str, Any]) -> None:
    series = summary.get("dynamics", {}).get("series", [])
    keys = [item["key"] for item in series]
    header = sheet.header(["Период"] + [INTERACTION_LABELS.get(key, key) for key in keys])
    by_date: dict[str, dict[str, int]] = {}
    for item in series:
        for point in item["points"]:
            by_date.setdefault(point["date"], {})[item["key"]] = point["value"]
    for day in sorted(by_date):
        sheet.line(
            [date.fromisoformat(day)] + [by_date[day].get(key, 0) for key in keys],
            [_DATE_FORMAT],
        )
    sheet.finish_table(header, len(keys) + 1)
    sheet.widths([14] + [14] * len(keys))


def render_kam_activity_xlsx(
    detail: Mapping[str, Any], summary: Mapping[str, Any], meta: ExportMeta
) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheets = [workbook.active] + [workbook.create_sheet() for _ in SHEET_NAMES[1:]]
    for sheet, name in zip(sheets, SHEET_NAMES, strict=True):
        sheet.title = name
    _write_summary_sheet(_Sheet(sheets[0]), detail, summary, meta)
    _write_kams_sheet(_Sheet(sheets[1]), detail)
    _write_universities_sheet(_Sheet(sheets[2]), detail)
    _write_requests_sheet(_Sheet(sheets[3]), detail)
    _write_dynamics_sheet(_Sheet(sheets[4]), summary)
    workbook.properties.title = REPORT_TITLE
    workbook.properties.creator = "Альма"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------


def _kpi_item(
    kpi: Mapping[str, Any],
    key: str,
    label: str,
    *,
    unit: str = "",
    inverted: bool = False,
    money: bool = False,
    hint: str | None = None,
) -> KpiItem:
    item = kpi.get(key) or {}
    value = item.get("value")
    delta = item.get("delta")
    if money:
        shown = _money(value)
    elif unit == "%":
        shown = _pct(value)
    else:
        shown = _num(value)
    if delta is None:
        return KpiItem(label, shown, None, None, inverted, hint or "нет базы для сравнения")
    delta_number = float(Decimal(str(delta)))
    delta_unit = " п.п." if unit == "%" else (" ₽" if money else "")
    if delta_number == 0:
        return KpiItem(label, shown, None, None, inverted, hint or "без изменений к прошлому")
    return KpiItem(
        label,
        shown,
        _signed(delta, delta_unit),
        delta_number > 0,
        inverted,
        hint,
    )


def _pdf_detail_rows(row: Mapping[str, Any]) -> list[dict[str, Any]]:
    items = []
    for node in row["universities"]:
        flags = []
        if node["assigned"]:
            flags.append("закреплён")
        if node["assigned"] and node["interactions"]["total"] == 0:
            flags.append("нет контакта")
        name = node.get("short_name") or node["name"]
        items.append(
            {
                "university": name + (f" ({', '.join(flags)})" if flags else ""),
                "requests": node["requests_count"],
                "stages": _stages_text(node["stages"]),
                "interactions": _interactions_short(node["interactions"]),
                "last": _moment(node["last_contact_at"]),
                "stuck": node["stuck"] or EMPTY,
            }
        )
    return items


def render_kam_activity_pdf(
    detail: Mapping[str, Any], summary: Mapping[str, Any], meta: ExportMeta
) -> bytes:
    generated = f"{meta.generated_at.astimezone(REPORT_TZ):%d.%m.%Y %H:%M} МСК"
    pdf = make_report_pdf(title=REPORT_TITLE, generated=generated)
    kpi = summary["kpi"]
    assigned = (kpi.get("universities_assigned") or {}).get("value") or 0
    touched = (kpi.get("assigned_touched") or {}).get("value") or 0

    # --- стр. 1 (вид Б): 4 KPI, выводы, вузы по этапам, риски ---
    pdf.title_block(
        REPORT_TITLE,
        [
            period_text(summary),
            f"Фильтры: {meta.filters_text}",
            f"Сформировал: {meta.generated_by}",
            "▲▼ — изменение к прошлому периоду той же длины (проценты — в п.п.); "
            "зелёный — к лучшему, красный — к худшему.",
        ],
    )
    conversion = (kpi.get("conversion") or {}).get("value")
    pdf.kpi_grid(
        [
            _kpi_item(
                kpi, "coverage_pct", "Охват вузов", unit="%", hint=f"{touched} из {assigned}"
            ),
            _kpi_item(kpi, "universities_with_contract", "Вузы с договором"),
            _kpi_item(kpi, "students_studying", "Обучаются"),
            _kpi_item(
                kpi,
                "won",
                "Выиграно заявок",
                hint=f"конверсия {_pct(conversion)}" if conversion is not None else None,
            ),
        ]
    )
    pdf.section("Ключевые выводы")
    highlights = summary.get("highlights") or ["Данных за период нет."]
    pdf.bullets(highlights)

    by_stage = summary.get("universities_by_stage")
    if by_stage:
        pdf.section(
            "Вузы B2B по этапам на конец периода",
            f"{by_stage['rule']} В скобках — на конец прошлого периода той же длины. "
            "Ось — от нуля.",
            keep=40,
        )
        stage_rows = universities_stage_rows(by_stage)
        top = max((max(row["value"], row["prev"]) for row in stage_rows), default=0)
        for row in stage_rows:
            names = ", ".join(row["universities"][:3])
            note = f"{row['value']} ({row['prev']})" + (f" · {names}" if names else "")
            pdf.hbar(row["name"], row["value"], top, row.get("color"), note)
        pdf.ln(2)

    pdf.section("Риски")
    risks = summary.get("risks") or []
    if risks:
        pdf.bullets([f"{risk['label']}: {risk['value']}. {risk_note(risk)}" for risk in risks])
    else:
        pdf.paragraph("Рисков за период не выявлено.")

    # --- стр. 2: показатели активности и конверсия по этапам ---
    pdf.add_page()
    pdf.section("Показатели активности КАМов", "Сравнение — с прошлым периодом той же длины.")
    pdf.kpi_grid(
        [
            _kpi_item(kpi, "interactions", "Взаимодействий"),
            _kpi_item(kpi, "new_requests", "Новые заявки"),
            _kpi_item(kpi, "conversion", "Конверсия", unit="%"),
            _kpi_item(kpi, "lost", "Отказы", inverted=True),
            _kpi_item(kpi, "stuck", "Зависших", inverted=True),
            _kpi_item(kpi, "talent_pool_added", "В пул талантов"),
        ],
        columns=3,
    )

    for workflow_type, stages in summary.get("stages", {}).items():
        if not stages:
            continue
        pdf.section(
            f"Конверсия по этапам — {WORKFLOW_LABELS.get(workflow_type, workflow_type)}",
            "Дошли до этапа (по истории переходов) · доля от входа · конверсия шага · "
            "среднее время на этапе · зависли сейчас. Ось — от нуля.",
        )
        top = max((stage["reached"] for stage in stages), default=0)
        for stage in stages:
            note_parts = [f"{stage['reached']}"]
            if stage["share_of_first_pct"] is not None:
                note_parts.append(_pct(stage["share_of_first_pct"]))
            if stage["step_conversion_pct"] is not None:
                note_parts.append(f"шаг {_pct(stage['step_conversion_pct'])}")
            if stage["avg_days"] is not None:
                note_parts.append(f"{_num(stage['avg_days'])} дн.")
            if stage["stuck"]:
                note_parts.append(f"зависли {stage['stuck']}")
            pdf.hbar(stage["name"], stage["reached"], top, stage["color"], " · ".join(note_parts))
        pdf.ln(2)

    # --- разрез по КАМам, «где застревают», студенты ---
    pdf.section(
        "Разрез по КАМам",
        "Взаимодействия — собственные действия КАМа; в скобках — прошлый период той же длины.",
    )
    pdf.data_table(
        [
            Column("kam", "КАМ", 3.0),
            Column("interactions", "Взаимодействия", 2.0, "R"),
            Column("coverage", "Вузы с контактом", 1.5, "R"),
            Column("new", "Новых", 0.9, "R"),
            Column("won_lost", "Выигр. / отказ", 1.3, "R"),
            Column("conversion", "Конв.", 0.9, "R"),
            Column("stuck", "Зависшие", 1.3, "R"),
            Column("last", "Последний контакт", 1.6),
        ],
        [
            {
                "kam": row["kam"] + ("" if row["is_active"] else " (неактивен)"),
                "interactions": f"{row['interactions']} ({row['interactions_prev']})",
                "coverage": f"{row['assigned_touched']} из {row['universities_assigned']}"
                if row["universities_assigned"]
                else str(row["universities_touched"]),
                "new": row["new_requests"],
                "won_lost": f"{row['won']} / {row['lost']}",
                "conversion": _pct(row["conversion"]),
                "stuck": row["stuck"] or EMPTY,
                "last": _moment(row["last_contact_at"]),
            }
            for row in summary.get("by_kam", [])
        ],
    )

    pdf.section(
        "Где застревают: топ-10 зависших заявок",
        "Заявки на нефинальном этапе дольше порога на конец периода.",
    )
    pdf.data_table(
        [
            Column("title", "Заявка", 3.6),
            Column("university", "Вуз / клиент", 1.6),
            Column("kam", "КАМ", 2.2),
            Column("stage", "Этап", 1.8),
            Column("days", "Дней / порог", 1.2, "R"),
        ],
        [
            {
                "title": item["title"],
                "university": item["university"],
                "kam": item["kam"],
                "stage": item["stage"]["name"],
                "days": f"{item['days_on_stage']} / {item['threshold_days']}",
            }
            for item in summary.get("top_stuck", [])
        ],
        empty_text="Зависших заявок на конец периода нет.",
    )

    talent = summary.get("talent") or {}
    pdf.section(
        "Студенты на программах (пул талантов)",
        "Статус на конец периода по истории воронки студентов; «за период» — перешли в "
        "статус внутри периода. Только агрегаты, без персональных данных.",
    )
    statuses = talent.get("by_status", [])
    top_status = max((item["value"] for item in statuses), default=0)
    for item in statuses:
        pdf.hbar(
            item["label"],
            item["value"],
            top_status,
            item["color"],
            f"{item['value']} · за период +{item['entered']}",
        )
    pdf.ln(2)
    if talent.get("programs"):
        pdf.data_table(
            [
                Column("name", "Программа", 3.4),
                Column("university", "Вуз", 1.4),
                Column("studying", "Обучаются", 1.1, "R"),
                Column("graduate", "Выпускники", 1.1, "R"),
                Column("talent_pool", "Пул талантов", 1.2, "R"),
                Column("total", "Всего", 0.9, "R"),
            ],
            talent["programs"],
        )

    # --- стр. 3+: детализация ---
    pdf.add_page()
    pdf.section(
        "Детализация по КАМам",
        "Вузы КАМа: П — переходы, К — комментарии, Ф — файлы, М — мероприятия. "
        "Пообъектный список заявок — в XLSX (лист «Заявки»).",
    )
    rows = detail.get("rows", [])
    if not rows:
        pdf.paragraph("За выбранный период взаимодействий нет.")
    for row in rows:
        interactions = row["interactions"]
        pdf.section(
            row["kam"] + ("" if row["is_active"] else " (неактивен)"),
            f"Вузов {row['universities_count']} (с контактом {row['universities_touched']}, "
            f"закреплено {row['universities_assigned']}) · заявок {row['requests_count']}, "
            f"новых {row['new_requests']} · взаимодействий {interactions['total']} "
            f"({_interactions_short(interactions)}) · зависших {row['stuck']} · "
            f"последний контакт {_moment(row['last_contact_at'])}",
            keep=48,
        )
        pdf.data_table(
            [
                Column("university", "Вуз", 2.6),
                Column("requests", "Заявок", 0.8, "R"),
                Column("stages", "Этапы остановки", 2.9),
                Column("interactions", "П / К / Ф / М", 1.9),
                Column("last", "Последний контакт", 1.4),
                Column("stuck", "Зависшие", 1.1, "R"),
            ],
            _pdf_detail_rows(row),
            empty_text="Вузов и заявок за период нет.",
        )

    pdf.section("Как считается")
    for text in METRIC_DEFINITIONS:
        pdf.paragraph(text, size=7.5)
        pdf.ln(0.8)
    return pdf.to_bytes()


# --------------------------------------------------------------------------------------
# Точка входа выгрузки
# --------------------------------------------------------------------------------------

_EXTENSIONS = {"xlsx": "xlsx", "pdf": "pdf", "csv": "csv", "json": "json"}


async def build_kam_activity_file(
    session: AsyncSession,
    *,
    user: AppUser,
    filters: Mapping[str, Any],
    export_format: str,
    options: Mapping[str, Any] | None = None,
) -> tuple[bytes, int, str]:
    """Файл отчёта: `(байты, строк заявок, имя файла)`; скоуп — как у API (403 на чужих)."""
    if export_format not in _EXTENSIONS:
        from app.core.errors import unprocessable

        raise unprocessable(f"Формат «{export_format}» не поддерживается для отчёта КАМов")
    params = KamActivityParams.from_filters(filters)
    detail = await build_kam_activity(session, user, params, force_requests=True)
    columns, flat_rows = kam_activity_flat_rows(detail)
    period = detail["period"]
    stamp = f"{date.fromisoformat(period['from']):%Y%m%d}-{date.fromisoformat(period['to']):%Y%m%d}"
    file_name = f"kam-activity_{stamp}.{_EXTENSIONS[export_format]}"

    if export_format in ("csv", "json"):
        from app.modules.import_export.export import render_export

        opts = options or {}
        data = await asyncio.to_thread(
            render_export,
            columns,
            _json_ready(flat_rows) if export_format == "json" else flat_rows,
            export_format=export_format,
            encoding=(opts.get("encoding") or "utf-8").lower(),
            delimiter=opts.get("csv_delimiter") or ";",
        )
        return data, len(flat_rows), file_name

    summary = await build_kam_activity_summary(session, user, params)
    meta = ExportMeta(
        generated_by=user.full_name or user.username,
        generated_at=_now(),
        filters_text=describe_filters(detail),
    )
    renderer = render_kam_activity_xlsx if export_format == "xlsx" else render_kam_activity_pdf
    data = await asyncio.to_thread(renderer, detail, summary, meta)
    return data, len(flat_rows), file_name
