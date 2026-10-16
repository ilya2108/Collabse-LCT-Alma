"""«Взаимодействия за период» (P0-6), область отчётов RBAC (P0-11) и фильтры выгрузок
(P0-2) на настоящем PostgreSQL.

`NOW` = 20.10.2026 12:00 МСК (подмена `kam_activity._now`); «30 дней» = 21.09–20.10.
Сквозные кейсы — через приложение (`api.as_user`): своя роль «КАМ B2C» и роль
«Аналитик» (только `reports:read` + `reports:all`) из каталога ролей в БД, выгрузка
`POST /export/jobs {entity_type: report:interactions}` в XLSX/PDF/CSV.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta
from typing import Any

import pytest

from app.models import AppUser, Product, Program
from app.modules.reporting import kam_activity as ka
from app.modules.reporting.interactions import (
    COLUMNS_BY_KEY,
    DEFAULT_COLUMNS,
    InteractionsParams,
    build_interactions,
)
from app.modules.reporting.kam_activity import KamActivityParams
from app.modules.reporting.schemas import (
    InteractionsReportOut,
    KamActivityReportOut,
    KamActivitySummaryOut,
)

from . import factories
from .rbac_factories import KAM_PERMISSIONS, make_role

API = "/api/v1"
NOW = datetime(2026, 10, 20, 12, 0, tzinfo=ka.REPORT_TZ)


def ago(days: float) -> datetime:
    return NOW - timedelta(days=days)


@pytest.fixture(autouse=True)
def _fixed_now(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ka, "_now", lambda: NOW)


@pytest.fixture
async def world(db_session, workflows) -> dict[str, Any]:
    """Два КАМа, два вуза, продукты и программы; заявки B2B и B2C в периоде и вне его."""
    wf, stages = workflows["b2b"]
    wf_c, stages_c = workflows["b2c"]
    make = factories.make_user
    head = await make(db_session, roles=("head_kam",), full_name="Морозов Дмитрий")
    kam1 = await make(db_session, roles=("kam",), full_name="Петров Константин")
    kam2 = await make(db_session, roles=("kam",), full_name="Иванова Мария")
    mipt = await factories.make_university(
        db_session, name="Московский физико-технический институт", short_name="МФТИ", kam=kam1
    )
    itmo = await factories.make_university(
        db_session, name="Университет ИТМО", short_name="ИТМО", kam=kam2
    )
    python = Product(code="python", name="Python-разработчик")
    devops = Product(code="devops", name="DevOps-инженер")
    db_session.add_all([python, devops])
    await db_session.flush()
    program = Program(name="Python для вузов", product_id=python.id, university_id=mipt.id)
    db_session.add(program)
    await db_session.flush()

    mipt_deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        university=mipt,
        title="МФТИ — Python",
        created_at=ago(100),
        path=[("new", ago(100)), ("first_contact", ago(60)), ("negotiation", ago(20))],
        product_id=python.id,
        program_id=program.id,
        amount="450000.00",
    )
    await factories.make_comment(db_session, deal=mipt_deal, author=kam1, at=ago(10))
    itmo_deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        university=itmo,
        title="ИТМО — DevOps",
        created_at=ago(15),
        path=[("new", ago(15)), ("first_contact", ago(5))],
        product_id=devops.id,
    )
    closed_long_ago = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        university=itmo,
        title="ИТМО — закрыта весной",
        created_at=ago(200),
        path=[("new", ago(200)), ("rejected", ago(150))],
    )
    client = await factories.make_counterparty(db_session, display_name="Смирнова О. В.")
    b2c_deal = await factories.make_deal(
        db_session,
        workflow=wf_c,
        stages=stages_c,
        responsible=kam1,
        counterparty=client,
        title="Курс Python — физлицо",
        created_at=ago(4),
        product_id=python.id,
    )
    await factories.make_activity(db_session, university=mipt, responsible=kam1, slots=[ago(3)])
    return {
        "head": head,
        "kam1": kam1,
        "kam2": kam2,
        "mipt": mipt,
        "itmo": itmo,
        "python": python,
        "devops": devops,
        "program": program,
        "mipt_deal": mipt_deal,
        "itmo_deal": itmo_deal,
        "closed": closed_long_ago,
        "b2c_deal": b2c_deal,
        "stages": stages,
    }


def _ids(report: dict[str, Any]) -> set[str]:
    return {row["request_id"] for row in report["rows"]}


def _kam_report_ids(report: dict[str, Any]) -> set[str]:
    return {
        request["request_id"]
        for row in report["rows"]
        for node in row["universities"]
        for request in node["requests"]
    }


async def _interactions(session, viewer: AppUser, **kwargs: Any) -> dict[str, Any]:
    base = {
        key: kwargs.pop(key)
        for key in ("date_from", "date_to", "workflow_type", "kam_ids", "university_ids")
        if key in kwargs
    }
    report = await build_interactions(
        session, viewer, InteractionsParams(base=KamActivityParams(**base), **kwargs)
    )
    InteractionsReportOut.model_validate(report)  # форма ответа = схема OpenAPI
    return report


# --- P0-6: одна семантика периода с отчётом КАМов -----------------------------------------


async def test_same_rows_and_period_as_kam_activity(db_session, world) -> None:
    head = world["head"]
    report = await _interactions(db_session, head)
    kam_report = await ka.build_kam_activity(db_session, head, KamActivityParams())
    assert report["period"] == kam_report["period"]
    assert _ids(report) == _kam_report_ids(kam_report)
    assert str(world["closed"].id) not in _ids(report)  # закрыта до периода — не в отчёте

    past = {"date_from": ago(90).date(), "date_to": ago(45).date()}
    assert _ids(await _interactions(db_session, head, **past)) == _kam_report_ids(
        await ka.build_kam_activity(db_session, head, KamActivityParams(**past))
    )


async def test_row_columns_and_values(db_session, world) -> None:
    report = await _interactions(db_session, world["head"])
    assert report["default_columns"] == list(DEFAULT_COLUMNS)
    assert [column["key"] for column in report["columns"]][:5] == [
        "university",
        "it_direction",
        "product",
        "status",
        "responsible",
    ]
    rows = {row["request_id"]: row for row in report["rows"]}
    mipt = rows[str(world["mipt_deal"].id)]
    assert mipt["university"] == "Московский физико-технический институт"
    assert mipt["it_direction"] is None  # до справочника ИТ-направлений
    assert mipt["product"] == "Python-разработчик" and mipt["program"] == "Python для вузов"
    assert mipt["stage"]["code"] == "negotiation"  # статус работы — этап на конец периода
    assert mipt["responsible"] == "Петров Константин"
    assert mipt["days_on_stage"] == 20
    # ручной переход −20 д и комментарий −10 д ответственного КАМа
    assert mipt["interactions"] == {
        "transitions": 1,
        "returns": 0,
        "comments": 1,
        "files": 0,
        "total": 2,
    }
    b2c = rows[str(world["b2c_deal"].id)]
    assert b2c["university"] is None and b2c["client"] == "Смирнова О. В."
    # B2B по алфавиту вузов, B2C — в конце
    assert report["rows"][-1]["request_id"] == str(world["b2c_deal"].id)
    assert report["totals"]["requests"] == 3 and report["totals"]["universities"] == 2
    stage_codes = {(option["workflow_type"], option["code"]) for option in report["stages"]}
    assert ("b2b", "negotiation") in stage_codes and ("b2c", "new") in stage_codes


async def test_filters_product_program_stage_kam_university(db_session, world) -> None:
    head = world["head"]
    by_product = await _interactions(db_session, head, product_ids=(world["python"].id,))
    assert _ids(by_product) == {str(world["mipt_deal"].id), str(world["b2c_deal"].id)}
    by_program = await _interactions(db_session, head, program_ids=(world["program"].id,))
    assert _ids(by_program) == {str(world["mipt_deal"].id)}
    by_stage = await _interactions(
        db_session, head, stage_ids=(world["stages"]["first_contact"].id,)
    )
    assert _ids(by_stage) == {str(world["itmo_deal"].id)}
    by_kam = await _interactions(db_session, head, kam_ids=(world["kam2"].id,))
    assert _ids(by_kam) == {str(world["itmo_deal"].id)}
    by_university = await _interactions(db_session, head, university_ids=(world["mipt"].id,))
    assert _ids(by_university) == {str(world["mipt_deal"].id)}
    b2c_only = await _interactions(db_session, head, workflow_type="b2c")
    assert _ids(b2c_only) == {str(world["b2c_deal"].id)}
    combined = await _interactions(
        db_session,
        head,
        product_ids=(world["python"].id, world["devops"].id),
        workflow_type="b2b",
    )
    assert _ids(combined) == {str(world["mipt_deal"].id), str(world["itmo_deal"].id)}


async def test_api_scope_and_validation(api, db_session, world) -> None:
    url = f"{API}/reports/interactions"
    kam1, kam2 = world["kam1"], world["kam2"]
    own = await api.as_user(kam1).get(url)
    assert own.status_code == 200, own.text
    InteractionsReportOut.model_validate(own.json())
    assert {row["responsible_id"] for row in own.json()["rows"]} == {str(kam1.id)}
    assert own.json()["scope"]["own_only"] is True

    foreign = await api.as_user(kam1).get(url, params={"kam_id": str(kam2.id)})
    assert foreign.status_code == 403
    many = await api.as_user(world["head"]).get(
        url, params=[("product_id", str(uuid.uuid4())) for _ in range(ka.MAX_FILTER_IDS + 1)]
    )
    assert many.status_code == 400
    assert many.json()["error"]["details"][0]["code"] == "too_many_values"
    multi = await api.as_user(world["head"]).get(
        url,
        params=[("stage_id", str(world["stages"]["negotiation"].id)), ("workflow_type", "b2b")],
    )
    assert multi.status_code == 200 and _ids(multi.json()) == {str(world["mipt_deal"].id)}
    assert (await api.anonymous().get(url)).status_code == 401


# --- P0-6: выгрузка report:interactions -----------------------------------------------------


async def test_export_xlsx_pdf_csv_by_selected_columns(
    api, db_session, world, memory_storage
) -> None:
    import openpyxl

    client = api.as_user(world["head"])
    filters = {"from": "2026-09-21", "to": "2026-10-20", "workflow_type": "b2b"}
    columns = ["university", "it_direction", "product", "status", "responsible", "created_at"]
    created = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:interactions",
            "format": "xlsx",
            "filters": filters,
            "columns": columns,
        },
    )
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["rows"] == 2 and job["file"]["file_name"] == "interactions_20260921-20261020.xlsx"
    content = (await client.get(f"{API}/export/jobs/{job['id']}/download")).content
    book = openpyxl.load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["Взаимодействия", "О выгрузке"]
    sheet = book["Взаимодействия"]
    assert [cell.value for cell in sheet[1]] == [COLUMNS_BY_KEY[key].label for key in columns]
    values = {row[0]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    mipt = values["Московский физико-технический институт"]
    assert mipt[1] == "—" and mipt[2] == "Python-разработчик" and mipt[3] == "Переговоры"
    assert isinstance(mipt[5], datetime)  # дата — настоящая дата Excel

    pdf = await client.post(
        f"{API}/export/jobs",
        json={"entity_type": "report:interactions", "format": "pdf", "filters": filters},
    )
    assert pdf.status_code == 201, pdf.text
    pdf_bytes = (await client.get(f"{API}/export/jobs/{pdf.json()['id']}/download")).content
    assert pdf_bytes.startswith(b"%PDF")

    csv_job = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:interactions",
            "format": "csv",
            "filters": {**filters, "product_id": [str(world["python"].id)]},
            "columns": ["university", "product"],
        },
    )
    assert csv_job.status_code == 201 and csv_job.json()["rows"] == 1
    text = (await client.get(f"{API}/export/jobs/{csv_job.json()['id']}/download")).content
    assert text.decode("utf-8-sig").splitlines()[0] == "Наименование вуза;ИТ-продукт"

    unknown = await client.post(
        f"{API}/export/jobs",
        json={"entity_type": "report:interactions", "format": "xlsx", "columns": ["salary"]},
    )
    assert unknown.status_code == 400
    assert unknown.json()["error"]["details"][0]["code"] == "unknown_column"
    bad_filter = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:interactions",
            "format": "xlsx",
            "filters": {"product_id": ["не-uuid"]},
        },
    )
    assert bad_filter.status_code == 422
    assert bad_filter.json()["error"]["details"][0]["code"] == "invalid_uuid"
    foreign = await api.as_user(world["kam1"]).post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:interactions",
            "format": "pdf",
            "filters": {"kam_id": [str(world["kam2"].id)]},
        },
    )
    assert foreign.status_code == 403


# --- P0-2: выгрузки с фильтрами-списками — не 500 -----------------------------------------


async def test_export_list_filters_apply_all_values(api, db_session, world, memory_storage) -> None:
    client = api.as_user(world["head"])
    both = [str(world["kam1"].id), str(world["kam2"].id)]
    dashboard = await client.post(
        f"{API}/export/jobs",
        json={"entity_type": "report:dashboard", "format": "xlsx", "filters": {"kam_id": both}},
    )
    assert dashboard.status_code == 201, dashboard.text
    workload = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:kam-workload",
            "format": "csv",
            "filters": {"kam_id": both, "university_id": [str(world["mipt"].id)]},
        },
    )
    assert workload.status_code == 201, workload.text

    requests = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "requests",
            "format": "json",
            "filters": {"assignee_id": both, "product_id": [str(world["python"].id)]},
        },
    )
    assert requests.status_code == 201, requests.text
    payload = (await client.get(f"{API}/export/jobs/{requests.json()['id']}/download")).json()
    titles = {item["title"] for item in payload["items"]}
    # оба КАМа и продукт Python: B2B МФТИ и B2C физлицо, а не «первый элемент списка»
    assert titles == {"МФТИ — Python", "Курс Python — физлицо"}

    for entity_type, filters in (
        ("report:dashboard", {"kam_id": ["не-uuid"]}),
        ("report:funnel", {"from": "вчера"}),
        ("requests", {"assignee_id": [both[0], "мусор"]}),
        ("requests", {"created_from": "32.13.2026"}),
        ("report:kam-activity", {"university_id": "123"}),
    ):
        bad = await client.post(
            f"{API}/export/jobs",
            json={"entity_type": entity_type, "format": "xlsx", "filters": filters},
        )
        assert bad.status_code == 422, (entity_type, bad.text)
        assert bad.json()["error"]["code"] == "unprocessable"
        assert bad.json()["error"]["details"], entity_type


# --- P0-11: область отчётов RBAC ------------------------------------------------------------


async def test_kam_b2c_sees_no_b2b_universities_activities_or_rows(
    api, db_session, workflows, world
) -> None:
    await make_role(
        db_session, code="kam_b2c", permissions=KAM_PERMISSIONS, workflow_types=("b2c",)
    )
    kam_c = await factories.make_user(db_session, roles=("kam_b2c",), full_name="Смирнова Анна")
    # КАМ B2C закреплён за вузом и проводит в нём мероприятие — это сущности B2B
    await factories.make_university(db_session, name="Закреплённый вуз", kam=kam_c)
    await factories.make_activity(
        db_session, university=world["itmo"], responsible=kam_c, slots=[ago(2)]
    )
    wf_c, stages_c = workflows["b2c"]
    own = await factories.make_deal(
        db_session, workflow=wf_c, stages=stages_c, responsible=kam_c, created_at=ago(3)
    )
    client = api.as_user(kam_c)

    detail = (await client.get(f"{API}/reports/kam-activity")).json()
    KamActivityReportOut.model_validate(detail)
    universities = [node for row in detail["rows"] for node in row["universities"]]
    assert [node["university_id"] for node in universities] == [None]
    assert detail["totals"]["universities_assigned"] == 0
    assert detail["totals"]["interactions"]["activities"] == 0

    summary = (await client.get(f"{API}/reports/kam-activity/summary")).json()
    KamActivitySummaryOut.model_validate(summary)
    assert summary["universities_by_stage"] is None  # главная диаграмма вида Б — про B2B
    assert summary["untouched_assigned"] == []
    assert summary["kpi"]["universities_with_contract"]["value"] == 0

    interactions = (await client.get(f"{API}/reports/interactions")).json()
    assert _ids(interactions) == {str(own.id)}
    assert {option["workflow_type"] for option in interactions["stages"]} == {"b2c"}
    for path in ("kam-activity", "kam-activity/summary", "interactions"):
        denied = await client.get(f"{API}/reports/{path}", params={"workflow_type": "b2b"})
        assert denied.status_code == 403, path


async def test_analyst_role_gets_wide_non_empty_reports(api, db_session, world) -> None:
    """Роль «Аналитик»: только `reports:read` + `reports:all` (без прав на заявки и без
    «ведёт портфель») — широкий отчёт по всем КАМам, не пустой."""
    await make_role(
        db_session,
        code="analyst",
        name="Аналитик",
        permissions=("reports:read", "reports:all"),
    )
    analyst = await factories.make_user(db_session, roles=("analyst",), full_name="Аналитиков А.")
    client = api.as_user(analyst)

    detail = await client.get(f"{API}/reports/kam-activity")
    assert detail.status_code == 200, detail.text
    body = detail.json()
    assert body["scope"]["own_only"] is False
    assert {str(world["kam1"].id), str(world["kam2"].id)} <= {row["kam_id"] for row in body["rows"]}
    assert str(analyst.id) not in {row["kam_id"] for row in body["rows"]}
    assert body["totals"]["requests_count"] == 3

    summary = await client.get(f"{API}/reports/kam-activity/summary")
    assert summary.status_code == 200
    assert summary.json()["universities_by_stage"]["total"] == 2

    interactions = await client.get(f"{API}/reports/interactions")
    assert interactions.status_code == 200
    assert len(interactions.json()["rows"]) == 3


async def test_reports_read_without_all_is_own_only(api, db_session, world) -> None:
    """Без `reports:all` отчёт сужается до своих заявок — даже у роли без «КАМ»."""
    await make_role(db_session, code="reader", permissions=("reports:read",))
    reader = await factories.make_user(db_session, roles=("reader",), full_name="Читатель Ч.")
    body = (await api.as_user(reader).get(f"{API}/reports/kam-activity")).json()
    assert body["scope"]["own_only"] is True
    assert body["params"]["kam_id"] == [str(reader.id)]
    assert body["totals"]["requests_count"] == 0
    denied = await api.as_user(reader).get(
        f"{API}/reports/interactions", params={"kam_id": str(world["kam1"].id)}
    )
    assert denied.status_code == 403
    no_reports = await factories.make_user(db_session, roles=("kam_without_reports",))
    assert (await api.as_user(no_reports).get(f"{API}/reports/interactions")).status_code == 403
