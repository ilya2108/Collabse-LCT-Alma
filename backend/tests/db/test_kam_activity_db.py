"""Отчёт «Активность КАМов» на настоящем PostgreSQL (F1): сервисы, API, выгрузки, файлы.

`NOW` = 20.10.2026 12:00 МСК (подмена `kam_activity._now`); «30 дней» = 21.09–20.10.
Данные строятся относительно NOW общими фабриками tests/db/factories.py; KeyDB в
тестах недоступен — заодно проверяется деградация кэша в чтение из PG.

Сквозные кейсы идут через приложение (`api.as_user`): область RBAC (`report_scope`,
своя роль «КАМ B2C»), выгрузка `POST /export/jobs` в XLSX/PDF с аудитом, файлы заявки
(`/requests/{id}/files`, скачивание через проверку области), D16 дашборда, индексы 0005.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta
from typing import Any

import pytest
import sqlalchemy as sa

from app.core.access import DealScope, reset_role_catalog
from app.core.errors import ApiError
from app.models import (
    AppSetting,
    AppUser,
    AuditLog,
    Contract,
    Deal,
    DealStageHistory,
    FileObject,
    Program,
    Student,
    StudentStatusHistory,
    WorkflowStage,
)
from app.modules.reporting import kam_activity as ka
from app.modules.reporting.kam_activity import KamActivityParams
from app.modules.reporting.kam_activity_export import (
    SHEET_NAMES,
    build_kam_activity_file,
)
from app.modules.reporting.schemas import KamActivityReportOut, KamActivitySummaryOut
from app.modules.reporting.service import ReportParams, build_report

from . import factories
from .rbac_factories import KAM_PERMISSIONS, make_role

API = "/api/v1"

NOW = datetime(2026, 10, 20, 12, 0, tzinfo=ka.REPORT_TZ)
TODAY = NOW.date()


def ago(days: float) -> datetime:
    return NOW - timedelta(days=days)


@pytest.fixture(autouse=True)
def _fixed_now(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ka, "_now", lambda: NOW)


@pytest.fixture
async def people(db_session) -> dict[str, AppUser]:
    make = factories.make_user
    head = await make(db_session, roles=("head_kam",), full_name="Морозов Дмитрий Александрович")
    return {
        "head": head,
        "kam1": await make(
            db_session, roles=("kam",), full_name="Петров Константин Андреевич", manager=head
        ),
        "kam2": await make(
            db_session, roles=("kam",), full_name="Иванова Мария Сергеевна", manager=head
        ),
        "observer": await make(db_session, roles=("observer",), full_name="Кузнецов Олег"),
    }


def _row(report: dict[str, Any], user: AppUser) -> dict[str, Any]:
    return next(row for row in report["rows"] if row["kam_id"] == str(user.id))


def _node(row: dict[str, Any], university_id: uuid.UUID | None) -> dict[str, Any]:
    key = str(university_id) if university_id else None
    return next(node for node in row["universities"] if node["university_id"] == key)


def _request(report: dict[str, Any], deal: Deal) -> dict[str, Any] | None:
    for row in report["rows"]:
        for node in row["universities"]:
            for request in node["requests"]:
                if request["request_id"] == str(deal.id):
                    return request
    return None


async def _detail(session, viewer: AppUser, **kwargs: Any) -> dict[str, Any]:
    report = await ka.build_kam_activity(session, viewer, KamActivityParams(**kwargs))
    KamActivityReportOut.model_validate(report)  # форма ответа = схема OpenAPI
    return report


async def _summary(session, viewer: AppUser, **kwargs: Any) -> dict[str, Any]:
    report = await ka.build_kam_activity_summary(session, viewer, KamActivityParams(**kwargs))
    KamActivitySummaryOut.model_validate(report)
    return report


@pytest.fixture
async def mipt_deal(db_session, workflows, people) -> dict[str, Any]:
    """Кейс 1: kam1/МФТИ — создана −100 д (new), −60 д → first_contact, −20 д → negotiation;
    комментарий −10 д, файл −5 д."""
    wf, stages = workflows["b2b"]
    kam1 = people["kam1"]
    mipt = await factories.make_university(
        db_session, name="Московский физико-технический институт", short_name="МФТИ", kam=kam1
    )
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        university=mipt,
        title="МФТИ — корпоративная школа Python",
        created_at=ago(100),
        path=[("new", ago(100)), ("first_contact", ago(60)), ("negotiation", ago(20))],
        amount="450000.00",
    )
    await factories.make_comment(db_session, deal=deal, author=kam1, at=ago(10))
    await factories.make_file(db_session, entity_id=deal.id, uploaded_by=kam1, at=ago(5))
    return {"deal": deal, "university": mipt}


# --- кейсы 1–2: этап на конец периода и взаимодействия -----------------------------------


async def test_case1_stage_days_and_interactions(db_session, people, mipt_deal) -> None:
    report = await _detail(db_session, people["head"])
    assert report["period"]["from"] == "2026-09-21" and report["period"]["to"] == "2026-10-20"
    request = _request(report, mipt_deal["deal"])
    assert request is not None
    assert request["stage"]["code"] == "negotiation"
    assert request["days_on_stage"] == 20
    assert request["interactions"] == {
        "transitions": 1,
        "returns": 0,
        "comments": 1,
        "files": 1,
        "activities": 0,
        "total": 3,
    }
    last = datetime.fromisoformat(request["last_contact_at"])
    assert abs(last - ago(5)) < timedelta(seconds=1)
    node = _node(_row(report, people["kam1"]), mipt_deal["university"].id)
    assert node["assigned"] is True and node["stages"][0]["code"] == "negotiation"


async def test_case2_stage_restored_from_history(db_session, people, mipt_deal) -> None:
    report = await _detail(
        db_session,
        people["head"],
        date_from=(NOW - timedelta(days=90)).date(),
        date_to=(NOW - timedelta(days=45)).date(),
    )
    request = _request(report, mipt_deal["deal"])
    assert request is not None
    assert request["stage"]["code"] == "first_contact"  # этап на 05.09, не сегодняшний
    assert request["days_on_stage"] == 15
    assert request["interactions"]["transitions"] == 1
    assert request["interactions"]["comments"] == 0


# --- кейсы 3–6: кто попадает в отчёт и что считается активностью --------------------------


async def test_case3_closed_before_period_excluded_unless_touched(
    db_session, workflows, people
) -> None:
    wf, stages = workflows["b2b"]
    kam1 = people["kam1"]
    path = [
        ("new", ago(130)),
        ("first_contact", ago(125)),
        ("negotiation", ago(120)),
        ("contract_prep", ago(115)),
        ("contract_signed", ago(110)),
        ("in_learning", ago(105)),
        ("completed", ago(100)),
    ]
    deal = await factories.make_deal(
        db_session, workflow=wf, stages=stages, responsible=kam1, created_at=ago(130), path=path
    )
    assert _request(await _detail(db_session, people["head"]), deal) is None
    await factories.make_comment(db_session, deal=deal, author=kam1, at=ago(3))
    request = _request(await _detail(db_session, people["head"]), deal)
    assert request is not None and request["closed_in_period"] is None
    assert request["interactions"]["comments"] == 1


async def test_case4_won_in_period(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam2 = people["kam2"]
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(60),
        path=[("new", ago(60)), ("in_learning", ago(30)), ("completed", ago(10))],
    )
    report = await _detail(db_session, people["head"], kam_ids=(kam2.id,))
    request = _request(report, deal)
    assert request["closed_in_period"] == "won" and request["stage"]["is_terminal"] is True
    row = _row(report, kam2)
    assert row["won"] == 1 and row["lost"] == 0 and row["conversion"] == 100.0
    assert report["totals"]["won"] == 1


async def test_case5_foreign_and_system_actions_are_not_activity(
    db_session, workflows, people
) -> None:
    wf, stages = workflows["b2b"]
    kam1, head = people["kam1"], people["head"]
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(90),
        path=[
            ("new", ago(90)),
            ("first_contact", ago(80)),
            # перевод схемы (F3): актор NULL, kind=system_migration
            ("negotiation", ago(7), None, "system_migration"),
            # переход руководителя в заявке КАМа — не активность КАМа
            ("contract_prep", ago(6), head),
        ],
    )
    await factories.make_comment(db_session, deal=deal, author=head, at=ago(4))
    request = _request(await _detail(db_session, head, kam_ids=(kam1.id,)), deal)
    assert request is not None
    assert request["stage"]["code"] == "contract_prep"  # этап меняется любыми записями
    assert request["interactions"]["total"] == 0
    # последний собственный контакт — ручной переход КАМа −80 д, а не действия системы/руководителя
    last = datetime.fromisoformat(request["last_contact_at"])
    assert abs(last - ago(80)) < timedelta(seconds=1)


async def test_case6_deleted_and_archived_are_ignored(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam1 = people["kam1"]
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(200),
        path=[("new", ago(200)), ("first_contact", ago(150))],
    )
    await factories.make_comment(db_session, deal=deal, author=kam1, at=ago(2), deleted=True)
    await factories.make_file(
        db_session, entity_id=deal.id, uploaded_by=kam1, at=ago(2), deleted=True
    )
    archived = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(5),
        status="archived",
    )
    await factories.make_comment(db_session, deal=archived, author=kam1, at=ago(1))
    report = await _detail(db_session, people["head"], kam_ids=(kam1.id,))
    request = _request(report, deal)
    assert request is not None  # висит в работе — в отчёте, но без взаимодействий
    assert request["interactions"]["total"] == 0
    assert _request(report, archived) is None


# --- кейс 7: зона видимости (через API, область RBAC report_scope) ------------------------


async def test_case7_scope(api, db_session, workflows, people, mipt_deal) -> None:
    wf, stages = workflows["b2b"]
    kam1, kam2, head = people["kam1"], people["kam2"], people["head"]
    other = await factories.make_deal(
        db_session, workflow=wf, stages=stages, responsible=kam2, created_at=ago(3)
    )
    url = f"{API}/reports/kam-activity"

    response = await api.as_user(kam1).get(url)
    assert response.status_code == 200, response.text
    own_report = response.json()
    KamActivityReportOut.model_validate(own_report)
    assert [row["kam_id"] for row in own_report["rows"]] == [str(kam1.id)]
    assert own_report["params"]["kam_id"] == [str(kam1.id)]
    assert own_report["scope"]["own_only"] is True
    assert _request(own_report, other) is None
    assert _request(own_report, mipt_deal["deal"]) is not None

    for path in (url, f"{url}/summary"):
        forbidden = await api.as_user(kam1).get(path, params={"kam_id": str(kam2.id)})
        assert forbidden.status_code == 403
        assert forbidden.json()["error"]["message"] == "КАМ видит отчёт только по своим заявкам"
    # свой id в фильтре — можно
    mine = await api.as_user(kam1).get(url, params={"kam_id": str(kam1.id)})
    assert mine.status_code == 200

    for viewer in (head, people["observer"]):
        response = await api.as_user(viewer).get(url)
        assert response.status_code == 200, response.text
        report = response.json()
        ids = {row["kam_id"] for row in report["rows"]}
        assert {str(kam1.id), str(kam2.id)} <= ids
        assert report["scope"]["own_only"] is False
        assert _request(report, other) is not None and _request(report, mipt_deal["deal"])
        summary = await api.as_user(viewer).get(f"{url}/summary")
        assert summary.status_code == 200
        KamActivitySummaryOut.model_validate(summary.json())

    # период и фильтры валидирует сервис: 400 с машинными кодами
    bad = await api.as_user(head).get(url, params={"from": "2026-10-10", "to": "2026-10-01"})
    assert bad.status_code == 400
    assert bad.json()["error"]["details"][0]["code"] == "period_inverted"
    many = await api.as_user(head).get(
        url, params=[("kam_id", str(uuid.uuid4())) for _ in range(ka.MAX_FILTER_IDS + 1)]
    )
    assert many.status_code == 400
    assert many.json()["error"]["details"][0]["code"] == "too_many_values"
    assert (await api.anonymous().get(url)).status_code == 401


async def test_scope_restricted_workflow_types(
    api, db_session, workflows, people, mipt_deal
) -> None:
    """Своя роль «КАМ B2C» (права КАМа только в b2c): b2b — 403, b2b-заявки, вузы и
    мероприятия вне области; сама роль — КАМ отчёта в процессе b2c."""
    await make_role(
        db_session, code="kam_b2c", permissions=KAM_PERMISSIONS, workflow_types=("b2c",)
    )
    kam_c = await factories.make_user(
        db_session, roles=("kam_b2c",), full_name="Смирнова Анна Олеговна"
    )
    wf_c, stages_c = workflows["b2c"]
    own_b2c = await factories.make_deal(
        db_session, workflow=wf_c, stages=stages_c, responsible=kam_c, created_at=ago(4)
    )
    url = f"{API}/reports/kam-activity"

    # руководитель: КАМ B2C — строка отчёта по b2c (право requests:kam в процессе), не по b2b
    head = api.as_user(people["head"])
    b2c_rows = (await head.get(url, params={"workflow_type": "b2c"})).json()
    assert str(kam_c.id) in {row["kam_id"] for row in b2c_rows["rows"]}
    b2b_rows = (await head.get(url, params={"workflow_type": "b2b"})).json()
    assert str(kam_c.id) not in {row["kam_id"] for row in b2b_rows["rows"]}
    assert {str(people["kam1"].id), str(people["kam2"].id)} <= {
        row["kam_id"] for row in b2b_rows["rows"]
    }

    # мероприятие с вузом — сущность B2B: в области «только b2c» не считается
    await factories.make_activity(
        db_session, university=mipt_deal["university"], responsible=kam_c, slots=[ago(2)]
    )
    denied = await api.as_user(kam_c).get(url, params={"workflow_type": "b2b"})
    assert denied.status_code == 403
    assert denied.json()["error"]["message"] == "Процесс B2B недоступен вашей роли"

    response = await api.as_user(kam_c).get(url)
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["scope"] == {
        "own_only": True,
        "own_user_id": str(kam_c.id),
        "workflow_types": ["b2c"],
    }
    assert _request(report, own_b2c) is not None
    assert _request(report, mipt_deal["deal"]) is None  # b2b-заявка вне области
    row = _row(report, kam_c)
    assert [node["university_id"] for node in row["universities"]] == [None]
    assert report["totals"]["interactions"]["activities"] == 0

    summary = await api.as_user(kam_c).get(f"{url}/summary")
    assert summary.status_code == 200
    assert list(summary.json()["stages"]) == ["b2c"]

    # служебный вызов с явной областью (выгрузки, тесты) — то же правило 403
    with pytest.raises(ApiError) as excinfo:
        await ka.build_kam_activity(
            db_session,
            people["head"],
            KamActivityParams(workflow_type="b2b"),
            scope=DealScope(rules=(("b2c", None),)),
        )
    assert excinfo.value.status_code == 403 and "B2B" in excinfo.value.message


# --- кейсы 8–11: фильтры, закреплённые вузы, пороги, мероприятия -------------------------


async def test_case8_b2c_filter_and_university_filter(
    db_session, workflows, people, mipt_deal
) -> None:
    wf_c, stages_c = workflows["b2c"]
    kam1 = people["kam1"]
    client = await factories.make_counterparty(db_session, display_name="Смирнова О. В.")
    b2c = await factories.make_deal(
        db_session,
        workflow=wf_c,
        stages=stages_c,
        responsible=kam1,
        counterparty=client,
        created_at=ago(4),
    )
    await factories.make_activity(
        db_session, university=mipt_deal["university"], responsible=kam1, slots=[ago(2)]
    )
    report = await _detail(db_session, people["head"], workflow_type="b2c", kam_ids=(kam1.id,))
    row = _row(report, kam1)
    assert [node["university_id"] for node in row["universities"]] == [None]
    node = row["universities"][0]
    assert node["name"] == ka.B2C_NODE_NAME and node["interactions"]["activities"] == 0
    request = _request(report, b2c)
    assert request["client"] == "Смирнова О. В." and request["workflow_type"] == "b2c"
    assert _request(report, mipt_deal["deal"]) is None

    filtered = await _detail(
        db_session, people["head"], university_ids=(mipt_deal["university"].id,)
    )
    assert _request(filtered, mipt_deal["deal"]) is not None
    assert _request(filtered, b2c) is None
    assert {row["kam_id"] for row in filtered["rows"]} == {str(kam1.id)}


async def test_case9_assigned_university_without_activity(db_session, people, mipt_deal) -> None:
    kam1 = people["kam1"]
    itmo = await factories.make_university(
        db_session, name="Университет ИТМО", short_name="ИТМО", kam=kam1
    )
    report = await _detail(db_session, people["head"], kam_ids=(kam1.id,))
    row = _row(report, kam1)
    node = _node(row, itmo.id)
    assert node["assigned"] is True and node["interactions"]["total"] == 0
    assert node["requests"] == [] and node["requests_count"] == 0
    assert row["assigned_untouched"] == 1 and row["universities_assigned"] == 2
    assert row["universities"][0]["university_id"] == str(itmo.id)  # проблемные сначала
    summary = await _summary(db_session, people["head"], kam_ids=(kam1.id,))
    assert summary["kpi"]["universities_assigned"]["value"] == 2
    assert summary["kpi"]["assigned_touched"]["value"] == 1
    assert summary["kpi"]["coverage_pct"]["value"] == 50.0
    assert summary["untouched_assigned"] == ["ИТМО"]
    assert any("ИТМО" in line for line in summary["highlights"])


async def _without_own_threshold(db_session, stage: WorkflowStage) -> None:
    """Каскад порога проверяется на этапе без своего порога: сид может задать этапу
    обучения свой (длинный) порог — тест от этого не зависит."""
    stage.stuck_threshold_days = None
    await db_session.flush()


async def test_case10_stuck_thresholds_cascade(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam2 = people["kam2"]
    await _without_own_threshold(db_session, stages["in_learning"])
    on_new = await factories.make_deal(
        db_session, workflow=wf, stages=stages, responsible=kam2, created_at=ago(10)
    )
    on_learning = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(50),
        path=[("new", ago(50)), ("in_learning", ago(20))],
    )
    report = await _detail(db_session, people["head"], kam_ids=(kam2.id,))
    first, second = _request(report, on_new), _request(report, on_learning)
    assert (first["is_stuck"], first["threshold_days"], first["days_on_stage"]) == (True, 7, 10)
    assert (second["is_stuck"], second["threshold_days"]) == (True, 14)
    assert _row(report, kam2)["stuck"] == 2


async def test_case11_activity_slots(db_session, people, mipt_deal) -> None:
    kam1 = people["kam1"]
    await factories.make_activity(
        db_session,
        university=mipt_deal["university"],
        responsible=kam1,
        slots=[ago(5), NOW + timedelta(days=10)],
        name="День карьеры на Физтехе",
    )
    report = await _detail(db_session, people["head"], kam_ids=(kam1.id,))
    node = _node(_row(report, kam1), mipt_deal["university"].id)
    assert node["interactions"]["activities"] == 1
    assert node["interactions"]["total"] == 4  # 1 переход + 1 комментарий + 1 файл + 1 слот
    last = datetime.fromisoformat(node["last_contact_at"])
    assert abs(last - ago(5)) < timedelta(seconds=1)  # будущий слот не учтён


# --- кейс 12: сводка ----------------------------------------------------------------------


async def test_case12_summary(db_session, workflows, people, mipt_deal) -> None:
    kam1 = people["kam1"]
    await _without_own_threshold(db_session, workflows["b2b"][1]["in_learning"])
    report = await _summary(db_session, people["head"], kam_ids=(kam1.id,))
    assert report["report"] == "kam-activity-summary"
    assert report["previous_period"]["from"] == "2026-08-22"
    assert report["previous_period"]["to"] == "2026-09-20"
    assert report["period"]["granularity"] == "day"

    stages = {stage["code"]: stage for stage in report["stages"]["b2b"]}
    assert stages["new"]["reached"] == 1 and stages["negotiation"]["reached"] == 1
    assert stages["first_contact"]["avg_days"] == 40.0  # −60 → −20
    assert stages["negotiation"]["on_stage_at_end"] == 1
    assert stages["first_contact"]["share_of_first_pct"] == 100.0
    # порог каскадом: этап (new — 7) → воронка (in_learning без своего — 14)
    assert stages["new"]["threshold_days"] == 7 and stages["in_learning"]["threshold_days"] == 14
    assert stages["completed"]["threshold_days"] is None
    assert [s["is_terminal"] for s in report["stages"]["b2b"]][-2:] == [True, True]
    assert report["stages"]["b2b"][-2]["outcome"] == "won"

    series = {item["key"]: item for item in report["dynamics"]["series"]}
    assert len(series["comments"]["points"]) == 30  # непрерывная сетка дней
    assert sum(p["value"] for p in series["transitions"]["points"]) == 1
    assert sum(p["value"] for p in series["comments"]["points"]) == 1
    assert sum(p["value"] for p in series["files"]["points"]) == 1

    kpi = report["kpi"]
    assert kpi["interactions"] == {"value": 3, "prev": 0, "delta": 3}
    assert kpi["requests_count"]["value"] == 1
    assert report["by_kam"][0]["kam_id"] == str(kam1.id)
    assert report["by_kam"][0]["interactions_prev"] == 0
    assert 1 <= len(report["highlights"]) <= 5


async def test_summary_top_stuck_sorted_and_deleted_stage(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam2 = people["kam2"]
    long_stuck = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(70),
        path=[("new", ago(70)), ("first_contact", ago(40))],
        title="Долго висит",
    )
    await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(30),
        path=[("new", ago(30)), ("first_contact", ago(20))],
        title="Висит поменьше",
    )
    # мягкое удаление этапа (демо F3): заявка в истории остаётся на нём
    stage = stages["first_contact"]
    await db_session.execute(
        sa.update(WorkflowStage).where(WorkflowStage.id == stage.id).values(deleted_at=ago(1))
    )
    summary = await _summary(db_session, people["head"], kam_ids=(kam2.id,))
    titles = [item["title"] for item in summary["top_stuck"]]
    assert titles == ["Долго висит", "Висит поменьше"]
    deleted = next(s for s in summary["stages"]["b2b"] if s["code"] == "first_contact")
    assert deleted["deleted"] is True and deleted["name"] == "Первичный контакт (удалён)"
    detail = await _detail(db_session, people["head"], kam_ids=(kam2.id,))
    request = _request(detail, long_stuck)
    assert request["stage"]["name"] == "Первичный контакт (удалён)"
    assert request["stage"]["deleted"] is True


async def test_summary_previous_period_and_deltas(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam2 = people["kam2"]
    for days in (40, 35):  # прошлый период: две новые заявки и отказ
        await factories.make_deal(
            db_session,
            workflow=wf,
            stages=stages,
            responsible=kam2,
            created_at=ago(days),
            path=[("new", ago(days)), ("first_contact", ago(days - 1))],
        )
    await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(45),
        path=[("new", ago(45)), ("first_contact", ago(44)), ("rejected", ago(33))],
    )
    won = await factories.make_deal(  # текущий период: выигрыш
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam2,
        created_at=ago(25),
        path=[("new", ago(25)), ("in_learning", ago(15)), ("completed", ago(2))],
        amount=100000,
    )
    assert won.status == "won"
    summary = await _summary(db_session, people["head"], kam_ids=(kam2.id,))
    kpi = summary["kpi"]
    assert kpi["new_requests"] == {"value": 1, "prev": 3, "delta": -2}
    assert kpi["won"]["value"] == 1 and kpi["lost"]["prev"] == 1
    assert kpi["conversion"] == {"value": 100.0, "prev": 0.0, "delta": 100.0}
    assert kpi["new_amount"]["value"] == "100000.00"
    assert any(line.startswith("Выиграно 1, отказов 0") for line in summary["highlights"])


async def test_summary_talent_block(db_session, workflows, people, mipt_deal) -> None:
    kam1 = people["kam1"]
    program = Program(name="Практикум по ML", university_id=mipt_deal["university"].id)
    db_session.add(program)
    await db_session.flush()

    async def student(n: int, path: list[tuple[str, datetime]]) -> None:
        row = Student(
            full_name=f"Студентов{n} Иван",
            email=f"s{n}@example.com",
            email_hmac=f"{n:064d}",
            display_name=f"Студентов{n} И.",
            university_id=mipt_deal["university"].id,
            program_id=program.id,
            funnel_status=path[-1][0],
            status_changed_at=path[-1][1],
        )
        db_session.add(row)
        await db_session.flush()
        previous = None
        for status, at in path:
            db_session.add(
                StudentStatusHistory(
                    student_id=row.id, from_status=previous, to_status=status, created_at=at
                )
            )
            previous = status
        await db_session.flush()

    await student(1, [("candidate", ago(90)), ("studying", ago(60))])
    await student(2, [("candidate", ago(90)), ("studying", ago(60)), ("talent_pool", ago(5))])
    await student(3, [("candidate", ago(90)), ("talent_pool", ago(40))])
    summary = await _summary(db_session, people["head"], kam_ids=(kam1.id,))
    talent = summary["talent"]
    by_status = {item["key"]: item for item in talent["by_status"]}
    assert by_status["studying"]["value"] == 1 and by_status["talent_pool"]["value"] == 2
    assert by_status["talent_pool"]["entered"] == 1
    assert talent["talent_pool_added"] == {"value": 1, "prev": 1, "delta": 0}
    assert talent["programs"][0]["name"] == "Практикум по ML"
    assert talent["programs"][0]["university"] == "МФТИ"
    assert summary["kpi"]["talent_pool_added"]["value"] == 1

    # прошлый период: статус на его конец восстановлен по истории
    past = await _summary(
        db_session,
        people["head"],
        kam_ids=(kam1.id,),
        date_from=(NOW - timedelta(days=70)).date(),
        date_to=(NOW - timedelta(days=50)).date(),
    )
    past_status = {item["key"]: item["value"] for item in past["talent"]["by_status"]}
    assert past_status == {"candidate": 1, "studying": 2, "graduate": 0, "talent_pool": 0}


# --- кейс 13: выгрузка end-to-end (уровень сервиса) -------------------------------------


async def test_case13_export_files(db_session, workflows, people, mipt_deal) -> None:
    import openpyxl

    head = people["head"]
    filters = {"from": "2026-09-21", "to": "2026-10-20", "kam_id": [str(people["kam1"].id)]}
    data, rows, name = await build_kam_activity_file(
        db_session, user=head, filters=filters, export_format="xlsx"
    )
    assert name == "kam-activity_20260921-20261020.xlsx" and rows == 1
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert tuple(book.sheetnames) == SHEET_NAMES
    requests = book["Заявки"]
    assert requests["C2"].value == "МФТИ — корпоративная школа Python"
    assert isinstance(requests["F2"].value, datetime)  # «На этапе с» — настоящая дата

    pdf, pdf_rows, pdf_name = await build_kam_activity_file(
        db_session, user=head, filters=filters, export_format="pdf"
    )
    assert pdf.startswith(b"%PDF") and pdf_rows == 1 and pdf_name.endswith(".pdf")

    csv_data, csv_rows, _ = await build_kam_activity_file(
        db_session, user=head, filters=filters, export_format="csv"
    )
    assert csv_rows == 1 and "МФТИ — корпоративная школа Python" in csv_data.decode("utf-8-sig")

    with pytest.raises(ApiError) as excinfo:  # КАМ не выгрузит чужой отчёт
        await build_kam_activity_file(
            db_session,
            user=people["kam2"],
            filters={"kam_id": str(people["kam1"].id)},
            export_format="xlsx",
        )
    assert excinfo.value.status_code == 403


async def test_case13_export_jobs_end_to_end(
    api, db_session, workflows, people, mipt_deal, memory_storage
) -> None:
    """`POST /export/jobs {entity_type: report:kam-activity}`: MinIO, аудит, журнал, скачивание."""
    import openpyxl

    head, kam1 = people["head"], people["kam1"]
    filters = {"from": "2026-09-21", "to": "2026-10-20", "kam_id": [str(kam1.id)]}
    client = api.as_user(head)

    created = await client.post(
        f"{API}/export/jobs",
        json={"entity_type": "report:kam-activity", "format": "xlsx", "filters": filters},
    )
    assert created.status_code == 201, created.text
    job = created.json()
    assert job["status"] == "done" and job["rows"] == 1 and job["format"] == "xlsx"
    assert job["file"]["file_name"] == "kam-activity_20260921-20261020.xlsx"
    stored = await db_session.get(FileObject, uuid.UUID(job["id"]))
    assert stored.entity_type == "report" and stored.uploaded_by == head.id
    assert (stored.s3_bucket, stored.s3_key) in memory_storage.objects
    audit = (
        await db_session.execute(
            sa.select(AuditLog).where(
                AuditLog.action == "export.created", AuditLog.entity_id == stored.id
            )
        )
    ).scalar_one()
    assert audit.after == {"entity_type": "report:kam-activity", "format": "xlsx", "rows": 1}

    download = await client.get(f"{API}/export/jobs/{job['id']}/download")
    assert download.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(download.content))
    assert tuple(book.sheetnames) == SHEET_NAMES

    pdf = await client.post(
        f"{API}/export/jobs",
        json={"entity_type": "report:kam-activity", "format": "pdf", "filters": filters},
    )
    assert pdf.status_code == 201, pdf.text
    assert pdf.json()["file"]["content_type"] == "application/pdf"
    pdf_bytes = await client.get(f"{API}/export/jobs/{pdf.json()['id']}/download")
    assert pdf_bytes.content.startswith(b"%PDF")

    # КАМ не выгрузит чужой отчёт — тот же 403, что у API отчёта
    foreign = await api.as_user(people["kam2"]).post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:kam-activity",
            "format": "pdf",
            "filters": {"kam_id": str(kam1.id)},
        },
    )
    assert foreign.status_code == 403
    # чужая выгрузка не видна автору другого отчёта (files:others только у admin)
    hidden = await api.as_user(people["kam2"]).get(f"{API}/export/jobs/{job['id']}")
    assert hidden.status_code == 404


async def test_export_pdf_rules(api, db_session, people, memory_storage) -> None:
    """PDF: реестрам — 422; табличный PDF дашборда — за фичефлагом report_pdf."""
    client = api.as_user(people["head"])
    registry = await client.post(
        f"{API}/export/jobs", json={"entity_type": "requests", "format": "pdf"}
    )
    assert registry.status_code == 422
    assert "только для отчётов" in registry.json()["error"]["message"]

    dashboard = await client.post(
        f"{API}/export/jobs", json={"entity_type": "report:dashboard", "format": "pdf"}
    )
    assert dashboard.status_code == 201, dashboard.text
    body = (await client.get(f"{API}/export/jobs/{dashboard.json()['id']}/download")).content
    assert body.startswith(b"%PDF")

    db_session.add(
        AppSetting(
            key="feature_flags",
            value={"auto_ranking": False, "report_pdf": False, "llm_features": False},
        )
    )
    await db_session.flush()
    disabled = await client.post(
        f"{API}/export/jobs", json={"entity_type": "report:funnel", "format": "pdf"}
    )
    assert disabled.status_code == 422
    assert "report_pdf" in disabled.json()["error"]["message"]
    # PDF отчёта КАМов флагом не закрыт (D15)
    kam_pdf = await client.post(
        f"{API}/export/jobs", json={"entity_type": "report:kam-activity", "format": "pdf"}
    )
    assert kam_pdf.status_code == 201, kam_pdf.text
    unknown = await client.post(
        f"{API}/export/jobs", json={"entity_type": "requests", "format": "docx"}
    )
    assert unknown.status_code == 400  # Literal формата — валидация тела


# --- кейс 14: файлы заявки (§4.7) --------------------------------------------------------


async def test_case14_request_files(api, db_session, workflows, people, memory_storage) -> None:
    wf, stages = workflows["b2b"]
    kam1, kam2, head = people["kam1"], people["kam2"], people["head"]
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(40),
        path=[("new", ago(40)), ("first_contact", ago(35))],
    )
    before_activity = deal.last_activity_at
    url = f"{API}/requests/{deal.id}/files"

    empty = await api.as_user(kam1).get(url)
    assert empty.status_code == 200 and empty.json() == {"items": [], "total": 0}

    uploaded = await api.as_user(kam1).post(
        url, files={"file": ("КП_МФТИ.pdf", b"%PDF-1.4 demo", "application/pdf")}
    )
    assert uploaded.status_code == 201, uploaded.text
    item = uploaded.json()
    assert item["file_name"] == "КП_МФТИ.pdf" and item["uploaded_by"] == str(kam1.id)
    file_id = uuid.UUID(item["id"])
    stored = await db_session.get(FileObject, file_id)
    assert stored.entity_type == "deal" and stored.entity_id == deal.id
    assert memory_storage.objects[(stored.s3_bucket, stored.s3_key)][0] == b"%PDF-1.4 demo"
    await db_session.refresh(deal)
    assert deal.last_activity_at is not None and deal.last_activity_at > before_activity
    audit = (
        await db_session.execute(
            sa.select(AuditLog).where(
                AuditLog.action == "request.file_uploaded", AuditLog.entity_id == deal.id
            )
        )
    ).scalar_one()
    assert audit.actor_user_id == kam1.id and audit.after["file_id"] == str(file_id)

    listed = (await api.as_user(head).get(url)).json()
    assert listed["total"] == 1 and listed["items"][0]["id"] == str(file_id)

    # скачивание: КАМ видит заявки процесса (записи all) — файл отдаётся
    download = await api.as_user(kam2).get(f"{API}/files/{file_id}/download")
    assert download.status_code == 200 and download.content == b"%PDF-1.4 demo"
    # наблюдатель видит заявку, но файлы заявок ему не скачать (нет files:business)
    observer = await api.as_user(people["observer"]).get(f"{API}/files/{file_id}/download")
    assert observer.status_code == 403
    no_upload = await api.as_user(people["observer"]).post(url, files={"file": ("a.txt", b"x")})
    assert no_upload.status_code == 403  # нет requests:files
    # роль «только свои заявки»: чужая заявка и её файл не существуют — 404
    await make_role(db_session, code="kam_own", permissions=KAM_PERMISSIONS, record_scope="own")
    reset_role_catalog()  # каталог уже загружен запросами выше — перечитать с новой ролью
    own_only = await factories.make_user(db_session, roles=("kam_own",))
    hidden = await api.as_user(own_only).get(f"{API}/files/{file_id}/download")
    assert hidden.status_code == 404
    assert (await api.as_user(own_only).get(url)).status_code == 404

    # загрузка КАМом — взаимодействие в отчёте (дата файла — внутри периода отчёта)
    await db_session.execute(
        sa.update(FileObject).where(FileObject.id == file_id).values(created_at=ago(1))
    )
    report = await _detail(db_session, head, kam_ids=(kam1.id,))
    assert _request(report, deal)["interactions"]["files"] == 1

    # удаление: чужой файл — 403 (не автор и нет requests:write_any), автор — 204
    foreign = await api.as_user(kam2).delete(f"{url}/{file_id}")
    assert foreign.status_code == 403
    missing = await api.as_user(kam1).delete(f"{url}/{uuid.uuid4()}")
    assert missing.status_code == 404
    deleted = await api.as_user(kam1).delete(f"{url}/{file_id}")
    assert deleted.status_code == 204
    await db_session.refresh(stored)
    assert stored.deleted_at is not None
    assert (await api.as_user(kam1).get(url)).json()["total"] == 0
    assert (await api.as_user(kam1).delete(f"{url}/{file_id}")).status_code == 404
    gone = await api.as_user(kam1).get(f"{API}/files/{file_id}/download")
    assert gone.status_code == 404
    deleted_audit = (
        await db_session.execute(
            sa.select(sa.func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "request.file_deleted", AuditLog.entity_id == deal.id)
        )
    ).scalar_one()
    assert deleted_audit == 1

    # руководитель удаляет чужие файлы (requests:write_any в процессе)
    second = await api.as_user(kam1).post(
        url,
        files={"file": ("Протокол.docx", b"PK\x03\x04 docx")},  # сигнатура ZIP/DOCX
    )
    assert second.status_code == 201
    by_head = await api.as_user(head).delete(f"{url}/{second.json()['id']}")
    assert by_head.status_code == 204


# --- кейс 15: D16 — «Завершено за период» дашборда по истории --------------------------


async def test_case15_dashboard_completed_uses_history(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam1, head = people["kam1"], people["head"]
    won_recent = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(60),
        path=[("new", ago(60)), ("in_learning", ago(30)), ("completed", ago(10))],
    )
    won_long_ago = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(200),
        path=[("new", ago(200)), ("in_learning", ago(150)), ("completed", ago(100))],
    )
    # правка карточки сдвигает updated_at: недавняя «старая», давняя — «свежая»
    await db_session.execute(
        sa.update(Deal).where(Deal.id == won_recent.id).values(updated_at=ago(100))
    )
    await db_session.execute(
        sa.update(Deal).where(Deal.id == won_long_ago.id).values(updated_at=ago(5))
    )
    params = ReportParams(
        workflow_type="b2b",
        date_from=(NOW - timedelta(days=30)).date(),
        date_to=TODAY,
    )
    report = await build_report(db_session, "dashboard", head, params)
    assert report["kpi"]["completed"]["value"] == 1
    assert report["kpi"]["conversion"]["value"] == 100.0
    # отчёт КАМов за тот же период — тот же «Выиграно»
    detail = await _detail(
        db_session, head, date_from=params.date_from, date_to=params.date_to, workflow_type="b2b"
    )
    assert detail["totals"]["won"] == 1


# --- миграция 0005 и OpenAPI ---------------------------------------------------------------


async def test_kam_activity_indexes_exist(db_session) -> None:
    rows = (
        await db_session.execute(
            sa.text(
                "SELECT indexname, indexdef FROM pg_indexes WHERE schemaname = 'public' "
                "AND indexname IN ('ix_deal_responsible_created', 'ix_activity_university', "
                "'ix_activity_responsible')"
            )
        )
    ).all()
    defs = dict(rows)
    assert set(defs) == {
        "ix_deal_responsible_created",
        "ix_activity_university",
        "ix_activity_responsible",
    }
    assert "(responsible_user_id, created_at)" in defs["ix_deal_responsible_created"]
    assert "archived" in defs["ix_deal_responsible_created"]


async def test_openapi_documents_report_schemas(api) -> None:
    schema = api.app.openapi()
    detail = schema["paths"]["/api/v1/reports/kam-activity"]["get"]
    summary = schema["paths"]["/api/v1/reports/kam-activity/summary"]["get"]
    ok = detail["responses"]["200"]["content"]["application/json"]["schema"]
    assert ok["$ref"].endswith("/KamActivityReportOut")
    assert summary["responses"]["200"]["content"]["application/json"]["schema"]["$ref"].endswith(
        "/KamActivitySummaryOut"
    )
    names = {param["name"] for param in detail["parameters"]}
    assert {"from", "to", "workflow_type", "kam_id", "university_id"} <= names
    files = schema["paths"]["/api/v1/requests/{request_id}/files"]
    assert {"get", "post"} <= set(files)
    assert "delete" in schema["paths"]["/api/v1/requests/{request_id}/files/{file_id}"]


# --- граница периода и SQL-специфика PostgreSQL -------------------------------------------


async def test_period_boundary_is_moscow_midnight(db_session, workflows, people) -> None:
    wf, stages = workflows["b2b"]
    kam1 = people["kam1"]
    deal = await factories.make_deal(
        db_session,
        workflow=wf,
        stages=stages,
        responsible=kam1,
        created_at=ago(200),
        path=[("new", ago(200)), ("first_contact", ago(150))],
    )
    edge = datetime(2026, 10, 1, 0, 0, tzinfo=ka.REPORT_TZ)  # 30.09 21:00 UTC
    await factories.make_comment(db_session, deal=deal, author=kam1, at=edge)
    await factories.make_comment(db_session, deal=deal, author=kam1, at=edge - timedelta(seconds=1))
    october = await _detail(
        db_session,
        people["head"],
        kam_ids=(kam1.id,),
        date_from=date(2026, 10, 1),
        date_to=date(2026, 10, 10),
    )
    assert _request(october, deal)["interactions"]["comments"] == 1
    september = await _detail(
        db_session,
        people["head"],
        kam_ids=(kam1.id,),
        date_from=date(2026, 9, 21),
        date_to=date(2026, 9, 30),
    )
    assert _request(september, deal)["interactions"]["comments"] == 1


async def test_many_deals_single_array_param(db_session, workflows, people) -> None:
    """deal_ids — один параметр-массив; динамика по неделям без GroupingError."""
    wf, stages = workflows["b2b"]
    kam1 = people["kam1"]
    for index in range(30):
        deal = await factories.make_deal(
            db_session,
            workflow=wf,
            stages=stages,
            responsible=kam1,
            created_at=ago(100 + index),
            path=[("new", ago(100 + index)), ("first_contact", ago(50 - index))],
        )
        await factories.make_comment(db_session, deal=deal, author=kam1, at=ago(index + 1))
    summary = await _summary(
        db_session,
        people["head"],
        kam_ids=(kam1.id,),
        date_from=date(2026, 7, 1),
        date_to=date(2026, 10, 20),
    )
    assert summary["period"]["granularity"] == "week"
    comments = next(s for s in summary["dynamics"]["series"] if s["key"] == "comments")
    assert sum(point["value"] for point in comments["points"]) == 30
    assert all(date.fromisoformat(p["date"]).weekday() == 0 for p in comments["points"])
    history_rows = (
        await db_session.execute(sa.select(sa.func.count()).select_from(DealStageHistory))
    ).scalar_one()
    assert history_rows >= 60


# --- демо-сид активности (tools/seed_activity.py) ----------------------------------------


async def test_seed_activity_history_is_idempotent(db_session, memory_storage) -> None:
    from app.tools import seed
    from app.tools.seed_activity import KAMS, seed_activity_history

    users = await seed.seed_users(db_session)
    admin = users["admin@demo"]
    workflows = await seed.seed_workflows(db_session, admin)
    interaction_types = await seed.seed_interaction_types(db_session, admin)
    await seed.seed_app_settings(db_session)
    universities = await seed.seed_universities(db_session, users)
    products, programs = await seed.seed_catalog(db_session, universities, users)
    contracts = await seed.seed_contracts(db_session, universities, products, users)
    counterparties = await seed.seed_counterparties(db_session)
    await seed.seed_deals(
        db_session,
        workflows,
        universities,
        products,
        programs,
        contracts,
        counterparties,
        interaction_types,
        users,
    )
    now = factories.utcnow()
    stats = await seed_activity_history(
        db_session, workflows, universities, products, programs, users, now=now
    )
    assert stats["deals"] >= 90 and stats["activities"] >= 20 and stats["students"] == 36
    assert stats["files"] > 0 and memory_storage.objects  # файлы легли в хранилище
    # договоры вузов, дошедших до «Договор подписан»: МФТИ уже с договором базового сида
    assert stats["contracts"] >= 5
    active_contracts = (
        await db_session.execute(
            sa.select(sa.func.count(sa.distinct(Contract.university_id))).where(
                Contract.status == "active"
            )
        )
    ).scalar_one()
    assert active_contracts == stats["contracts"] + 1
    again = await seed_activity_history(
        db_session, workflows, universities, products, programs, users, now=now
    )
    assert again == {
        **again,
        "deals": 0,
        "contracts": 0,
        "base_comments": 0,
        "activities": 0,
        "students": 0,
    }
    assert again["files"] == 0

    kam_count = (
        await db_session.execute(
            sa.select(sa.func.count()).select_from(AppUser).where(AppUser.roles.any("kam"))
        )
    ).scalar_one()
    assert kam_count == len(KAMS) == 20

    head = users["head@demo"]
    monkey_now = NOW  # отчёт «на 20.10» поверх данных относительно реального «сейчас»
    assert ka._now() == monkey_now
    report = await _detail(db_session, head)
    assert len(report["rows"]) >= 19  # все активные КАМы + уволившийся, если попал в период
    totals = report["totals"]
    assert totals["interactions"]["total"] > 100 and totals["interactions"]["files"] > 0
    assert 0 < totals["coverage_pct"] < 100  # ИТМО и ТГУ закреплены, но без контакта
    summary = await _summary(db_session, head)
    assert "ИТМО" in summary["untouched_assigned"] and "ТГУ" in summary["untouched_assigned"]
    assert len(summary["highlights"]) == 5
    assert summary["talent"]["students_total"] >= 36  # студенты сида активности
