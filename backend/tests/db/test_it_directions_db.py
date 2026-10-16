"""Справочник ИТ-направлений (P1-1) на настоящем PostgreSQL.

Каталог `it_direction` с правами из каталога F2 (читают все, правит `admin:dictionaries`),
направление программы (N:1) и продукта (M:N), единое правило «направление записи»
(`crm.directions.direction_match`/`direction_names`) в реестре программ, на дашборде, в
«Активности КАМов» и «Взаимодействиях за период» (фильтр и колонка), выгрузка, сид.

Мир: DevOps (sort 10), QA (20), «Анализ данных» (30);
- продукт «Облако» — DevOps + QA, «BI» — анализ данных, «Без направления» — ничего;
- программа «DevOps-практикум» (Облако, своё направление DevOps), «Тестирование ПО»
  (Облако, своего нет → направления продукта DevOps + QA), «Аналитика» (BI, своё —
  анализ данных);
- заявки: A — программа «DevOps-практикум» (DevOps, хотя продукт ещё и QA), B — только
  продукт «Облако» (DevOps + QA), C — «Аналитика», D — продукт без направлений,
  E — «Тестирование ПО» (DevOps + QA по продукту).
"""

from __future__ import annotations

import inspect
import io
import uuid
from datetime import datetime, timedelta
from typing import Any

import pytest
import sqlalchemy as sa

from app.models import ItDirection, Product, ProductDirection, Program
from app.modules.reporting import kam_activity as ka
from app.modules.reporting.interactions import InteractionsParams, build_interactions
from app.modules.reporting.kam_activity import KamActivityParams
from app.modules.reporting.schemas import InteractionsReportOut, KamActivityReportOut
from app.modules.reporting.service import ReportParams, build_report

from . import factories

API = "/api/v1"
NOW = datetime(2026, 10, 20, 12, 0, tzinfo=ka.REPORT_TZ)


def ago(days: float) -> datetime:
    return NOW - timedelta(days=days)


@pytest.fixture(autouse=True)
def _fixed_now(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ka, "_now", lambda: NOW)


@pytest.fixture
async def world(db_session, workflows) -> dict[str, Any]:
    wf, stages = workflows["b2b"]
    make = factories.make_user
    head = await make(db_session, roles=("head_kam",), full_name="Морозов Дмитрий")
    kam1 = await make(db_session, roles=("kam",), full_name="Петров Константин")
    kam2 = await make(db_session, roles=("kam",), full_name="Иванова Мария")
    mipt = await factories.make_university(
        db_session, name="Московский физико-технический институт", short_name="МФТИ", kam=kam1
    )
    itmo = await factories.make_university(
        db_session, name="Университет ИТМО", short_name="ИТМО", kam=kam2
    )
    devops = ItDirection(code="devops", name="DevOps", sort_order=10)
    qa = ItDirection(code="qa", name="Тестирование (QA)", sort_order=20)
    data = ItDirection(code="data-analysis", name="Анализ данных", sort_order=30)
    db_session.add_all([devops, qa, data])
    await db_session.flush()

    cloud = Product(code="cloud", name="Облако")
    bi = Product(code="bi", name="BI")
    plain = Product(code="plain", name="Без направления")
    db_session.add_all([cloud, bi, plain])
    await db_session.flush()
    db_session.add_all(
        [
            ProductDirection(product_id=cloud.id, direction_id=devops.id),
            ProductDirection(product_id=cloud.id, direction_id=qa.id),
            ProductDirection(product_id=bi.id, direction_id=data.id),
        ]
    )
    devops_program = Program(name="DevOps-практикум", product_id=cloud.id, direction_id=devops.id)
    qa_program = Program(name="Тестирование ПО", product_id=cloud.id)
    data_program = Program(name="Аналитика", product_id=bi.id, direction_id=data.id)
    db_session.add_all([devops_program, qa_program, data_program])
    await db_session.flush()

    async def deal(title: str, kam, university, **refs: Any):  # noqa: ANN001, ANN202
        return await factories.make_deal(
            db_session,
            workflow=wf,
            stages=stages,
            responsible=kam,
            university=university,
            title=title,
            created_at=ago(12),
            path=[("new", ago(12)), ("first_contact", ago(6))],
            **refs,
        )

    deal_a = await deal("A", kam1, mipt, product_id=cloud.id, program_id=devops_program.id)
    deal_b = await deal("B", kam1, mipt, product_id=cloud.id)
    deal_c = await deal("C", kam2, itmo, product_id=bi.id, program_id=data_program.id)
    deal_d = await deal("D", kam2, itmo, product_id=plain.id)
    deal_e = await deal("E", kam1, mipt, program_id=qa_program.id)  # продукт — от программы

    devops_event = await factories.make_activity(
        db_session, university=mipt, responsible=kam1, slots=[ago(3)]
    )
    devops_event.program_id = devops_program.id
    await factories.make_activity(db_session, university=mipt, responsible=kam1, slots=[ago(2)])
    await db_session.flush()
    return {
        "head": head,
        "kam1": kam1,
        "kam2": kam2,
        "mipt": mipt,
        "itmo": itmo,
        "devops": devops,
        "qa": qa,
        "data": data,
        "cloud": cloud,
        "bi": bi,
        "plain": plain,
        "devops_program": devops_program,
        "qa_program": qa_program,
        "data_program": data_program,
        "A": deal_a,
        "B": deal_b,
        "C": deal_c,
        "D": deal_d,
        "E": deal_e,
        "stages": stages,
    }


def _titles(world: dict[str, Any], ids: set[str]) -> set[str]:
    by_id = {str(world[key].id): key for key in "ABCDE"}
    return {by_id[item] for item in ids if item in by_id}


# --- CRUD справочника и права -----------------------------------------------------------------


async def test_directions_crud_and_permissions(api, admin_user, world) -> None:
    admin = api.as_user(admin_user)
    kam = api.as_user(world["kam1"])
    url = f"{API}/it-directions"

    listed = await kam.get(url)  # читают все роли
    assert listed.status_code == 200, listed.text
    items = listed.json()["items"]
    assert [item["code"] for item in items] == ["devops", "qa", "data-analysis"]  # sort_order
    devops = next(item for item in items if item["code"] == "devops")
    assert devops["programs_count"] == 1 and devops["products_count"] == 1
    assert (await api.anonymous().get(url)).status_code == 401
    assert (
        await kam.post(url, json={"code": "gamedev", "name": "Разработка игр"})
    ).status_code == 403
    assert (
        await api.as_user(world["head"]).post(url, json={"code": "x1", "name": "X1"})
    ).status_code == 403

    created = await admin.post(
        url, json={"code": "gamedev", "name": "Разработка игр", "sort_order": 90}
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["is_active"] is True and body["programs_count"] == 0
    duplicate_code = await admin.post(url, json={"code": "GameDev", "name": "Геймдев"})
    assert duplicate_code.status_code == 409
    duplicate_name = await admin.post(url, json={"code": "games", "name": "разработка ИГР"})
    assert duplicate_name.status_code == 409
    bad_code = await admin.post(url, json={"code": "с пробелом", "name": "Что-то"})
    assert bad_code.status_code == 400  # схема тела: код без пробелов

    item_url = f"{url}/{body['id']}"
    stale = await admin.patch(item_url, json={"version": body["version"] + 1, "name": "Игры"})
    assert stale.status_code == 409
    renamed = await admin.patch(
        item_url, json={"version": body["version"], "name": "Разработка игр и XR"}
    )
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["name"] == "Разработка игр и XR"
    assert (await kam.get(item_url)).json()["name"] == "Разработка игр и XR"

    # направление активной программы не деактивируется (программа — одно направление)
    devops_url = f"{url}/{world['devops'].id}"
    assert (await admin.delete(devops_url)).status_code == 409
    current = (await admin.get(devops_url)).json()
    blocked = await admin.patch(
        devops_url, json={"version": current["version"], "is_active": False}
    )
    assert blocked.status_code == 409
    # связи с продуктами деактивации не мешают; удаление — мягкое
    qa_url = f"{url}/{world['qa'].id}"
    assert (await kam.delete(qa_url)).status_code == 403
    assert (await admin.delete(qa_url)).status_code == 200
    active = (await kam.get(url, params={"is_active": "true"})).json()["items"]
    assert "qa" not in {item["code"] for item in active}
    assert (await kam.get(qa_url)).json()["products_count"] == 1  # связь сохранена
    searched = (await kam.get(url, params={"search": "devo"})).json()
    assert [item["code"] for item in searched["items"]] == ["devops"]


async def test_program_direction_and_inherited(api, admin_user, world) -> None:
    head = api.as_user(world["head"])
    response = await head.get(f"{API}/programs/{world['qa_program'].id}")
    assert response.status_code == 200, response.text
    out = response.json()
    assert out["direction_id"] is None and out["direction"] is None
    assert [d["code"] for d in out["product_directions"]] == ["devops", "qa"]
    own = (await head.get(f"{API}/programs/{world['devops_program'].id}")).json()
    assert own["direction"]["name"] == "DevOps"

    # фильтр реестра — по тому же правилу, что отчёты
    by_qa = await head.get(f"{API}/programs", params={"direction_id": str(world["qa"].id)})
    assert {item["name"] for item in by_qa.json()["items"]} == {"Тестирование ПО"}
    by_devops = await head.get(
        f"{API}/programs",
        params=[("direction_id", str(world["devops"].id)), ("direction_id", str(world["data"].id))],
    )
    assert {item["name"] for item in by_devops.json()["items"]} == {
        "DevOps-практикум",
        "Тестирование ПО",
        "Аналитика",
    }

    created = await head.post(
        f"{API}/programs",
        json={
            "name": "QA-автоматизатор",
            "product_id": str(world["cloud"].id),
            "direction_id": str(world["qa"].id),
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["direction"]["code"] == "qa"
    unknown = await head.post(
        f"{API}/programs", json={"name": "Неизвестное", "direction_id": str(uuid.uuid4())}
    )
    assert unknown.status_code == 400
    assert unknown.json()["error"]["details"][0]["code"] == "lookup_not_found"

    program = created.json()
    cleared = await head.patch(
        f"{API}/programs/{program['id']}",
        json={"version": program["version"], "direction_id": None},
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["direction"] is None
    assert [d["code"] for d in cleared.json()["product_directions"]] == ["devops", "qa"]

    # деактивированное направление нельзя выбрать заново
    gamedev = await api.as_user(admin_user).post(
        f"{API}/it-directions", json={"code": "gamedev", "name": "Разработка игр"}
    )
    await api.as_user(admin_user).delete(f"{API}/it-directions/{gamedev.json()['id']}")
    inactive = await head.patch(
        f"{API}/programs/{program['id']}",
        json={"version": cleared.json()["version"], "direction_id": gamedev.json()["id"]},
    )
    assert inactive.status_code == 400
    assert inactive.json()["error"]["details"][0]["code"] == "inactive"


async def test_product_directions_m2m(api, db_session, world) -> None:
    head = api.as_user(world["head"])
    created = await head.post(
        f"{API}/products",
        json={
            "code": "mobile-kit",
            "name": "Мобильный SDK",
            "direction_ids": [str(world["qa"].id), str(world["devops"].id)],
        },
    )
    assert created.status_code == 201, created.text
    product = created.json()
    assert [d["code"] for d in product["directions"]] == ["devops", "qa"]  # по sort_order
    assert set(product["direction_ids"]) == {str(world["devops"].id), str(world["qa"].id)}

    replaced = await head.patch(
        f"{API}/products/{product['id']}",
        json={"version": product["version"], "direction_ids": [str(world["data"].id)]},
    )
    assert replaced.status_code == 200, replaced.text
    assert [d["code"] for d in replaced.json()["directions"]] == ["data-analysis"]
    assert replaced.json()["version"] != product["version"]  # смена набора — новая версия
    stale = await head.patch(
        f"{API}/products/{product['id']}",
        json={"version": product["version"], "direction_ids": []},
    )
    assert stale.status_code == 409
    untouched = await head.patch(
        f"{API}/products/{product['id']}",
        json={"version": replaced.json()["version"], "description": "SDK для вузов"},
    )
    assert [d["code"] for d in untouched.json()["directions"]] == ["data-analysis"]

    by_direction = await head.get(f"{API}/products", params={"direction_id": str(world["data"].id)})
    assert {item["code"] for item in by_direction.json()["items"]} == {"bi", "mobile-kit"}
    unknown = await head.post(
        f"{API}/products",
        json={"code": "x", "name": "X", "direction_ids": [str(uuid.uuid4())]},
    )
    assert unknown.status_code == 400
    # удаление продукта физически снимает связи каскадом (на уровне БД)
    await db_session.execute(sa.delete(Product).where(Product.code == "mobile-kit"))
    left = (
        await db_session.execute(
            sa.select(sa.func.count())
            .select_from(ProductDirection)
            .where(ProductDirection.direction_id == world["data"].id)
        )
    ).scalar_one()
    assert left == 1  # осталась связь «BI»


# --- дашборд --------------------------------------------------------------------------------


async def test_dashboard_sections_filter_by_direction(db_session, world) -> None:
    head = world["head"]

    async def funnel_total(**kwargs: Any) -> int:
        report = await build_report(db_session, "funnel", head, ReportParams(**kwargs))
        return sum(item["value"] for item in report["series"])

    assert await funnel_total() == 5
    assert await funnel_total(direction_id=[world["qa"].id]) == 2  # B и E, не A
    assert await funnel_total(direction_id=[world["devops"].id]) == 3  # A, B, E
    assert await funnel_total(direction_id=[world["data"].id]) == 1
    assert await funnel_total(direction_id=[world["qa"].id, world["data"].id]) == 3

    dashboard = await build_report(
        db_session, "dashboard", head, ReportParams(direction_id=world["devops"].id)
    )
    assert dashboard["params"]["direction_id"] == [str(world["devops"].id)]
    assert dashboard["kpi"]["active"]["value"] == 3
    workload = {row["kam_id"]: row for row in dashboard["kam_workload"]["rows"]}
    assert workload[str(world["kam1"].id)]["active_requests"] == 3
    assert workload[str(world["kam2"].id)]["active_requests"] == 0
    demand = {row["name"] for row in dashboard["programs_demand"]["rows"]}
    assert demand == {"DevOps-практикум", "Тестирование ПО"}

    # без фильтра ключ direction_id в эхо параметров не появляется (прежний контракт)
    plain = await build_report(db_session, "funnel", head, ReportParams())
    assert "direction_id" not in plain["params"]

    # выгрузка: filters → те же параметры
    params = ReportParams.from_filters({"direction_id": [str(world["qa"].id)]})
    assert params.direction_ids == (world["qa"].id,)
    from app.core.errors import ApiError

    with pytest.raises(ApiError):
        ReportParams.from_filters({"direction_id": ["не-uuid"]})


# --- «Активность КАМов» -------------------------------------------------------------------------


def _kam_request_ids(report: dict[str, Any]) -> set[str]:
    return {
        request["request_id"]
        for row in report["rows"]
        for node in row["universities"]
        for request in node["requests"]
    }


async def test_kam_activity_direction_filter(db_session, world) -> None:
    head = world["head"]
    full = await ka.build_kam_activity(db_session, head, KamActivityParams())
    assert _titles(world, _kam_request_ids(full)) == {"A", "B", "C", "D", "E"}
    assert "directions" not in full and "direction_id" not in full["params"]

    qa = await ka.build_kam_activity(
        db_session, head, KamActivityParams(direction_ids=(world["qa"].id,))
    )
    KamActivityReportOut.model_validate(qa)
    assert _titles(world, _kam_request_ids(qa)) == {"B", "E"}
    # КАМ без заявок направления строкой не выводится; закреплённые вузы — только с заявками
    assert {row["kam_id"] for row in qa["rows"]} == {str(world["kam1"].id)}
    assert qa["directions"] == [
        {"id": str(world["qa"].id), "code": "qa", "name": "Тестирование (QA)"}
    ]
    assert qa["params"]["direction_id"] == [str(world["qa"].id)]

    data = await ka.build_kam_activity(
        db_session, head, KamActivityParams(direction_ids=(world["data"].id,))
    )
    assert _titles(world, _kam_request_ids(data)) == {"C"}
    universities = {node["short_name"] for row in data["rows"] for node in row["universities"]}
    assert universities == {"ИТМО"}

    # мероприятия — по направлению их программы: «DevOps-практикум» входит, без программы — нет
    devops = await ka.build_kam_activity(
        db_session, head, KamActivityParams(direction_ids=(world["devops"].id,))
    )
    mipt = next(
        node
        for row in devops["rows"]
        for node in row["universities"]
        if node["short_name"] == "МФТИ"
    )
    assert mipt["interactions"]["activities"] == 1
    assert _titles(world, _kam_request_ids(devops)) == {"A", "B", "E"}

    summary = await ka.build_kam_activity_summary(
        db_session, head, KamActivityParams(direction_ids=(world["devops"].id,))
    )
    assert summary["directions"][0]["name"] == "DevOps"
    assert summary["kpi"]["requests_count"]["value"] == 3

    too_many = KamActivityParams(
        direction_ids=tuple(uuid.uuid4() for _ in range(ka.MAX_FILTER_IDS + 1))
    )
    from app.core.errors import ApiError

    with pytest.raises(ApiError):
        await ka.build_kam_activity(db_session, head, too_many)
    parsed = KamActivityParams.from_filters({"direction": str(world["qa"].id)})
    assert parsed.direction_ids == (world["qa"].id,)


# --- «Взаимодействия за период»: колонка и фильтр -------------------------------------------------


async def test_interactions_column_and_filter(api, db_session, world) -> None:
    head = world["head"]
    report = await build_interactions(db_session, head, InteractionsParams())
    InteractionsReportOut.model_validate(report)
    by_title = {row["title"]: row["it_direction"] for row in report["rows"]}
    assert by_title == {
        "A": "DevOps",  # у программы своё направление
        "B": "DevOps, Тестирование (QA)",  # только продукт — все его направления
        "C": "Анализ данных",
        "D": None,  # у продукта направлений нет — «—»
        "E": "DevOps, Тестирование (QA)",  # программа без своего — направления её продукта
    }

    filtered = await build_interactions(
        db_session,
        head,
        InteractionsParams(base=KamActivityParams(direction_ids=(world["qa"].id,))),
    )
    rows = filtered["rows"]
    assert {row["title"] for row in rows} == {"B", "E"}
    assert all("Тестирование (QA)" in row["it_direction"] for row in rows)
    assert filtered["directions"][0]["code"] == "qa"

    response = await api.as_user(world["head"]).get(
        f"{API}/reports/interactions", params={"direction_id": str(world["data"].id)}
    )
    assert response.status_code == 200, response.text
    assert [row["title"] for row in response.json()["rows"]] == ["C"]
    own = await api.as_user(world["kam1"]).get(
        f"{API}/reports/interactions", params={"direction_id": str(world["data"].id)}
    )
    assert own.status_code == 200 and own.json()["rows"] == []  # «C» — заявка другого КАМа


async def test_interactions_export_direction_column(api, world, memory_storage) -> None:
    import openpyxl

    client = api.as_user(world["head"])
    created = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:interactions",
            "format": "xlsx",
            "filters": {
                "from": "2026-09-21",
                "to": "2026-10-20",
                "direction_id": [str(world["qa"].id)],
            },
            "columns": ["university", "it_direction", "product"],
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["rows"] == 2
    content = (await client.get(f"{API}/export/jobs/{created.json()['id']}/download")).content
    book = openpyxl.load_workbook(io.BytesIO(content))
    sheet = book["Взаимодействия"]
    assert [cell.value for cell in sheet[1]] == [
        "Наименование вуза",
        "ИТ-направление",
        "ИТ-продукт",
    ]
    directions = {row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert directions == {"DevOps, Тестирование (QA)"}
    about = "\n".join(
        str(cell) for row in book["О выгрузке"].iter_rows(values_only=True) for cell in row if cell
    )
    assert "ИТ-направления: Тестирование (QA)" in about

    kam_job = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "report:kam-activity",
            "format": "xlsx",
            "filters": {"direction_id": [str(world["data"].id)]},
        },
    )
    assert kam_job.status_code == 201, kam_job.text


# --- сид демо-стенда -----------------------------------------------------------------------------


async def test_seed_directions_idempotent(db_session) -> None:
    from app.tools.seed import (
        IT_DIRECTIONS,
        PRODUCT_DIRECTIONS,
        PROGRAM_DIRECTIONS,
        seed_it_directions,
    )

    assert 8 <= len(IT_DIRECTIONS) <= 12
    products = {code: Product(code=code, name=f"Продукт {code}") for code in PRODUCT_DIRECTIONS}
    db_session.add_all(products.values())
    await db_session.flush()
    programs = {name: Program(name=name) for name in PROGRAM_DIRECTIONS}
    db_session.add_all(programs.values())
    await db_session.flush()

    directions = await seed_it_directions(db_session, products, programs)
    assert set(directions) == {code for code, _name, _description in IT_DIRECTIONS}
    assert programs["MLOps для инженеров"].direction_id == directions["devops"].id
    manual = programs["Корпоративная школа Python"]
    manual.direction_id = directions["qa"].id  # ручная правка не перетирается повтором
    await db_session.flush()

    again = await seed_it_directions(db_session, products, programs)
    assert {code: row.id for code, row in again.items()} == {
        code: row.id for code, row in directions.items()
    }
    assert manual.direction_id == directions["qa"].id
    total = (
        await db_session.execute(sa.select(sa.func.count()).select_from(ItDirection))
    ).scalar_one()
    assert total == len(IT_DIRECTIONS)
    links = (
        await db_session.execute(sa.select(sa.func.count()).select_from(ProductDirection))
    ).scalar_one()
    assert links == sum(len(codes) for codes in PRODUCT_DIRECTIONS.values())


# --- после diff интегратора: query-параметр отчётов, выгрузки реестров, импорт заявок ----------
# Файлы вне зоны потока (reporting/router.py, import_export/export.py и importers.py) правит
# интегратор; тесты включаются сами, когда правка внесена.


def _router_has_direction() -> bool:
    from app.modules.reporting.router import kam_activity_params, report_params

    return all(
        "direction_id" in inspect.signature(fn).parameters
        for fn in (report_params, kam_activity_params)
    )


def _export_has_direction() -> bool:
    from app.modules.import_export import export

    return "direction_id" in inspect.getsource(export._dataset_requests)


def _importer_has_direction() -> bool:
    from app.modules.import_export import importers

    return hasattr(importers, "_lookup_direction")


@pytest.mark.skipif(not _router_has_direction(), reason="ждёт diff reporting/router.py")
async def test_reports_query_param_direction(api, world) -> None:
    client = api.as_user(world["head"])
    response = await client.get(
        f"{API}/reports/funnel", params={"direction_id": str(world["qa"].id)}
    )
    assert response.status_code == 200, response.text
    assert sum(item["value"] for item in response.json()["series"]) == 2
    assert response.json()["params"]["direction_id"] == [str(world["qa"].id)]
    kam = await client.get(
        f"{API}/reports/kam-activity", params={"direction_id": str(world["data"].id)}
    )
    assert kam.status_code == 200, kam.text
    KamActivityReportOut.model_validate(kam.json())
    assert _titles(world, _kam_request_ids(kam.json())) == {"C"}
    summary = await client.get(
        f"{API}/reports/kam-activity/summary", params={"direction_id": str(world["qa"].id)}
    )
    assert summary.status_code == 200, summary.text
    many = await client.get(
        f"{API}/reports/dashboard",
        params=[("direction_id", str(uuid.uuid4())) for _ in range(51)],
    )
    assert many.status_code == 400


@pytest.mark.skipif(not _export_has_direction(), reason="ждёт diff import_export/export.py")
async def test_registry_exports_direction(api, world, memory_storage) -> None:
    import openpyxl

    client = api.as_user(world["head"])
    created = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "requests",
            "format": "xlsx",
            "filters": {"direction_id": [str(world["qa"].id)]},
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["rows"] == 2
    content = (await client.get(f"{API}/export/jobs/{created.json()['id']}/download")).content
    sheet = openpyxl.load_workbook(io.BytesIO(content)).active
    header = [cell.value for cell in sheet[1]]
    column = header.index("ИТ-направление")
    values = {row[column] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert values == {"DevOps, Тестирование (QA)"}
    programs = await client.post(
        f"{API}/export/jobs",
        json={
            "entity_type": "programs",
            "format": "csv",
            "filters": {"direction_id": [str(world["qa"].id)]},
        },
    )
    assert programs.status_code == 201, programs.text
    assert programs.json()["rows"] == 1


@pytest.mark.skipif(not _importer_has_direction(), reason="ждёт diff import_export/importers.py")
async def test_import_requests_with_direction(db_session, world) -> None:
    from app.core.pii import get_pii_crypto
    from app.models import Deal
    from app.modules.import_export.importers import PreparedRow, get_importer, prepare_context

    ctx = await prepare_context(
        db_session,
        entity_type="requests_b2b",
        entries=[],
        raw_options={},
        actor=world["head"],
        crypto=get_pii_crypto(),
        import_session_id=uuid.uuid4(),
    )
    rows = [
        PreparedRow(
            row_number=1,
            values={
                "university": world["mipt"].name,
                "product": "plain",
                "it_direction": "QA",  # по коду, без регистра
                "title": "Импорт QA",
            },
        ),
        PreparedRow(
            row_number=2,
            values={"university": world["mipt"].name, "it_direction": "Геймдев", "title": "X"},
        ),
    ]
    importer = get_importer("requests_b2b")
    await importer.validate(ctx, rows)
    assert rows[0].valid and not rows[1].valid
    assert rows[1].issues[0].column == "it_direction"
    await importer.apply(ctx, [rows[0]])
    links = (
        (
            await db_session.execute(
                sa.select(ProductDirection.direction_id).where(
                    ProductDirection.product_id == world["plain"].id
                )
            )
        )
        .scalars()
        .all()
    )
    assert links == [world["qa"].id]  # продукт строки теперь учит и по QA
    deal = (await db_session.execute(sa.select(Deal).where(Deal.title == "Импорт QA"))).scalar_one()
    assert deal.product_id == world["plain"].id
