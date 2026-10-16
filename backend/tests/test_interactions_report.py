"""Отчёт «Взаимодействия за период» (P0-6), вход в этап без строк перевода (P0-1), вид Б
сводки для заказчика (P1-3) и фильтры выгрузок (P0-2): чистые функции, SQL, рендер файлов
без БД."""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

import pytest
from sqlalchemy.dialects import postgresql

from app.core.access import DealScope
from app.core.errors import ApiError
from app.modules.import_export import export as export_module
from app.modules.import_export.export import (
    filter_date,
    filter_uuids,
    report_params_from_filters,
)
from app.modules.reporting import interactions as it
from app.modules.reporting import kam_activity as ka
from app.modules.reporting.interactions_export import (
    SHEET_NAMES,
    cell_value,
    describe_filters,
    flat_rows,
    render_interactions_pdf,
    render_interactions_xlsx,
)
from app.modules.reporting.kam_activity import DealRow, KamActivityParams, PersonRef, StageRef
from app.modules.reporting.schemas import InteractionsReportOut

MSK = ka.REPORT_TZ
NOW = datetime(2026, 10, 20, 12, 0, tzinfo=MSK)
TODAY = NOW.date()


def _sql(stmt: Any) -> str:
    return str(stmt.compile(dialect=postgresql.dialect()))


def _bounds(date_from: date, date_to: date) -> ka.PeriodBounds:
    params = ka.resolve_params(KamActivityParams(date_from=date_from, date_to=date_to), today=TODAY)
    return ka.resolve_bounds(params, now=NOW)


CURRENT = _bounds(date(2026, 9, 21), TODAY)
PAST = _bounds(date(2026, 7, 1), date(2026, 9, 30))


# --- P0-1: момент входа в этап -------------------------------------------------------------


class TestStageEntry:
    def test_current_and_past_bounds(self) -> None:
        assert CURRENT.is_current is True and CURRENT.as_of == NOW
        assert PAST.is_current is False
        assert ka.previous_bounds(CURRENT).is_current is False

    def test_current_period_uses_stage_entered_at(self) -> None:
        sql = _sql(ka.deal_rows_stmt(KamActivityParams(), CURRENT, 14))
        assert sql.count("LATERAL") == 5 and "dsh_enter" not in sql
        assert "deal.stage_entered_at AS stage_since" in sql

    def test_past_period_entry_lateral_skips_migration_rows(self) -> None:
        stmt = ka.deal_rows_stmt(KamActivityParams(), PAST, 14)
        sql = _sql(stmt)
        assert sql.count("LATERAL") == 6
        enter = sql.split("AS dsh_enter", 1)[1].split(") AS h_enter", 1)[0]
        assert "dsh_enter.kind != " in enter
        assert "dsh_enter.from_stage_id IS DISTINCT FROM dsh_enter.to_stage_id" in enter
        assert "coalesce(h_enter.since, deal.stage_entered_at) AS stage_since" in sql
        # этап на конец периода — по-прежнему последняя строка истории любого вида
        end = sql.split("AS dsh_end", 1)[1].split(") AS h_end", 1)[0]
        assert "kind" not in end
        assert "system_migration" in stmt.compile(dialect=postgresql.dialect()).params.values()

    def test_reached_counts_only_entries(self) -> None:
        sql = _sql(ka.stage_reached_stmt([uuid.uuid4()], PAST))
        assert "deal_stage_history.kind != " in sql
        assert "deal_stage_history.from_stage_id IS DISTINCT FROM" in sql

    def test_dwell_segments_by_entries(self) -> None:
        sql = _sql(ka.stage_dwell_stmt([uuid.uuid4()], PAST))
        # сегмент — от строки-входа до следующей строки-входа; этап — на конце сегмента
        assert "sum(CASE WHEN (deal_stage_history.kind != " in sql
        assert "ORDER BY marked.created_at DESC, marked.row_id DESC))[1]" in sql
        assert "lead(segments.entered) OVER (PARTITION BY segments.deal_id" in sql

    def test_extra_filters_of_deal_rows(self) -> None:
        stage = uuid.uuid4()
        sql = _sql(
            ka.deal_rows_stmt(
                KamActivityParams(),
                CURRENT,
                14,
                stage_ids=(stage,),
                where=[ka.Deal.product_id.in_([uuid.uuid4()])],
            )
        )
        assert "end_stage.id IN" in sql and "deal.product_id IN" in sql


# --- P0-6: параметры, колонки, запрос -------------------------------------------------------


class TestInteractionsParams:
    def test_columns_default_and_unknown(self) -> None:
        assert it.resolve_columns(None) == it.DEFAULT_COLUMNS
        assert it.DEFAULT_COLUMNS[:5] == (
            "university",
            "it_direction",
            "product",
            "status",
            "responsible",
        )
        assert it.resolve_columns(["status", "university", "status"]) == ("status", "university")
        with pytest.raises(ApiError) as excinfo:
            it.resolve_columns(["salary"])
        assert excinfo.value.status_code == 400
        assert excinfo.value.details[0]["code"] == "unknown_column"

    def test_from_filters_lists_and_url_keys(self) -> None:
        a, b, c = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
        params = it.InteractionsParams.from_filters(
            {
                "from": "2026-09-01",
                "to": "2026-09-30",
                "wf": "b2b",
                "kam": [str(a)],
                "product_id": [str(b), str(c)],
                "program": str(a),
                "status": [str(b)],
            }
        )
        assert params.base.workflow_type == "b2b" and params.base.kam_ids == (a,)
        assert params.product_ids == (b, c) and params.program_ids == (a,)
        assert params.stage_ids == (b,)
        assert params.as_json()["product_id"] == sorted([str(b), str(c)])
        with pytest.raises(ApiError) as excinfo:
            it.InteractionsParams.from_filters({"product_id": ["не-uuid"]})
        assert excinfo.value.details[0]["code"] == "invalid_uuid"

    def test_resolve_params_limits(self) -> None:
        many = tuple(uuid.uuid4() for _ in range(ka.MAX_FILTER_IDS + 1))
        with pytest.raises(ApiError) as excinfo:
            it.resolve_params(it.InteractionsParams(stage_ids=many), today=TODAY)
        assert excinfo.value.details[0] == {
            "field": "stage_id",
            "code": "too_many_values",
            "message": f"Не больше {ka.MAX_FILTER_IDS} значений фильтра",
        }
        resolved = it.resolve_params(it.InteractionsParams(), today=TODAY)
        assert resolved.base.date_from == date(2026, 9, 21)  # как у отчёта КАМов

    def test_statement_reuses_kam_activity_query(self) -> None:
        params = it.InteractionsParams(
            product_ids=(uuid.uuid4(),), program_ids=(uuid.uuid4(),), stage_ids=(uuid.uuid4(),)
        )
        sql = _sql(it.interactions_stmt(params, CURRENT, 14, DealScope.unrestricted()))
        assert "dsh_own.actor_user_id = deal.responsible_user_id" in sql  # D4 — тот же запрос
        assert "deal.product_id IN" in sql and "deal.program_id IN" in sql
        assert "end_stage.id IN" in sql
        assert "LEFT OUTER JOIN product ON product.id = deal.product_id" in sql
        assert "LEFT OUTER JOIN app_user AS responsible" in sql

    def test_cache_key_depends_on_filters_and_scope(self) -> None:
        own = DealScope(rules=(("b2b", uuid.uuid4()),))
        first = it._cache_key(it.InteractionsParams(), DealScope.unrestricted())
        assert first.startswith("dash:interactions:")
        assert first != it._cache_key(it.InteractionsParams(), own)
        assert first != it._cache_key(
            it.InteractionsParams(product_ids=(uuid.uuid4(),)), DealScope.unrestricted()
        )


# --- P0-6: строки и файлы ---------------------------------------------------------------------

KAM = uuid.uuid4()
MIPT = uuid.uuid4()
STAGE = uuid.uuid4()


def _mapping(**overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "deal_id": uuid.uuid4(),
        "title": "МФТИ — Python",
        "pipeline": "b2b",
        "university_id": MIPT,
        "university_name": "Московский физико-технический институт",
        "university_short": "МФТИ",
        "university_region": None,
        "counterparty_id": None,
        "client": None,
        "responsible_id": KAM,
        "amount": Decimal("450000.00"),
        "status": "open",
        "created_at": NOW - timedelta(days=100),
        "stage_id": STAGE,
        "stage_code": "negotiation",
        "stage_current_name": "Переговоры",
        "stage_history_name": "Переговоры",
        "stage_color": "#1677ff",
        "stage_is_terminal": False,
        "stage_outcome": None,
        "stage_sort": 2,
        "stage_deleted_at": None,
        "stage_since": NOW - timedelta(days=20),
        "threshold": 14,
        "transitions": 1,
        "returns": 0,
        "comments": 1,
        "files": 0,
        "last_contact_at": NOW - timedelta(days=10),
        "product_id": uuid.uuid4(),
        "product_name": "Python-разработчик",
        "program_id": None,
        "program_name": None,
        "responsible_full_name": "Петров Константин",
        "responsible_username": "kam1",
    }
    base.update(overrides)
    return base


def sample_report() -> dict[str, Any]:
    rows = [
        it.interaction_row(_mapping(), CURRENT),
        it.interaction_row(
            _mapping(
                title="Курс — физлицо",
                pipeline="b2c",
                university_id=None,
                university_name=None,
                university_short=None,
                client="Смирнова О. В.",
                amount=None,
                product_id=None,
                product_name=None,
                stage_is_terminal=True,
                stage_outcome="won",
                stage_code="completed",
                stage_current_name="Завершена",
                stage_since=NOW - timedelta(days=3),
            ),
            CURRENT,
        ),
    ]
    stage_options = [
        {
            "id": str(STAGE),
            "code": "negotiation",
            "name": "Переговоры",
            "color": None,
            "workflow_type": "b2b",
            "is_terminal": False,
        }
    ]
    return {
        "report": "interactions",
        "generated_at": NOW.isoformat(),
        "params": {
            "from": "2026-09-21",
            "to": "2026-10-20",
            "workflow_type": None,
            "kam_id": [str(KAM)],
            "university_id": [str(MIPT)],
            "product_id": [],
            "program_id": [],
            "stage_id": [str(STAGE)],
        },
        "period": CURRENT.as_json(),
        "scope": {"own_only": False, "own_user_id": None, "workflow_types": ["b2b", "b2c"]},
        "columns": it.columns_payload(),
        "default_columns": list(it.DEFAULT_COLUMNS),
        "stages": stage_options,
        "totals": it.totals_of(rows),
        "truncated": False,
        "rows": rows,
    }


class TestInteractionRows:
    def test_row_shape_matches_schema(self) -> None:
        report = sample_report()
        InteractionsReportOut.model_validate(report)
        row = report["rows"][0]
        assert row["it_direction"] is None and row["product"] == "Python-разработчик"
        assert row["days_on_stage"] == 20 and row["is_stuck"] is True
        assert row["interactions"]["total"] == 2  # П + К + Ф (мероприятия — на уровне вуза)
        assert report["totals"]["requests"] == 2 and report["totals"]["universities"] == 1
        assert report["totals"]["won"] == 1 and report["totals"]["amount"] == "450000.00"

    def test_cell_values_for_files(self) -> None:
        b2b, b2c = sample_report()["rows"]
        assert cell_value(b2b, "university") == "Московский физико-технический институт"
        assert cell_value(b2c, "university") == "Смирнова О. В. (B2C)"
        assert cell_value(b2b, "it_direction") == "—"
        assert cell_value(b2b, "status") == "Переговоры"
        assert cell_value(b2b, "created_at") == (NOW - timedelta(days=100)).date()
        assert cell_value(b2c, "days_on_stage") is None  # финальный этап — без «дней»
        assert cell_value(b2c, "closed_in_period") == "Выиграна"
        assert cell_value(b2b, "amount") == 450000.0 and cell_value(b2b, "workflow_type") == "B2B"

    def test_describe_filters_uses_names_from_report(self) -> None:
        text = describe_filters(sample_report())
        assert "Ответственные: Петров Константин" in text
        assert "Вузы: Московский физико-технический институт" in text
        assert "Статусы: Переговоры" in text

    def test_xlsx_selected_columns(self) -> None:
        import openpyxl

        columns = ("university", "it_direction", "status", "stage_since", "amount")
        data = render_interactions_xlsx(
            sample_report(), columns, generated_by="Морозов Д.", generated_at=NOW
        )
        book = openpyxl.load_workbook(io.BytesIO(data))
        assert tuple(book.sheetnames) == SHEET_NAMES
        sheet = book[SHEET_NAMES[0]]
        assert [cell.value for cell in sheet[1]] == [
            "Наименование вуза",
            "ИТ-направление",
            "Статус работы с вузом",
            "В статусе с",
            "Сумма, ₽",
        ]
        assert sheet["B2"].value == "—" and isinstance(sheet["D2"].value, datetime)
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:E3"
        about = [row[0] for row in book[SHEET_NAMES[1]].iter_rows(values_only=True) if row[0]]
        assert any(str(line).startswith("Период: 21.09.2026") for line in about)

    def test_xlsx_formula_injection_is_text(self) -> None:
        import openpyxl

        report = sample_report()
        report["rows"][0]["title"] = '=HYPERLINK("http://evil")'
        data = render_interactions_xlsx(report, ("title",), generated_by="Тест", generated_at=NOW)
        cell = openpyxl.load_workbook(io.BytesIO(data))[SHEET_NAMES[0]]["A2"]
        assert cell.data_type == "s"

    def test_pdf_renders_cyrillic(self) -> None:
        data = render_interactions_pdf(
            sample_report(), it.DEFAULT_COLUMNS, generated_by="Морозов Д.", generated_at=NOW
        )
        assert data.startswith(b"%PDF")
        empty = sample_report() | {"rows": []}
        assert render_interactions_pdf(
            empty, ("university",), generated_by="Тест", generated_at=NOW
        ).startswith(b"%PDF")

    def test_flat_rows_csv_ready(self) -> None:
        labels, rows = flat_rows(sample_report(), ("university", "product"))
        assert labels == [("university", "Наименование вуза"), ("product", "ИТ-продукт")]
        assert rows[1] == {"university": "Смирнова О. В. (B2C)", "product": None}


# --- P1-3: вид Б — вузы по этапам, риски ------------------------------------------------------

KAM1, KAM2 = uuid.uuid4(), uuid.uuid4()
U1, U2, U3, U4 = (uuid.uuid4() for _ in range(4))
NEG, SIGNED, WON, LOST, NEW = (uuid.uuid4() for _ in range(5))
STAGES = [
    StageRef(NEW, "b2b", "new", "Новая", None, 0, False, None, False, True, 7),
    StageRef(NEG, "b2b", "negotiation", "Переговоры", None, 2, False, None, False),
    StageRef(SIGNED, "b2b", "contract_signed", "Договор подписан", None, 4, False, None, False),
    StageRef(WON, "b2b", "completed", "Завершена", None, 6, True, "won", False),
    StageRef(LOST, "b2b", "rejected", "Отказ", None, 7, True, "lost", False),
    StageRef(uuid.uuid4(), "b2b", "old", "Старый", None, 3, False, None, True),
]
_BY_ID = {stage.stage_id: stage for stage in STAGES}


def _row(university: uuid.UUID | None, stage: uuid.UUID, **overrides: Any) -> DealRow:
    ref = _BY_ID[stage]
    base: dict[str, Any] = {
        "deal_id": uuid.uuid4(),
        "title": "Заявка",
        "pipeline": "b2b",
        "university_id": university,
        "university_name": f"Вуз {str(university)[:4]}",
        "university_short": None,
        "university_region": None,
        "counterparty_id": None,
        "client": None,
        "responsible_id": KAM1,
        "amount": None,
        "status": "open",
        "created_at": NOW - timedelta(days=60),
        "stage_id": stage,
        "stage_code": ref.code,
        "stage_name": ref.name,
        "stage_color": None,
        "stage_is_terminal": ref.is_terminal,
        "stage_outcome": ref.outcome,
        "stage_sort": ref.sort_order,
        "stage_deleted": False,
        "stage_since": NOW - timedelta(days=5),
        "threshold": 14,
    }
    base.update(overrides)
    return DealRow(**base)


PEOPLE = {
    KAM1: PersonRef(KAM1, "Петров К.", ("kam",)),
    KAM2: PersonRef(KAM2, "Иванова М.", ("kam",)),
}


def _tree(rows: list[DealRow], assigned: list[ka.UniversityRef] | None = None) -> dict[str, Any]:
    return ka.assemble_tree(rows, [], assigned or [], PEOPLE, CURRENT)


class TestUniversitiesByStage:
    def test_most_advanced_stage_rule(self) -> None:
        tree = _tree(
            [
                # U1: отказ и переговоры → «Переговоры» (отказ — только если других нет)
                _row(U1, LOST),
                _row(U1, NEG),
                # U2: договор подписан у одного КАМа и выигрыш у другого → «Завершена»
                _row(U2, SIGNED),
                _row(U2, WON, responsible_id=KAM2),
                # U3: только отказ
                _row(U3, LOST),
                # B2C-заявка вуза не имеет — в диаграмму вузов не попадает
                _row(None, NEW, pipeline="b2c", client="Клиент К."),
            ],
            assigned=[ka.UniversityRef(U4, "Вуз без заявок", None, None, KAM1)],
        )
        prev = _tree([_row(U1, NEW)])
        block = ka.universities_by_stage(tree, prev, STAGES)
        values = {row["code"]: row for row in block["rows"]}
        assert [row["code"] for row in block["rows"]] == [
            "new",
            "negotiation",
            "contract_signed",
            "completed",
            "rejected",
        ]  # удалённый этап без вузов скрыт, финальные — в конце
        assert values["negotiation"]["value"] == 1
        assert values["completed"]["value"] == 1 and values["contract_signed"]["value"] == 0
        assert values["rejected"]["value"] == 1
        assert values["new"] == values["new"] | {"value": 0, "prev": 1}
        assert block["total"] == 3 and block["total_prev"] == 1
        assert block["without_requests"]["value"] == 1
        assert block["without_requests"]["universities"] == ["Вуз без заявок"]
        assert block["workflow_type"] == "b2b" and "один раз" in block["rule"]

    def test_risks_only_non_zero(self) -> None:
        kpi = {
            "stuck": {"value": 2, "prev": 1, "delta": 1},
            "lost": {"value": 0, "prev": 1, "delta": -1},
        }
        summary = {
            "kpi": kpi,
            "top_stuck": [
                {
                    "title": "МФТИ — Python",
                    "university": "МФТИ",
                    "stage": {"name": "Переговоры"},
                    "days_on_stage": 40,
                }
            ],
            "untouched_assigned": ["ИТМО"],
        }
        risks = ka.build_risks(summary)
        assert [risk["key"] for risk in risks] == ["stuck", "untouched"]
        assert risks[0]["items"] == ["МФТИ — Python: «Переговоры», 40 дн."]
        kpi["lost"] = {"value": 3, "prev": 1, "delta": 2}
        assert ka.build_risks(summary)[-1] | {} == {
            "key": "lost",
            "label": "Отказы за период",
            "value": 3,
            "note": "К прошлому периоду +2",
            "items": [],
        }

    def test_contract_active_on_rule(self) -> None:
        sql = _sql(ka.contract_active_on(date(2026, 9, 30)))
        assert "coalesce(contract.valid_from, contract.signed_at)" in sql
        assert "contract.status = " in sql and "contract.valid_to >= " in sql


# --- P0-2: фильтры тела выгрузки -------------------------------------------------------------


class TestExportFilters:
    def test_uuid_lists_and_scalars(self) -> None:
        a, b = uuid.uuid4(), uuid.uuid4()
        assert filter_uuids({"kam_id": str(a)}, "kam_id") == (a,)
        assert filter_uuids({"kam_id": [str(a), str(b), str(a)]}, "kam_id") == (a, b)
        assert filter_uuids({"federal_project": str(a)}, "x", "federal_project") == (a,)
        assert filter_uuids({"kam_id": []}, "kam_id") == ()

    @pytest.mark.parametrize(
        "filters",
        [{"kam_id": "мусор"}, {"kam_id": [str(uuid.uuid4()), "мусор"]}, {"kam_id": [123]}],
    )
    def test_bad_uuid_is_422(self, filters: dict[str, Any]) -> None:
        with pytest.raises(ApiError) as excinfo:
            filter_uuids(filters, "kam_id")
        error = excinfo.value
        assert (error.status_code, error.code) == (422, "unprocessable")
        assert error.details[0]["field"] == "filters.kam_id"

    def test_bad_date_is_422(self) -> None:
        assert filter_date({"from": "2026-10-01"}, "from") == date(2026, 10, 1)
        with pytest.raises(ApiError) as excinfo:
            filter_date({"from": "01.10.2026"}, "from")
        assert excinfo.value.status_code == 422

    def test_report_params_keep_whole_list(self) -> None:
        a, b = uuid.uuid4(), uuid.uuid4()
        params = report_params_from_filters({"kam_id": [str(a), str(b)], "workflow_type": "b2b"})
        kam = getattr(params, "kam_ids", params.kam_id)
        # не «первый элемент списка»: либо оба значения, либо явная 422 (одиночное поле)
        assert set(kam if isinstance(kam, tuple) else (kam,)) == {a, b}

    def test_report_params_bad_value_is_422(self) -> None:
        for filters in ({"kam_id": ["не-uuid"]}, {"from": "вчера"}, {"workflow_type": "b2x"}):
            with pytest.raises(ApiError) as excinfo:
                report_params_from_filters(filters)
            assert excinfo.value.status_code == 422, filters

    def test_legacy_single_value_fields_reject_lists(self, monkeypatch) -> None:  # noqa: ANN001
        """Без единого слоя фильтров (скалярный ReportParams) — 422 вместо сужения."""
        from dataclasses import dataclass

        @dataclass(frozen=True)
        class ScalarParams:
            workflow_type: str | None = None
            date_from: date | None = None
            date_to: date | None = None
            kam_id: uuid.UUID | None = None
            product_id: uuid.UUID | None = None
            university_id: uuid.UUID | None = None
            program_id: uuid.UUID | None = None
            region: str | None = None
            federal_project_id: uuid.UUID | None = None

        monkeypatch.setattr(export_module, "ReportParams", ScalarParams)
        a = uuid.uuid4()
        assert report_params_from_filters({"kam_id": [str(a)]}).kam_id == a
        with pytest.raises(ApiError) as excinfo:
            report_params_from_filters({"kam_id": [str(a), str(uuid.uuid4())]})
        assert excinfo.value.status_code == 422
        assert excinfo.value.details[0]["code"] == "too_many_values"
