"""Выгрузка «Просрочек» без БД: XLSX (листы, A1, шапки, форматы, CWE-1236), модель PDF
(«Сводка без имён» не содержит имён КАМов), CSV-колонки и имя файла `overdue_*`.

Отчёт собирается настоящими функциями `sla.team_body` из синтетического набора данных —
та же сборка, что у API, но без PostgreSQL.
"""

from __future__ import annotations

import importlib.util
import io
import uuid
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from typing import Any

import pytest

from app.modules.reporting import sla
from app.modules.reporting.kam_activity import REPORT_TZ, PersonRef, _bounds, previous_bounds
from app.modules.reporting.kam_activity_export import ExportMeta
from app.modules.reporting.sla import (
    DealInfo,
    HistoryRow,
    OwnerTimeline,
    SlaDataset,
    SnapshotGrid,
    StageInfo,
)
from app.modules.reporting.sla_export import (
    FLAT_COLUMNS,
    SHEET_NAMES,
    file_name,
    pdf_kam_sections,
    pdf_sections,
    render_sla_pdf,
    render_sla_xlsx,
    section_texts,
    sheet_names,
    sla_flat_rows,
)

NOW = datetime(2026, 10, 20, 12, 0, tzinfo=REPORT_TZ)
WF = "b2b"
NEW = uuid.UUID(int=1)
NEG = uuid.UUID(int=2)
STAGES = {
    NEW: StageInfo(NEW, WF, "new", "Новая", "#4C9AFF", 0, False, None, False, 7, "stage"),
    NEG: StageInfo(
        NEG, WF, "negotiation", "Переговоры", "#F59E0B", 2, False, None, False, 14, "workflow"
    ),
}
KAMS = {
    uuid.UUID(int=101): "Павлов Роман Евгеньевич",
    uuid.UUID(int=102): "Иванова Мария Сергеевна",
    uuid.UUID(int=103): "Фролов Кирилл Олегович",
}


def ago(days: float) -> datetime:
    return NOW - timedelta(days=days)


def make_dataset(*, titles: dict[int, str] | None = None) -> SlaDataset:
    """Набор данных отчёта из чистых функций `sla` (как `load_dataset`, но без запросов)."""
    bounds = _bounds(date(2026, 9, 21), date(2026, 10, 20), NOW)
    prev = previous_bounds(bounds)
    granularity, window_start, points = sla.trend_grid(bounds)
    owners = list(KAMS)
    deals: dict[uuid.UUID, DealInfo] = {}
    visits: list[sla.Visit] = []
    plan = [(0, 40, NEG), (0, 20, NEG), (1, 10, NEW), (1, 3, NEW), (2, 60, NEG), (2, 16, NEG)]
    for number, (owner, days, stage) in enumerate(plan):
        deal_id = uuid.UUID(int=1000 + number)
        info = DealInfo(
            deal_id=deal_id,
            title=(titles or {}).get(number, f"Заявка №{number}"),
            pipeline=WF,
            university_id=uuid.UUID(int=900),
            university_name="Национальный исследовательский университет «МЭИ»",
            university_short="МЭИ",
            client=None,
            responsible_id=owners[owner],
            stage_id=stage,
            stage_entered_at=ago(days),
            created_at=ago(days + 5),
        )
        deals[deal_id] = info
        rows = [HistoryRow(deal_id, ago(days + 5), None, NEW, "manual", True)]
        if stage != NEW:
            rows.append(HistoryRow(deal_id, ago(days), NEW, stage, "manual", True))
        visits += sla.build_visits(rows, info, STAGES, is_current=True)
    timelines = {deal_id: OwnerTimeline((), info.responsible_id) for deal_id, info in deals.items()}
    snap = sla.snapshot(visits, bounds.as_of, timelines)
    snap_prev = sla.snapshot(visits, prev.as_of, timelines)
    grid = SnapshotGrid([point.snapshot_at for point in points])
    for visit in visits:
        grid.add_visit(visit, timelines[visit.deal_id])
    since = min(prev.start, window_start)
    return SlaDataset(
        bounds=bounds,
        prev=prev,
        granularity=granularity,
        window_start=window_start,
        points=points,
        workflows=(WF,),
        stages=STAGES,
        deals=deals,
        visits=visits,
        timelines=timelines,
        snap=snap,
        snap_prev=snap_prev,
        last_actions={},
        no_action=sla.no_action_ids(snap, {}),
        no_action_prev=sla.no_action_ids(snap_prev, {}),
        decisions=sla.decided_in(visits, timelines, since, bounds.as_of),
        closures=sla.closed_late_in(visits, timelines, prev.start, bounds.as_of),
        inflow_cur=Counter(),
        inflow_prev=Counter(),
        grid=grid,
        visit_reminders={},
        episodes=[],
        journal_from=None,
        check_enabled=True,
        disabled_since=None,
        rules_active=True,
        notes=["Журнал напоминаний ведётся с 01.10.2026."],
        people={
            user_id: PersonRef(user_id=user_id, full_name=name, roles=("kam",))
            for user_id, name in KAMS.items()
        },
    )


def make_report(**kwargs: Any) -> dict[str, Any]:
    body = sla.team_body(make_dataset(**kwargs), include_late=True, limit_items=False)
    return {
        "report": "sla",
        "generated_at": NOW.isoformat(),
        "params": {
            "from": "2026-09-21",
            "to": "2026-10-20",
            "workflow_type": None,
            "kam_id": [],
            "university_id": [],
        },
        "scope": {"own_only": False, "own_user_id": None, "workflow_types": ["b2b", "b2c"]},
        **body,
    }


META = ExportMeta(
    generated_by="Морозов Дмитрий Александрович",
    generated_at=datetime(2026, 10, 20, 9, 0, tzinfo=timezone.utc),
    filters_text="Процесс: B2B и B2C · КАМ: все",
)


def test_report_shape() -> None:
    report = make_report()
    assert report["kpi"]["overdue"]["value"] == 5  # «Переговоры» −40, −20, −60, −16; «Новая» −10
    assert report["overdue"]["total"] == 5
    names = {row["kam"] for row in report["kams"]}
    assert names == set(KAMS.values())


def test_xlsx_sheets_headline_formats_and_formula_as_text() -> None:
    import openpyxl

    report = make_report(titles={0: '=HYPERLINK("http://evil.example","x")'})
    book = openpyxl.load_workbook(io.BytesIO(render_sla_xlsx(report, META)))
    # журнал напоминаний пуст — листа «Напоминания» нет
    assert tuple(book.sheetnames) == tuple(n for n in SHEET_NAMES if n != "Напоминания")
    assert sheet_names(report) == tuple(book.sheetnames)
    assert book["Сводка"]["A1"].value == report["headline"]
    assert book["Сводка"]["A2"].value == sla.DEFINITION
    sheet = book["Просрочено сейчас"]
    assert [cell.value for cell in sheet[1]] == [label for _key, label in FLAT_COLUMNS]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref.startswith("A1:")
    delay_col = [label for _key, label in FLAT_COLUMNS].index("Сверх нормы, дн.") + 1
    assert sheet.cell(row=2, column=delay_col).number_format == "0.0"
    since_col = [label for _key, label in FLAT_COLUMNS].index("На этапе с") + 1
    assert sheet.cell(row=2, column=since_col).number_format == "DD.MM.YYYY"
    titles = [
        cell
        for row in sheet.iter_rows(min_row=2)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK")
    ]
    assert titles and all(cell.data_type == "s" for cell in titles)  # CWE-1236: текстом
    kams = book["КАМы"]
    assert kams["A1"].value == "КАМ" and kams["C1"].value == "В работе"
    assert [cell.value for cell in book["Тренд"][1]][:3] == ["Неделя с", "по", "Дата среза"]
    definitions = [row[0].value for row in book["Как считается"].iter_rows(min_row=4)]
    assert "Просрочено" in definitions


def test_xlsx_with_reminders_sheet() -> None:
    import openpyxl

    report = make_report()
    report["reminders"]["info"]["coverage"] = "full"
    report["reminders"]["team"] = dict.fromkeys(("sent", "delivered"), 0)
    book = openpyxl.load_workbook(io.BytesIO(render_sla_xlsx(report, META)))
    assert tuple(book.sheetnames) == SHEET_NAMES


def test_pdf_sections_without_names_and_with_names() -> None:
    report = make_report()
    unnamed = " ".join(section_texts(pdf_sections(report, False, META)))
    for name in KAMS.values():
        assert name not in unnamed
    assert "Просрочек больше, чем по команде:" in unnamed
    named = " ".join(section_texts(pdf_sections(report, True, META)))
    assert any(name in named for name in KAMS.values())
    texts = section_texts(pdf_sections(report, True, META))
    assert not any(sla.NO_ADVICE_RE.search(text) for text in texts)
    assert not any("SLA" in text for text in texts)
    assert report["headline"] in texts


@pytest.mark.skipif(importlib.util.find_spec("fpdf") is None, reason="fpdf2 не установлен")
def test_pdf_renders() -> None:
    report = make_report()
    assert render_sla_pdf(report, META, by_kam=False).startswith(b"%PDF")
    assert render_sla_pdf(report, META, by_kam=True).startswith(b"%PDF")


def test_kam_pdf_sections() -> None:
    dataset = make_dataset()
    person = dataset.people[uuid.UUID(int=101)]
    body = sla.kam_body(dataset, person, with_team=False)
    report = {"params": {}, "scope": {"own_only": True}, **body}
    texts = section_texts(pdf_kam_sections(report, META))
    assert any(text.startswith("У Павлова Р. Е.:") for text in texts)
    assert "Журнал напоминаний пуст." in texts


def test_csv_columns_and_file_name() -> None:
    from app.modules.import_export.export import render_export

    report = make_report()
    columns, rows = sla_flat_rows(report)
    assert [key for key, _label in columns] == [key for key, _label in FLAT_COLUMNS]
    assert [label for _key, label in columns][:3] == ["КАМ", "Заявка", "Процесс"]
    assert len(rows) == report["overdue"]["total"]
    assert rows[0]["delay_days"] >= rows[-1]["delay_days"]  # по убыванию задержки
    assert rows[0]["link"].endswith(f"request={report['overdue']['items'][0]['request_id']}")
    data = render_export(columns, rows, export_format="csv", encoding="utf-8", delimiter=";")
    header = data.decode("utf-8-sig").splitlines()[0]
    assert header.startswith("КАМ;Заявка;Процесс;Вуз / клиент;Этап")
    assert file_name(report, "xlsx") == "overdue_20260921-20261020.xlsx"
    assert file_name(report, "csv") == "overdue_20260921-20261020.csv"
