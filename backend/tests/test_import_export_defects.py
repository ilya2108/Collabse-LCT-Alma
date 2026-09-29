"""Регрессионные тесты дефектов импорта/экспорта (QA-прогон 29.09).

Покрытие: маскирование email в тексте ошибки (ПДн, data-model §8.5), маски ПДн
в экспорте students для ЛЮБОЙ роли (api-contract §8.2), развёртка KPI дашборда
в скаляры, защита от formula injection (CWE-1236), чанкование lookup по
email_hmac (лимит asyncpg 32767 bind-параметров), фиксация status=failed
в отдельной сессии и предупреждение no_data_rows.
"""

from __future__ import annotations

import io
import uuid

from cryptography.fernet import Fernet
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.core.pii import PiiCrypto
from app.models.service import ImportSession
from app.models.talent import Student
from app.models.user import AppUser
from app.modules.import_export import export as export_module
from app.modules.import_export import importers, service
from app.modules.import_export.export import (
    _dataset_report,
    _dataset_students,
    render_export,
)
from app.modules.import_export.importers import (
    ImportContext,
    ImportOptions,
    PreparedRow,
    _apply_students,
    _validate_students,
)
from app.modules.import_export.mapping import STUDENT_FIELDS

HMAC_KEY = "unit-test-hmac-key-0123456789abcdef-0123456789abcdef"


def _crypto() -> PiiCrypto:
    return PiiCrypto([Fernet.generate_key().decode()], HMAC_KEY)


class _ScalarsAll:
    def __init__(self, items: list) -> None:
        self._items = items

    def all(self) -> list:
        return self._items


class _FakeResult:
    def __init__(self, items: list) -> None:
        self._items = items

    def scalars(self) -> _ScalarsAll:
        return _ScalarsAll(self._items)

    def all(self) -> list:
        return self._items


class _FakeSession:
    """Мини-двойник AsyncSession: считает execute и копит add."""

    def __init__(self, result_items: list | None = None) -> None:
        self.executed: list = []
        self.added: list = []
        self._result_items = result_items or []

    async def execute(self, stmt):  # noqa: ANN001
        self.executed.append(stmt)
        return _FakeResult(self._result_items)

    def add(self, obj) -> None:  # noqa: ANN001
        self.added.append(obj)

    async def flush(self) -> None:
        return None


def _ctx(session: _FakeSession) -> ImportContext:
    return ImportContext(
        session=session,  # type: ignore[arg-type]
        spec=STUDENT_FIELDS,
        options=ImportOptions(),
        entries=[],
        actor=AppUser(id=uuid.uuid4(), username="qa"),
        crypto=_crypto(),
        import_session_id=uuid.uuid4(),
    )


class TestEmailMaskedInErrorMessage:
    """Дефект 2: сырой email в тексте invalid_format утекал в БД и XLSX-отчёт."""

    async def test_students_message_masks_email(self) -> None:
        row = PreparedRow(
            row_number=3,
            values={"full_name": "Битый Емейл Строкович", "email": "ivanov@"},
        )
        await _validate_students(_ctx(_FakeSession()), [row])
        issue = next(i for i in row.issues if i.code == "invalid_format")
        assert "ivanov@" not in issue.message
        assert "i***v@" in issue.message  # маска сохраняет узнаваемость для UX

    async def test_b2c_message_masks_email(self) -> None:
        from app.modules.import_export.importers import _validate_requests_b2c

        ctx = _ctx(_FakeSession())
        ctx.lookups = {"itype_by_code": {}, "itype_by_name": {}}
        row = PreparedRow(
            row_number=2,
            values={
                "client_kind": "физлицо",
                "full_name": "Тестов Тест",
                "email": "ivan.petrov@mail,ru",
            },
        )
        await _validate_requests_b2c(ctx, [row])
        issue = next(i for i in row.issues if i.code == "invalid_format")
        assert "ivan.petrov@mail,ru" not in issue.message
        assert "***" in issue.message


class TestStudentsExportMasked:
    """Дефект 4: экспорт students отдавал открытые ПДн для admin/head_kam/kam."""

    async def test_admin_gets_masked_pii(self) -> None:
        student = Student(
            id=uuid.uuid4(),
            full_name="Белова Дарья Сергеевна",
            email="student14@example.com",
            phone="+7 923 101-32-23",
            display_name="Белова Д.С.",
            funnel_status="studying",
        )
        session = _FakeSession(result_items=[])
        session._result_items = [(student, None, None, None)]
        viewer = AppUser(id=uuid.uuid4(), username="admin", roles=["admin"])
        _columns, rows = await _dataset_students(session, {}, viewer)  # type: ignore[arg-type]
        assert rows[0]["full_name"] != "Белова Дарья Сергеевна"
        assert "***" in rows[0]["full_name"]
        assert rows[0]["email"] != "student14@example.com"
        assert "***" in rows[0]["email"]
        assert "***" in rows[0]["phone"]
        # немаскированные колонки остаются полезными
        assert rows[0]["display_name"] == "Белова Д.С."


class TestDashboardKpiExport:
    """Дефекты 6/10: рассинхрон ключей KPI — XLSX 500, CSV с Python-repr словаря."""

    _KPI = {
        "active": {"value": 10, "delta": None},
        "completed": {"value": 5, "delta": 2},
        "conversion": {"value": 50.0, "delta": 1.5},
        "stuck": {"value": 6, "delta": None},
        "period": {"from": "2026-01-01", "to": "2026-01-31"},
    }

    async def _dataset(self, monkeypatch) -> tuple[list, list]:  # noqa: ANN001
        async def _stub_build_report(session, code, viewer, params):  # noqa: ANN001
            return {"kpi": dict(self._KPI)}

        monkeypatch.setattr(export_module, "build_report", _stub_build_report)
        viewer = AppUser(id=uuid.uuid4(), username="admin", roles=["admin"])
        return await _dataset_report(None, "dashboard", {}, viewer)  # type: ignore[arg-type]

    async def test_kpi_flattened_to_scalars(self, monkeypatch) -> None:  # noqa: ANN001
        columns, rows = await self._dataset(monkeypatch)
        assert [key for key, _ in columns] == ["metric", "value", "delta"]
        by_metric = {row["metric"]: row for row in rows}
        assert by_metric["Зависших"]["value"] == 6  # раньше — {'value': 6, 'delta': None}
        assert by_metric["Завершено за период"]["delta"] == 2
        assert by_metric["Конверсия, %"]["value"] == 50.0
        assert all(not isinstance(row["value"], dict) for row in rows)

    async def test_kpi_xlsx_renders_without_error(self, monkeypatch) -> None:  # noqa: ANN001
        columns, rows = await self._dataset(monkeypatch)
        data = render_export(columns, rows, export_format="xlsx")  # раньше — ValueError
        sheet = load_workbook(io.BytesIO(data)).active
        assert sheet.max_row == len(rows) + 1

    async def test_kpi_csv_has_no_dict_repr(self, monkeypatch) -> None:  # noqa: ANN001
        columns, rows = await self._dataset(monkeypatch)
        text = render_export(columns, rows, export_format="csv").decode("utf-8-sig")
        assert "{'value'" not in text
        assert "Зависших;6" in text


class TestFormulaInjection:
    """Дефект 9: CWE-1236 — значения из БД не должны исполняться Excel."""

    COLUMNS = [("a", "A"), ("b", "B")]
    ROWS = [
        {"a": "=cmd|' /c calc'!A1", "b": "@SUM(1,2)"},
        {"a": "-5", "b": "+вперёд"},
    ]

    def test_csv_escapes_formula_prefixes(self) -> None:
        text = render_export(
            self.COLUMNS, self.ROWS, export_format="csv"
        ).decode("utf-8-sig")
        lines = text.splitlines()
        assert lines[1].startswith("'=cmd")
        assert "'@SUM(1,2)" in lines[1]
        assert lines[2].startswith("-5")  # число со знаком — не формула, не портим
        assert "'+вперёд" in lines[2]

    def test_xlsx_stores_formulas_as_text(self) -> None:
        data = render_export(self.COLUMNS, self.ROWS, export_format="xlsx")
        sheet = load_workbook(io.BytesIO(data)).active
        cell = sheet["A2"]
        assert cell.data_type != "f"  # раньше — живая формула (DDE-вектор)
        assert cell.value == "=cmd|' /c calc'!A1"  # значение не искажено


class TestApplyStudentsChunkedLookup:
    """Дефект 7: >32767 bind-параметров — InterfaceError asyncpg на больших файлах."""

    async def test_lookup_is_chunked(self, monkeypatch) -> None:  # noqa: ANN001
        monkeypatch.setattr(importers, "_HMAC_LOOKUP_CHUNK", 2)
        session = _FakeSession()
        rows = [
            PreparedRow(
                row_number=i + 2,
                values={"full_name": f"Тестов Тест {i}", "email": f"u{i}@example.com"},
            )
            for i in range(5)
        ]
        stats = await _apply_students(_ctx(session), rows)
        assert stats.created == 5
        # 5 уникальных hmac при чанке 2 → ровно 3 SELECT-а (ceil(5/2))
        select_count = sum(
            1 for stmt in session.executed if "email_hmac IN" in str(stmt)
        )
        assert select_count == 3

    async def test_duplicate_hmacs_not_reselected(self, monkeypatch) -> None:  # noqa: ANN001
        monkeypatch.setattr(importers, "_HMAC_LOOKUP_CHUNK", 10)
        session = _FakeSession()
        rows = [
            PreparedRow(
                row_number=i + 2,
                values={"full_name": "Тестов Тест", "email": "same@example.com"},
            )
            for i in range(3)
        ]
        stats = await _apply_students(_ctx(session), rows)
        # один hmac → один SELECT; первая строка создаёт, остальные обновляют
        assert stats.created == 1
        assert stats.updated == 2


class TestMarkSessionFailed:
    """Дефект 7 (часть 2): status=failed должен переживать сломанное соединение."""

    class _Recovery:
        def __init__(self, obj) -> None:  # noqa: ANN001
            self.obj = obj
            self.committed = False

        async def __aenter__(self):  # noqa: ANN204
            return self

        async def __aexit__(self, *args):  # noqa: ANN002, ANN204
            return False

        async def get(self, model, pk):  # noqa: ANN001
            assert model is ImportSession
            return self.obj

        async def commit(self) -> None:
            self.committed = True

    async def test_marks_failed_in_fresh_session(self, monkeypatch) -> None:  # noqa: ANN001
        import app.core.db as db

        import_session = ImportSession(
            id=uuid.uuid4(), entity_type="students", status="applying"
        )
        recovery = self._Recovery(import_session)
        monkeypatch.setattr(db, "get_session_factory", lambda: lambda: recovery)
        await service._mark_session_failed(import_session.id)
        assert import_session.status == "failed"
        assert import_session.error_message == "Ошибка применения импорта"
        assert import_session.finished_at is not None
        assert recovery.committed

    async def test_completed_session_not_overwritten(self, monkeypatch) -> None:  # noqa: ANN001
        import app.core.db as db

        import_session = ImportSession(
            id=uuid.uuid4(), entity_type="students", status="completed"
        )
        recovery = self._Recovery(import_session)
        monkeypatch.setattr(db, "get_session_factory", lambda: lambda: recovery)
        await service._mark_session_failed(import_session.id)
        assert import_session.status == "completed"
        assert not recovery.committed

    async def test_persist_error_is_swallowed(self, monkeypatch) -> None:  # noqa: ANN001
        import app.core.db as db

        def _boom():  # noqa: ANN202
            raise RuntimeError("pool is dead")

        monkeypatch.setattr(db, "get_session_factory", _boom)
        # исходная ошибка применения важнее — новая не должна её маскировать
        await service._mark_session_failed(uuid.uuid4())


class TestUniversitiesCoverageSql:
    """Дефекты 5/11: GROUP BY должен ссылаться на то же выражение, что и SELECT."""

    async def test_group_by_uses_label(self) -> None:
        from app.modules.reporting.service import ReportParams, report_universities_coverage

        session = _FakeSession(result_items=[])
        await report_universities_coverage(session, ReportParams(), None)  # type: ignore[arg-type]
        stmt = session.executed[0]
        sql = str(stmt.compile(dialect=postgresql.dialect()))
        select_part = sql.split("GROUP BY", 1)[0]
        group_by = sql.split("GROUP BY", 1)[1].split("ORDER BY", 1)[0].strip()
        # причина GroupingError: литерал «Без региона» рендерился РАЗНЫМИ
        # bind-параметрами в SELECT и GROUP BY; одно labeled-выражение даёт
        # один общий параметр — Postgres сопоставляет выражения дословно
        assert group_by.startswith("coalesce(university.region")
        param = group_by.split("%(")[1].split(")")[0]
        assert f"%({param})s" in select_part  # тот же параметр, что и в SELECT
        assert sql.count("%(coalesce") == 2  # ровно одно имя в двух позициях


class TestNoDataRowsWarning:
    """Дефект 12: headers-only файл — не 500, но и не молчаливый «успех нуля строк»."""

    def test_warning_shape(self) -> None:
        warning = service._no_data_rows_warning()
        assert warning["code"] == "no_data_rows"
        assert "нет строк данных" in warning["message"]
