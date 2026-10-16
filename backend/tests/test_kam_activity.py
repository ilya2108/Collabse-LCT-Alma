"""Отчёт «Активность КАМов» (F1): чистые функции, компиляция SQL и рендер файлов без БД."""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Any

import pytest
from pydantic import ValidationError
from sqlalchemy.dialects import postgresql

from app.core.access import DealScope, compute_access, static_catalog
from app.core.errors import ApiError
from app.models.user import AppUser
from app.modules.import_export.schemas import ExportJobBody
from app.modules.reporting import kam_activity as ka
from app.modules.reporting.kam_activity import (
    ActivityStat,
    DealRow,
    KamActivityParams,
    PersonRef,
    StageRef,
    StageStat,
    UniversityRef,
)
from app.modules.reporting.kam_activity_export import (
    SHEET_NAMES,
    ExportMeta,
    describe_filters,
    kam_activity_flat_rows,
    render_kam_activity_pdf,
    render_kam_activity_xlsx,
)
from app.modules.reporting.pdf import render_table_pdf
from app.modules.reporting.schemas import KamActivityReportOut, KamActivitySummaryOut
from app.modules.reporting.service import ReportParams, _closed_between_stmt, report_scope

MSK = ka.REPORT_TZ
NOW = datetime(2026, 10, 20, 12, 0, tzinfo=MSK)
TODAY = NOW.date()


def _user(*roles: str) -> AppUser:
    return AppUser(
        id=uuid.uuid4(),
        keycloak_sub=uuid.uuid4(),
        username=f"{roles[0] if roles else 'none'}@test",
        full_name="Тестов Тест",
        roles=list(roles),
        is_active=True,
    )


def _bounds(date_from: date, date_to: date, now: datetime = NOW) -> ka.PeriodBounds:
    params = ka.resolve_params(KamActivityParams(date_from=date_from, date_to=date_to), today=TODAY)
    return ka.resolve_bounds(params, now=now)


# --- 1–3: период, предыдущий период, гранулярность ----------------------------------------


class TestPeriod:
    def test_default_is_last_30_days(self) -> None:
        params = ka.resolve_params(KamActivityParams(), today=TODAY)
        assert (params.date_from, params.date_to) == (date(2026, 9, 21), TODAY)
        only_to = ka.resolve_params(KamActivityParams(date_to=date(2026, 9, 30)), today=TODAY)
        assert only_to.date_from == date(2026, 9, 1)

    @pytest.mark.parametrize(
        ("date_from", "date_to", "code", "field"),
        [
            (date(2026, 10, 10), date(2026, 10, 1), "period_inverted", "from"),
            (date(2024, 1, 1), date(2026, 1, 2), "period_too_long", "to"),
            (date(2026, 11, 1), date(2026, 11, 30), "period_in_future", "from"),
        ],
    )
    def test_validation_errors(self, date_from, date_to, code, field) -> None:  # noqa: ANN001
        with pytest.raises(ApiError) as excinfo:
            ka.resolve_params(KamActivityParams(date_from=date_from, date_to=date_to), today=TODAY)
        error = excinfo.value
        assert error.status_code == 400 and error.code == "validation_error"
        assert any(d["code"] == code and d["field"] == field for d in error.details)

    def test_two_years_is_the_limit(self) -> None:
        ka.resolve_params(
            KamActivityParams(date_from=date(2024, 10, 20), date_to=date(2026, 10, 20)),
            today=TODAY,
        )  # 731 день — можно

    def test_too_many_ids(self) -> None:
        ids = tuple(uuid.uuid4() for _ in range(ka.MAX_FILTER_IDS + 1))
        with pytest.raises(ApiError) as excinfo:
            ka.resolve_params(KamActivityParams(kam_ids=ids), today=TODAY)
        assert excinfo.value.details[0]["code"] == "too_many_values"

    def test_bounds_are_moscow_midnight(self) -> None:
        bounds = _bounds(date(2026, 7, 1), date(2026, 9, 30))
        assert bounds.start.utcoffset() == timedelta(hours=3)
        assert bounds.start == datetime(2026, 7, 1, tzinfo=MSK)
        assert bounds.end == datetime(2026, 10, 1, tzinfo=MSK)
        assert bounds.start.astimezone(timezone.utc) == datetime(
            2026, 6, 30, 21, tzinfo=timezone.utc
        )
        assert bounds.as_of == bounds.end and bounds.days == 92
        current = _bounds(date(2026, 10, 1), TODAY)
        assert current.as_of == NOW  # период включает сегодня — «данные на» сейчас

    def test_previous_bounds_same_length(self) -> None:
        q3 = _bounds(date(2026, 7, 1), date(2026, 9, 30))
        prev = ka.previous_bounds(q3)
        assert (prev.date_from, prev.date_to) == (date(2026, 3, 31), date(2026, 6, 30))
        assert prev.days == q3.days and prev.as_of == prev.end

    @pytest.mark.parametrize(
        ("days", "expected"), [(1, "day"), (31, "day"), (32, "week"), (184, "week"), (185, "month")]
    )
    def test_granularity(self, days: int, expected: str) -> None:
        assert ka.pick_granularity(days) == expected


# --- 4–5: зона видимости, фильтры выгрузки, ключ кэша --------------------------------------


def _own(user_id: uuid.UUID) -> DealScope:
    return DealScope(rules=(("b2b", user_id), ("b2c", user_id)))


def _report_scope(roles: tuple[str, ...], user_id: uuid.UUID) -> DealScope:
    """Область отчётов по правам ролей из статического каталога (как `report_scope`)."""
    access = compute_access(roles, static_catalog())
    return access.deal_scope("reports:read", user_id, wide_permission="reports:all")


class TestScope:
    def test_kam_without_filter_sees_only_self(self) -> None:
        kam = _user("kam")
        scope = _report_scope(("kam",), kam.id)
        params = ka._resolve_scope(kam, KamActivityParams(), scope)
        assert params.kam_ids == (kam.id,)
        # без явной области — та же, что у дашборда (report_scope зрителя)
        assert ka._resolve_scope(kam, KamActivityParams()).kam_ids == (kam.id,)
        assert ka.viewer_scope(kam) == report_scope(kam) == scope

    def test_kam_foreign_id_is_403(self) -> None:
        kam = _user("kam")
        scope = _report_scope(("kam",), kam.id)
        with pytest.raises(ApiError) as excinfo:
            ka._resolve_scope(kam, KamActivityParams(kam_ids=(kam.id, uuid.uuid4())), scope)
        assert excinfo.value.status_code == 403
        assert excinfo.value.message == "КАМ видит отчёт только по своим заявкам"

    def test_kam_own_id_is_ok(self) -> None:
        kam = _user("kam")
        scope = _report_scope(("kam",), kam.id)
        params = ka._resolve_scope(kam, KamActivityParams(kam_ids=(kam.id,)), scope)
        assert params.kam_ids == (kam.id,)

    @pytest.mark.parametrize(
        "roles", [("head_kam",), ("observer",), ("admin",), ("kam", "observer")]
    )
    def test_wide_roles_keep_filters(self, roles: tuple[str, ...]) -> None:
        other = uuid.uuid4()
        viewer = _user(*roles)
        scope = _report_scope(roles, viewer.id)
        assert scope.is_unrestricted
        assert ka._resolve_scope(viewer, KamActivityParams(), scope).kam_ids == ()
        assert ka._resolve_scope(viewer, KamActivityParams(kam_ids=(other,)), scope).kam_ids == (
            other,
        )

    def test_workflow_outside_scope_is_403(self) -> None:
        scope = DealScope(rules=(("b2c", None),))
        with pytest.raises(ApiError) as excinfo:
            ka._resolve_scope(_user("head_kam"), KamActivityParams(workflow_type="b2b"), scope)
        assert excinfo.value.status_code == 403
        assert excinfo.value.message == "Процесс B2B недоступен вашей роли"

    def test_empty_scope_is_403(self) -> None:
        with pytest.raises(ApiError) as excinfo:
            ka._resolve_scope(_user("kam"), KamActivityParams(), DealScope(rules=()))
        assert excinfo.value.status_code == 403

    def test_mixed_scope_keeps_foreign_kams(self) -> None:
        """b2b — все, b2c — только свои: чужой КАМ в фильтре допустим (b2c режет apply)."""
        me, other = uuid.uuid4(), uuid.uuid4()
        scope = DealScope(rules=(("b2b", None), ("b2c", me)))
        params = ka._resolve_scope(_user("kam"), KamActivityParams(kam_ids=(other,)), scope)
        assert params.kam_ids == (other,)
        sql = _sql(ka.deal_rows_stmt(params, PERIOD, 14, scope))
        assert "deal.pipeline = " in sql and "deal.responsible_user_id = " in sql

    def test_kam_role_codes_follow_catalog(self) -> None:
        both = ka.kam_role_codes(KamActivityParams(), DealScope.unrestricted())
        assert both == ("kam",)  # head_kam/admin без requests:kam — не КАМы отчёта
        assert ka.kam_role_codes(KamActivityParams(), DealScope(rules=())) == ()

    def test_from_filters(self) -> None:
        a, b = uuid.uuid4(), uuid.uuid4()
        params = KamActivityParams.from_filters(
            {"from": "2026-07-01", "to": "2026-09-30", "kam_id": str(a), "university_id": [str(b)]}
        )
        assert params.kam_ids == (a,) and params.university_ids == (b,)
        assert params.date_from == date(2026, 7, 1)
        url_state = KamActivityParams.from_filters({"wf": "b2c", "kam": [str(a), str(b)]})
        assert url_state.workflow_type == "b2c" and set(url_state.kam_ids) == {a, b}
        with pytest.raises(ApiError) as excinfo:
            KamActivityParams.from_filters({"kam_id": ["не-uuid"]})
        assert excinfo.value.status_code == 400
        assert excinfo.value.details[0]["code"] == "invalid_uuid"
        with pytest.raises(ApiError):
            KamActivityParams.from_filters({"workflow_type": "b2x"})

    def test_as_json_and_cache_key(self) -> None:
        a, b = uuid.uuid4(), uuid.uuid4()
        first = KamActivityParams(kam_ids=(a, b))
        second = KamActivityParams(kam_ids=(b, a))
        assert first.as_json()["kam_id"] == sorted([str(a), str(b)])
        wide, own = DealScope.unrestricted(), _own(a)
        assert ka._cache_key("x", first, wide) == ka._cache_key("x", second, wide)
        assert ka._cache_key("x", first, wide) != ka._cache_key("x", first, own)
        assert ka._cache_key("x", first, wide).startswith("dash:x:")


# --- 6: сборка дерева ---------------------------------------------------------------------

KAM1, KAM2 = uuid.uuid4(), uuid.uuid4()
MIPT, BMSTU, ITMO = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
STAGE_NEG, STAGE_NEW, STAGE_WON, STAGE_LOST = (uuid.uuid4() for _ in range(4))
B2C_NEW = uuid.uuid4()
PERIOD = _bounds(date(2026, 9, 21), TODAY)


def deal_row(**overrides: Any) -> DealRow:
    base: dict[str, Any] = {
        "deal_id": uuid.uuid4(),
        "title": "МФТИ — школа Python",
        "pipeline": "b2b",
        "university_id": MIPT,
        "university_name": "Московский физико-технический институт",
        "university_short": "МФТИ",
        "university_region": "Московская область",
        "counterparty_id": None,
        "client": None,
        "responsible_id": KAM1,
        "amount": Decimal("450000.00"),
        "status": "open",
        "created_at": NOW - timedelta(days=100),
        "stage_id": STAGE_NEG,
        "stage_code": "negotiation",
        "stage_name": "Переговоры",
        "stage_color": "#1677ff",
        "stage_is_terminal": False,
        "stage_outcome": None,
        "stage_sort": 2,
        "stage_deleted": False,
        "stage_since": NOW - timedelta(days=6),
        "threshold": 14,
        "transitions": 1,
        "returns": 0,
        "comments": 2,
        "files": 1,
        "last_contact_at": NOW - timedelta(days=2),
    }
    base.update(overrides)
    return DealRow(**base)


PEOPLE = {
    KAM1: PersonRef(KAM1, "Петров Константин Андреевич", ("kam",)),
    KAM2: PersonRef(KAM2, "Иванова Мария Сергеевна", ("kam",)),
}


def sample_rows() -> list[DealRow]:
    return [
        deal_row(),
        deal_row(  # зависшая на «Новая» (порог 7)
            title="МФТИ — стажировки",
            stage_id=STAGE_NEW,
            stage_code="new",
            stage_name="Новая",
            stage_sort=0,
            stage_since=NOW - timedelta(days=10),
            threshold=7,
            transitions=0,
            comments=0,
            files=0,
            last_contact_at=None,
            created_at=NOW - timedelta(days=10),
        ),
        deal_row(  # выиграна в периоде
            title="Бауманка — ИИ-инженер",
            university_id=BMSTU,
            university_name="МГТУ им. Н.Э. Баумана",
            university_short="Бауманка",
            stage_id=STAGE_WON,
            stage_code="completed",
            stage_name="Завершена",
            stage_is_terminal=True,
            stage_outcome="won",
            stage_sort=6,
            status="won",
            stage_since=NOW - timedelta(days=3),
        ),
        deal_row(  # B2C kam2
            title="Смирнова О.В. — Практикум",
            pipeline="b2c",
            university_id=None,
            university_name=None,
            university_short=None,
            counterparty_id=uuid.uuid4(),
            client="Смирнова О. В.",
            responsible_id=KAM2,
            stage_id=B2C_NEW,
            stage_code="new",
            stage_name="Новая заявка",
            stage_sort=0,
            created_at=NOW - timedelta(days=5),
            stage_since=NOW - timedelta(days=5),
        ),
    ]


def sample_tree(include_requests: bool = True) -> dict[str, Any]:
    activities = [
        ActivityStat(KAM1, MIPT, "МФТИ", "МФТИ", None, 2, NOW - timedelta(days=1)),
        # давнее мероприятие без заявок и закрепления — строку не заводит
        ActivityStat(KAM2, BMSTU, "Бауманка", "Бауманка", None, 0, NOW - timedelta(days=200)),
    ]
    assigned = [
        UniversityRef(MIPT, "Московский физико-технический институт", "МФТИ", None, KAM1),
        UniversityRef(ITMO, "Университет ИТМО", "ИТМО", "Санкт-Петербург", KAM1),
    ]
    return ka.assemble_tree(
        sample_rows(), activities, assigned, PEOPLE, PERIOD, include_requests=include_requests
    )


class TestAssembleTree:
    def test_grouping_and_rollups(self) -> None:
        tree = sample_tree()
        kam1, kam2 = tree["rows"]
        assert [r["kam"] for r in tree["rows"]] == sorted(
            (p.full_name for p in PEOPLE.values()), key=str.casefold
        )
        kam1 = next(r for r in tree["rows"] if r["kam_id"] == str(KAM1))
        names = [node["short_name"] for node in kam1["universities"]]
        assert names[0] == "ИТМО"  # закреплён и без контакта — первым
        mipt = next(n for n in kam1["universities"] if n["university_id"] == str(MIPT))
        assert mipt["assigned"] is True and mipt["requests_count"] == 2 and mipt["stuck"] == 1
        assert mipt["interactions"]["activities"] == 2
        assert mipt["interactions"]["total"] == 1 + 2 + 1 + 2
        assert mipt["last_contact_at"] == (NOW - timedelta(days=1)).isoformat()
        assert [chip["code"] for chip in mipt["stages"]] == ["new", "negotiation"]
        assert mipt["requests"][0]["is_stuck"] is True  # зависшие сначала
        itmo = kam1["universities"][0]
        assert itmo["interactions"]["total"] == 0 and itmo["requests"] == []
        assert kam1["universities_assigned"] == 2 and kam1["assigned_untouched"] == 1
        assert kam1["coverage_pct"] == 50.0
        assert kam1["won"] == 1 and kam1["lost"] == 0 and kam1["conversion"] == 100.0
        assert kam1["new_requests"] == 1 and kam1["new_amount"] == "450000.00"

    def test_b2c_node_and_no_closed(self) -> None:
        tree = sample_tree()
        kam2 = next(r for r in tree["rows"] if r["kam_id"] == str(KAM2))
        assert [n["university_id"] for n in kam2["universities"]] == [None]
        node = kam2["universities"][0]
        assert node["name"] == ka.B2C_NODE_NAME
        assert node["requests"][0]["client"] == "Смирнова О. В."
        assert kam2["conversion"] is None  # закрытых нет — «—»

    def test_request_metrics(self) -> None:
        tree = sample_tree()
        requests = {
            r["title"]: r
            for row in tree["rows"]
            for node in row["universities"]
            for r in node["requests"]
        }
        stuck = requests["МФТИ — стажировки"]
        assert stuck["is_stuck"] is True and stuck["days_on_stage"] == 10
        assert stuck["created_in_period"] is True
        won = requests["Бауманка — ИИ-инженер"]
        assert won["closed_in_period"] == "won" and won["is_stuck"] is False
        assert not any(key.startswith("_") for key in won)

    def test_totals(self) -> None:
        totals = sample_tree()["totals"]
        assert totals["kams"] == 2
        assert totals["universities_assigned"] == 2 and totals["assigned_touched"] == 1
        assert totals["universities_touched"] == 2  # МФТИ и Бауманка
        assert totals["coverage_pct"] == 50.0
        assert totals["requests_count"] == 4 and totals["stuck"] == 1

    def test_truncation(self) -> None:
        tree = ka.public_tree(sample_tree(include_requests=False))
        nodes = [node for row in tree["rows"] for node in row["universities"]]
        assert all(node["requests_truncated"] for node in nodes)
        assert all(node["requests"] == [] for node in nodes)
        assert any(node["requests_count"] > 0 for node in nodes)

    def test_person_without_data_is_a_row(self) -> None:
        idle = uuid.uuid4()
        tree = ka.assemble_tree(
            [], [], [], {**PEOPLE, idle: PersonRef(idle, "Сидоров Иван", ("kam",))}, PERIOD
        )
        assert {row["kam_id"] for row in tree["rows"]} == {str(KAM1), str(KAM2), str(idle)}
        assert all(row["interactions"]["total"] == 0 for row in tree["rows"])

    def test_deleted_stage_name(self) -> None:
        row = deal_row(stage_name="Первичный контакт (удалён)", stage_deleted=True)
        request = ka.request_item(row, PERIOD)
        assert request["stage"]["deleted"] is True
        assert request["stage"]["name"].endswith("(удалён)")


# --- 7: KPI, дельты, выводы ---------------------------------------------------------------


def _kpi(**overrides: Any) -> ka.PeriodKpi:
    base: dict[str, Any] = {
        "coverage_pct": 70.0,
        "universities_touched": 7,
        "universities_assigned": 10,
        "assigned_touched": 7,
        "requests_count": 21,
        "new_requests": 6,
        "new_amount": Decimal("3120000.00"),
        "won": 3,
        "lost": 1,
        "conversion": 75.0,
        "interactions": 51,
        "stuck": 4,
        "talent_pool_added": 0,
    }
    base.update(overrides)
    return ka.PeriodKpi(**base)


class TestKpiAndHighlights:
    def test_kpi_with_delta(self) -> None:
        cur = _kpi()
        prev = _kpi(
            coverage_pct=50.0,
            assigned_touched=5,
            new_requests=4,
            conversion=None,
            new_amount=Decimal("1000000.00"),
        )
        kpi = ka.kpi_with_delta(cur, prev)
        assert kpi["coverage_pct"] == {"value": 70.0, "prev": 50.0, "delta": 20.0}
        assert kpi["new_requests"] == {"value": 6, "prev": 4, "delta": 2}
        assert kpi["conversion"]["delta"] is None  # нет базы — нет дельты
        assert kpi["new_amount"] == {
            "value": "3120000.00",
            "prev": "1000000.00",
            "delta": "2120000.00",
        }

    def test_summarize_period(self) -> None:
        kpi = ka.summarize_period(sample_tree(), talent_pool_added=3)
        assert kpi.requests_count == 4 and kpi.won == 1 and kpi.talent_pool_added == 3
        assert kpi.new_amount == Decimal("900000.00")

    @pytest.mark.parametrize(
        ("n", "form"),
        [(1, "вуз"), (2, "вуза"), (5, "вузов"), (11, "вузов"), (21, "вуз"), (22, "вуза")],
    )
    def test_plural(self, n: int, form: str) -> None:
        assert ka._plural(n, ("вуз", "вуза", "вузов")) == form

    def _summary(self, cur: ka.PeriodKpi, prev: ka.PeriodKpi, **extra: Any) -> dict[str, Any]:
        return {"kpi": ka.kpi_with_delta(cur, prev), **extra}

    def test_highlights_rules(self) -> None:
        stages = {
            "b2b": [
                {
                    "name": "Переговоры",
                    "is_terminal": False,
                    "avg_days": 12.5,
                    "intervals": 4,
                    "stuck": 2,
                },
                {
                    "name": "Новая",
                    "is_terminal": False,
                    "avg_days": 30.0,
                    "intervals": 1,
                    "stuck": 1,
                },
            ]
        }
        lines = ka.build_highlights(
            self._summary(
                _kpi(),
                _kpi(assigned_touched=5, new_requests=4, conversion=50.0),
                stages=stages,
                untouched_assigned=["ИТМО", "ТГУ", "НГУ", "ЮФУ"],
                talent={"talent_pool_added": {"value": 2}},
            )
        )
        assert len(lines) == 5  # лимит: вывод про пул талантов не влез
        assert lines[0] == ("Работали с 7 из 10 закреплённых вузов (70%), к прошлому периоду +2.")
        assert lines[1] == "Новых заявок: 6 (+2) на сумму 3 120 000 ₽."
        assert lines[2] == "Выиграно 3, отказов 1, конверсия 75% (+25 п.п.)."
        # «Новая» — 1 интервал, не в счёт; берётся «Переговоры»
        assert lines[3].startswith("Дольше всего заявки задерживаются на этапе «Переговоры»")
        assert (
            lines[4] == "Без взаимодействий остались закреплённые вузы: ИТМО, ТГУ, НГУ (и ещё 1)."
        )

    def test_highlights_no_closed_and_singular(self) -> None:
        lines = ka.build_highlights(
            self._summary(
                _kpi(
                    won=0,
                    lost=0,
                    conversion=None,
                    universities_assigned=1,
                    assigned_touched=1,
                    coverage_pct=100.0,
                    new_requests=0,
                ),
                _kpi(
                    won=0,
                    lost=0,
                    conversion=None,
                    universities_assigned=1,
                    assigned_touched=1,
                    coverage_pct=100.0,
                ),
                talent={"talent_pool_added": {"value": 1}},
            )
        )
        assert lines[0] == "Работали с 1 из 1 закреплённого вуза (100%), к прошлому периоду ±0."
        assert "Новых заявок за период нет." in lines
        assert "Закрытых заявок за период нет." in lines
        assert lines[-1] == "В пул талантов за период добавлен 1 студент."


# --- 8–9: этапы и сетка динамики -----------------------------------------------------------


def _stage(
    code: str,
    sort: int,
    *,
    terminal: bool = False,
    outcome: str | None = None,
    initial: bool = False,
    deleted: bool = False,
) -> StageRef:
    return StageRef(
        uuid.uuid4(), "b2b", code, code.title(), None, sort, terminal, outcome, deleted, initial
    )


class TestStagesAndBuckets:
    def test_stage_rows(self) -> None:
        new = _stage("new", 0, initial=True)
        contact = _stage("contact", 1)
        offer = _stage("offer", 2)
        gone = _stage("gone", 3, deleted=True)  # удалён и в истории нет — не показываем
        won = _stage("won", 6, terminal=True, outcome="won")
        lost = _stage("lost", 7, terminal=True, outcome="lost")
        stats = {
            new.stage_id: StageStat(10, 2.0, 10),
            contact.stage_id: StageStat(8, 5.5, 8),
            offer.stage_id: StageStat(0, None, 0),
            won.stage_id: StageStat(2, None, 2),
            lost.stage_id: StageStat(3, None, 3),
        }
        tree = {"rows": []}
        rows = ka.stage_rows({"b2b": [lost, won, offer, gone, contact, new]}, stats, tree)["b2b"]
        assert [r["code"] for r in rows] == ["new", "contact", "offer", "won", "lost"]
        assert rows[0]["share_of_first_pct"] == 100.0 and rows[0]["step_conversion_pct"] is None
        assert rows[1]["step_conversion_pct"] == 80.0 and rows[1]["share_of_first_pct"] == 80.0
        assert rows[2]["step_conversion_pct"] == 0.0
        assert rows[3]["step_conversion_pct"] is None  # 2 / 0 — деление на ноль
        assert rows[3]["share_of_first_pct"] == 20.0 and rows[3]["avg_days"] is None
        assert rows[4]["step_conversion_pct"] is None

    def test_stage_rows_no_entries(self) -> None:
        new = _stage("new", 0, initial=True)
        rows = ka.stage_rows({"b2b": [new]}, {}, {"rows": []})["b2b"]
        assert rows[0]["reached"] == 0 and rows[0]["share_of_first_pct"] is None

    def test_fill_buckets(self) -> None:
        days = ka.fill_buckets(_bounds(date(2026, 9, 21), date(2026, 10, 20)), "day")
        assert len(days) == 30 and days[0] == date(2026, 9, 21) and days[-1] == TODAY
        weeks = ka.fill_buckets(_bounds(date(2026, 7, 1), date(2026, 9, 30)), "week")
        assert weeks[0] == date(2026, 6, 29) and weeks[-1] == date(2026, 9, 28)
        assert all((b - a).days == 7 for a, b in zip(weeks, weeks[1:], strict=False))
        months = ka.fill_buckets(_bounds(date(2025, 11, 15), date(2026, 10, 1)), "month")
        assert months[0] == date(2025, 11, 1) and months[-1] == date(2026, 10, 1)
        assert len(months) == 12

    def test_dynamics_series_fills_gaps(self) -> None:
        tree = sample_tree()
        events = {"comments": {date(2026, 10, 1): 4}}
        series = {s["key"]: s for s in ka.dynamics_series(events, tree, PERIOD, "day")}
        comments = series["comments"]["points"]
        assert len(comments) == 30 and sum(p["value"] for p in comments) == 4
        assert {p["date"]: p["value"] for p in comments}["2026-10-01"] == 4
        assert sum(p["value"] for p in series["new_requests"]["points"]) == 2
        assert sum(p["value"] for p in series["activities"]["points"]) == 0


# --- 10: компиляция SQL ---------------------------------------------------------------------


def _sql(stmt: Any) -> str:
    return str(stmt.compile(dialect=postgresql.dialect()))


class TestSql:
    def test_deal_rows_stmt(self) -> None:
        sql = _sql(ka.deal_rows_stmt(KamActivityParams(), PERIOD, 14, DealScope.unrestricted()))
        assert sql.count("LATERAL") == 5
        assert "dsh_own.actor_user_id = deal.responsible_user_id" in sql
        assert "deal_comment.author_user_id = deal.responsible_user_id" in sql
        assert "file_object.uploaded_by = deal.responsible_user_id" in sql
        assert "deal.status != " in sql
        assert "deal.pipeline =" not in sql and "deal.responsible_user_id IN" not in sql
        compiled = ka.deal_rows_stmt(KamActivityParams(), PERIOD, 14).compile(
            dialect=postgresql.dialect()
        )
        assert "manual" in compiled.params.values() and "archived" in compiled.params.values()

    def test_deal_rows_filters_only_when_set(self) -> None:
        own = uuid.uuid4()
        params = KamActivityParams(
            workflow_type="b2b", kam_ids=(own,), university_ids=(uuid.uuid4(),)
        )
        sql = _sql(ka.deal_rows_stmt(params, PERIOD, 14, _own(own)))
        assert "deal.pipeline =" in sql
        assert "deal.responsible_user_id IN" in sql
        assert "deal.university_id IN" in sql
        assert "deal.responsible_user_id = " in sql  # скоуп «только свои»

    def test_dynamics_group_by_alias_and_literal_unit(self) -> None:
        stmt = ka.dynamics_stmt(
            [uuid.uuid4(), uuid.uuid4()],
            KamActivityParams(),
            PERIOD,
            "week",
            activities_visible=True,
        )
        sql = _sql(stmt)
        assert "date_trunc('week', timezone('Europe/Moscow'" in sql
        group_by = sql.rsplit("GROUP BY", 1)[1]
        assert group_by.strip() == "events.kind, events.bucket"
        assert "date_trunc" not in group_by
        assert sql.count("UNION ALL") == 3
        assert "%(deal_ids)s" in sql  # один параметр-массив на все подзапросы
        assert sql.count("= ANY (%(deal_ids)s") == 3  # история, комментарии, файлы
        with pytest.raises(ValueError):
            ka.dynamics_stmt(
                [], KamActivityParams(), PERIOD, "week'; drop", activities_visible=False
            )

    def test_stage_stats_single_array_param(self) -> None:
        ids = [uuid.uuid4() for _ in range(40_000)]  # больше лимита asyncpg на параметры
        compiled = ka.stage_reached_stmt(ids, PERIOD).compile(dialect=postgresql.dialect())
        assert list(compiled.params) == ["deal_ids", "created_at_1"]
        dwell = _sql(ka.stage_dwell_stmt(ids, PERIOD))
        assert "lead(deal_stage_history.created_at) OVER (PARTITION BY" in dwell


# --- 14: D16 — «Завершено за период» дашборда по истории переходов -----------------------


class TestDashboardClosedKpi:
    def test_closed_between_uses_stage_history(self) -> None:
        stmt = _closed_between_stmt(
            "won",
            date(2026, 7, 1),
            date(2026, 9, 30),
            ReportParams(workflow_type="b2b"),
            DealScope.unrestricted(),
        )
        sql = _sql(stmt)
        where = sql.split("WHERE", 1)[1]
        assert "deal_stage_history" in where
        assert "workflow_stage.terminal_outcome = " in where
        assert "updated_at" not in sql
        compiled = stmt.compile(dialect=postgresql.dialect())
        assert "won" in compiled.params.values()

    def test_closed_between_keeps_scope_and_filters(self) -> None:
        own = uuid.uuid4()
        sql = _sql(_closed_between_stmt("lost", None, None, ReportParams(kam_id=own), _own(own)))
        assert "deal.responsible_user_id = " in sql
        assert "deal_stage_history" not in sql  # без периода подзапрос не нужен


# --- 11–12: выгрузка ---------------------------------------------------------------------


def sample_reports(rows: list[DealRow] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    rows = sample_rows() if rows is None else rows
    tree = ka.assemble_tree(rows, [], [], PEOPLE, PERIOD)
    prev_tree = ka.assemble_tree([], [], [], PEOPLE, ka.previous_bounds(PERIOD))
    stage_refs = {
        "b2b": [
            StageRef(STAGE_NEW, "b2b", "new", "Новая", "#8c8c8c", 0, False, None, False, True),
            StageRef(
                STAGE_NEG, "b2b", "negotiation", "Переговоры", "#1677ff", 2, False, None, False
            ),
            StageRef(STAGE_WON, "b2b", "completed", "Завершена", "#237804", 6, True, "won", False),
            StageRef(STAGE_LOST, "b2b", "rejected", "Отказ", "#f5222d", 7, True, "lost", False),
        ]
    }
    stats = {
        STAGE_NEW: StageStat(3, 4.0, 3),
        STAGE_NEG: StageStat(2, 9.5, 2),
        STAGE_WON: StageStat(1, None, 1),
    }
    params = KamActivityParams(date_from=PERIOD.date_from, date_to=PERIOD.date_to)
    scope = {"own_only": False, "own_user_id": None, "workflow_types": ["b2b", "b2c"]}
    detail = {
        "report": ka.CODE_DETAIL,
        "generated_at": NOW.isoformat(),
        "params": params.as_json(),
        "period": PERIOD.as_json(),
        "scope": scope,
        **ka.public_tree(tree),
    }
    talent = ka._empty_talent()
    talent["programs"] = [
        {
            "program_id": str(uuid.uuid4()),
            "name": "Практикум по ML",
            "university": "МФТИ",
            "candidate": 1,
            "studying": 4,
            "graduate": 2,
            "talent_pool": 3,
            "total": 10,
        }
    ]
    body: dict[str, Any] = {
        "period": {**PERIOD.as_json(), "granularity": "day"},
        "previous_period": ka.previous_bounds(PERIOD).as_json(),
        "scope": scope,
        "kpi": ka.kpi_with_delta(ka.summarize_period(tree), ka.summarize_period(prev_tree)),
        "stages": ka.stage_rows(stage_refs, stats, tree),
        "dynamics": {"granularity": "day", "series": ka.dynamics_series({}, tree, PERIOD, "day")},
        "by_kam": ka.by_kam_rows(tree, prev_tree),
        "top_stuck": ka.top_stuck(tree),
        "untouched_assigned": ka.untouched_assigned(tree),
        "talent": talent,
    }
    body["highlights"] = ka.build_highlights(body)
    summary = {
        "report": ka.CODE_SUMMARY,
        "generated_at": NOW.isoformat(),
        "params": params.as_json(),
        **body,
    }
    KamActivityReportOut.model_validate(detail)
    KamActivitySummaryOut.model_validate(summary)
    return detail, summary


META = ExportMeta(generated_by="Морозов Дмитрий", generated_at=NOW, filters_text="КАМ: все")


class TestExport:
    def test_xlsx_sheets_types_and_formula_injection(self) -> None:
        import openpyxl

        rows = sample_rows()
        rows[0] = deal_row(title="=cmd|' /c calc'!A1")
        detail, summary = sample_reports(rows)
        book = openpyxl.load_workbook(io.BytesIO(render_kam_activity_xlsx(detail, summary, META)))
        assert tuple(book.sheetnames) == SHEET_NAMES
        requests = book["Заявки"]
        header = [cell.value for cell in requests[1]]
        assert header[:3] == ["КАМ", "Вуз / клиент", "Заявка"]
        titles = {requests.cell(row=r, column=3).value: r for r in range(2, requests.max_row + 1)}
        injected = titles["=cmd|' /c calc'!A1"]
        assert requests.cell(row=injected, column=3).data_type == "s"
        assert isinstance(requests.cell(row=injected, column=6).value, datetime)  # дата
        assert requests.freeze_panes == "A2" and requests.auto_filter.ref
        summary_sheet = book["Сводка"]
        assert summary_sheet["A1"].value == "Отчёт об активности КАМов"
        assert summary_sheet._charts  # нативная диаграмма «Взаимодействия по КАМам»
        kams = book["КАМы"]
        assert kams["A1"].value == "КАМ" and kams.max_row == 3

    def test_pdf_renders_cyrillic_and_arrows(self) -> None:
        detail, summary = sample_reports()
        data = render_kam_activity_pdf(detail, summary, META)
        assert data.startswith(b"%PDF") and len(data) > 5000

    def test_empty_report_renders(self) -> None:
        detail, summary = sample_reports([])
        assert render_kam_activity_pdf(detail, summary, META).startswith(b"%PDF")
        assert render_kam_activity_xlsx(detail, summary, META)[:2] == b"PK"
        columns, flat = kam_activity_flat_rows(detail)
        assert flat == [] and columns[0] == ("kam", "КАМ")

    def test_table_pdf(self) -> None:
        data = render_table_pdf(
            "Нагрузка по КАМам ▲▼",
            [("kam", "КАМ"), ("value", "Заявок")],
            [{"kam": "Петров К. А.", "value": 3}, {"kam": None, "value": 0}],
            subtitle="Период: 01.07.2026 — 30.09.2026",
        )
        assert data.startswith(b"%PDF")

    def test_flat_rows_and_filters_text(self) -> None:
        detail, _summary = sample_reports()
        columns, flat = kam_activity_flat_rows(detail)
        assert len(flat) == 4 and len(columns) == 18
        b2c = next(row for row in flat if row["workflow_type"] == "B2C")
        assert b2c["university"] == "Смирнова О. В."
        detail["params"]["kam_id"] = [str(KAM1)]
        assert "Петров Константин Андреевич" in describe_filters(detail)

    def test_export_job_body_accepts_pdf(self) -> None:
        assert ExportJobBody(entity_type="report:kam-activity", format="pdf").format == "pdf"
        with pytest.raises(ValidationError):
            ExportJobBody(entity_type="report:kam-activity", format="docx")  # type: ignore[arg-type]


# --- демо-план активности (tools/seed_activity.py) ---------------------------------------


class TestActivityPlan:
    NOW = datetime(2026, 10, 16, 15, 0, tzinfo=timezone.utc)

    def _plan(self):  # noqa: ANN202
        from app.tools.seed_activity import build_activity_plan

        return build_activity_plan(self.NOW)

    def test_window_and_determinism(self) -> None:
        from app.tools.seed_activity import activity_window

        plan = self._plan()
        assert plan.window_start == datetime(2026, 9, 1, tzinfo=MSK)
        later = datetime(2027, 3, 1, tzinfo=timezone.utc)
        assert activity_window(later)[0] == later - timedelta(days=45)
        again = self._plan()
        assert [d.title for d in plan.deals] == [d.title for d in again.deals]
        assert [d.steps for d in plan.deals] == [d.steps for d in again.deals]

    def test_roster_and_volume(self) -> None:
        plan = self._plan()
        assert len(plan.kams) == 20
        assert len({username for username, *_ in plan.kams}) == 20
        assert len(plan.universities) == 48
        assert 80 <= len(plan.deals) <= 130
        assert len({deal.title for deal in plan.deals}) == len(plan.deals)
        assert len(plan.students) == 36 and len(plan.clients) == 16

    def test_histories_follow_workflow_graph(self) -> None:
        from app.tools.seed_activity import _workflow_seed_codes

        allowed = {}
        for code in ("b2b", "b2c"):
            import json

            from app.tools.seed import SEEDS_DIR

            spec = json.loads((SEEDS_DIR / f"{code}.json").read_text("utf-8"))
            allowed[code] = {(t["from"], t["to"]): t for t in spec["transitions"]}
            assert set(_workflow_seed_codes()[code]) >= {s["code"] for s in spec["stages"]}
        plan = self._plan()
        for deal in plan.deals:
            for previous, step in zip(deal.steps, deal.steps[1:], strict=False):
                transition = allowed[deal.pipeline][(previous.stage, step.stage)]
                assert previous.at < step.at <= self.NOW
                assert step.is_return == transition["is_return"]
                if transition["requires_comment"]:
                    assert step.comment, (deal.title, step.stage)
                if step.actor not in (None, "head@demo"):
                    assert "kam" in transition["allowed_roles"]
            for at, _author, _body in deal.comments:
                assert deal.created_at <= at <= self.NOW

    def test_personas_shape_activity(self) -> None:
        plan = self._plan()
        start = plan.window_start

        def own_in_window(username: str) -> int:
            total = 0
            for deal in plan.deals:
                if deal.kam != username:
                    continue
                total += sum(1 for s in deal.steps[1:] if s.actor == username and s.at >= start)
                total += sum(
                    1 for at, author, _b in deal.comments if author == username and at >= start
                )
            return total

        assert own_in_window("kam18@demo") == 0  # «пропавший» КАМ
        assert own_in_window("kam20@demo") == 0  # уволился в конце августа
        assert own_in_window("kam03@demo") > own_in_window("kam09@demo") > 0
        assert any(s.kind == "integration" and s.actor is None for d in plan.deals for s in d.steps)
        assert any(s.actor == "head@demo" for d in plan.deals for s in d.steps)
        talent_in_window = [
            at
            for s in plan.students
            for status, at in s.path
            if status == "talent_pool" and at >= start
        ]
        assert talent_in_window
