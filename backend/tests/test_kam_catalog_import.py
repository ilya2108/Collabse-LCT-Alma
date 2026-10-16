"""Импорт «Таблица КАМов» (P1-12): маппинг заголовков, разбор ячеек и кривые файлы.

Чистые проверки без БД: реестр типа `kam_catalog`, автопредложение маппинга по синонимам
(в том числе нестрогое — «статус по передаче» ↔ «статус передачи»), разбор дат и срока
лицензии, статуса передачи, ФИО КАМа и контактов вуза, маскирование ПДн, права на тип
импорта и парсинг «кривых» файлов: cp1251, лишние колонки, пустые строки, объединённые
ячейки Excel. Сквозной сценарий с PostgreSQL — tests/db/test_kam_catalog_import_db.py.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.core.errors import ApiError
from app.models.crm import CONTRACT_PRODUCT_TRANSFER_STATUSES
from app.models.user import AppUser
from app.modules.import_export.importers import (
    TRANSFER_STATUS_LABELS,
    ApplyStats,
    LicenseTerm,
    add_years,
    get_importer,
    license_end,
    mask_contacts,
    merge_license,
    parse_contacts,
    parse_license_date,
    parse_license_term,
    person_name_keys,
    product_vendor,
    resolve_transfer_status,
    split_direction_names,
    with_vendor,
)
from app.modules.import_export.mapping import (
    KAM_CATALOG_FIELDS,
    entity_fields,
    map_row,
    match_confidence,
    normalize_header,
    parse_mapping_entries,
    suggest_mapping,
    supported_entity_types,
    validate_mapping,
)
from app.modules.import_export.parsers import _fill_merged, parse_table
from app.modules.import_export.service import ensure_entity_access, mask_pii_values

FIXTURE = Path(__file__).parent / "fixtures" / "kam_catalog_example.xlsx"

ALL_FIELDS = {
    "university",
    "it_direction",
    "vendor",
    "product",
    "contract_number",
    "license_signed_at",
    "license_term",
    "transfer_status",
    "kam",
    "university_contacts",
    "comment",
}


def _suggested(columns: list[str]) -> dict[str, str]:
    """{заголовок: целевое поле} автопредложения для таблицы КАМов."""
    return {
        columns[s["source_column"]]: s["target_field"]
        for s in suggest_mapping(columns, KAM_CATALOG_FIELDS)
    }


class TestRegistry:
    def test_registered_with_required_fields_and_key(self) -> None:
        assert "kam_catalog" in supported_entity_types()
        spec = entity_fields("kam_catalog")
        assert spec is KAM_CATALOG_FIELDS
        assert set(spec.by_name()) == ALL_FIELDS
        assert spec.required_fields() == ("university", "product", "contract_number")
        assert spec.default_match_by == ("contract_number",)
        assert spec.pii_fields() == {"university_contacts": "contacts"}

    def test_importer_registered(self) -> None:
        importer = get_importer("kam_catalog")
        assert importer.validate.__name__ == "_validate_kam_catalog"
        assert importer.apply.__name__ == "_apply_kam_catalog"

    def test_transfer_statuses_match_model_and_check(self) -> None:
        assert tuple(TRANSFER_STATUS_LABELS) == CONTRACT_PRODUCT_TRANSFER_STATUSES

    def test_mapping_requires_key_fields(self) -> None:
        entries = parse_mapping_entries([{"source_column": 0, "target_field": "university"}])
        with pytest.raises(ApiError) as exc:
            validate_mapping(entries, 3, KAM_CATALOG_FIELDS)
        messages = " ".join(d["message"] for d in exc.value.details or [])
        assert "ИТ-продукт" in messages and "Номер договора" in messages

    def test_match_or_create_rule_means_create_missing(self) -> None:
        """Правило мастера «Создавать отсутствующие» приходит как lookup, не как флаг."""
        entries = parse_mapping_entries(
            [
                {"source_column": 0, "target_field": "university", "lookup": "match_or_create"},
                {"source_column": 1, "target_field": "product", "lookup": "strict"},
                {"source_column": 2, "target_field": "contract_number", "create_missing": True},
            ]
        )
        assert [e.create_missing for e in entries] == [True, False, True]


class TestEntityAccess:
    """Составной импорт пишет вузы и договоры — нужны их права из каталога F2."""

    @pytest.mark.parametrize(
        ("role", "allowed"),
        [("admin", True), ("head_kam", True), ("kam", True), ("observer", False)],
    )
    def test_system_roles(self, role: str, allowed: bool) -> None:
        user = AppUser(username=f"{role}@test", roles=[role])
        if allowed:
            ensure_entity_access(KAM_CATALOG_FIELDS, user)
            return
        with pytest.raises(ApiError) as exc:
            ensure_entity_access(KAM_CATALOG_FIELDS, user)
        assert exc.value.status_code == 403
        assert "Создание и правка вузов и контактов" in exc.value.message

    def test_plain_types_need_nothing_extra(self) -> None:
        ensure_entity_access(entity_fields("universities"), AppUser(roles=["observer"]))


class TestSuggestMapping:
    def test_fixture_headers(self) -> None:
        columns = parse_table(FIXTURE.read_bytes()).columns
        suggested = _suggested(columns)
        assert set(suggested.values()) == ALL_FIELDS  # 11 полей из 12 колонок
        assert "Служебная колонка" not in suggested  # лишняя колонка не маппится
        assert suggested["ПО"] == "product"
        assert suggested["№ договора"] == "contract_number"
        assert suggested["Менеджер (Фамилия Имя)"] == "kam"
        assert suggested["Контакты вуза"] == "university_contacts"

    def test_organizer_style_headers(self) -> None:
        columns = [
            "Название ВУЗа",
            "Вендор",
            "ПО",
            "Номер договора",
            "Подписание лицензии",
            "Срок действия лицензии, год",
            "Статус по передаче лицензии",
            "ФИО менеджера",
            "Ответственные от ВУЗа",
            "Комментарий",
        ]
        assert _suggested(columns) == {
            "Название ВУЗа": "university",
            "Вендор": "vendor",
            "ПО": "product",
            "Номер договора": "contract_number",
            "Подписание лицензии": "license_signed_at",
            "Срок действия лицензии, год": "license_term",
            "Статус по передаче лицензии": "transfer_status",
            "ФИО менеджера": "kam",
            "Ответственные от ВУЗа": "university_contacts",
            "Комментарий": "comment",
        }

    def test_stem_match_tolerates_word_forms_and_order(self) -> None:
        header = normalize_header("Лицензия: статус по передачи")
        assert match_confidence(header, "статус передачи лицензии") == 0.7
        assert match_confidence(normalize_header("Статус"), "статус передачи") == 0.8

    def test_short_synonym_only_whole_word(self) -> None:
        """«по» не должно находиться внутри «Подписание», «кам» — внутри «Камера»."""
        assert match_confidence(normalize_header("Подписание лицензии"), "по") is None
        assert match_confidence(normalize_header("Камера"), "кам") is None
        assert match_confidence(normalize_header("ИТ-продукт (ПО)"), "по") == 0.8
        assert _suggested(["Подписание лицензии"]) == {"Подписание лицензии": "license_signed_at"}

    def test_exact_match_wins_over_earlier_partial(self) -> None:
        """Жадное назначение по уверенности: «Статус договора» (частичное совпадение с
        «договор») не забирает поле у точного «Номер договора», стоящего правее."""
        suggested = _suggested(["Статус договора", "Номер договора"])
        assert suggested == {"Номер договора": "contract_number"}

    def test_one_field_per_column_and_column_per_field(self) -> None:
        suggestions = suggest_mapping(["Вуз", "Вуз", "ПО"], KAM_CATALOG_FIELDS)
        assert [s["target_field"] for s in suggestions] == ["university", "product"]
        assert [s["source_column"] for s in suggestions] == [0, 2]


class TestLicenseCells:
    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("01.02.2026", date(2026, 2, 1)),
            ("1.2.26", date(2026, 2, 1)),
            ("01/02/2026", date(2026, 2, 1)),
            ("2026-02-01", date(2026, 2, 1)),
            ("2026-02-01T00:00:00", date(2026, 2, 1)),
            ("01.02.2026 г.", date(2026, 2, 1)),
        ],
    )
    def test_dates(self, raw: str, expected: date) -> None:
        assert parse_license_date(raw) == expected

    @pytest.mark.parametrize("raw", ["32.13.2026", "вчера", "01.02.1800", ""])
    def test_bad_dates(self, raw: str) -> None:
        with pytest.raises(ValueError):
            parse_license_date(raw)

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            ("1", LicenseTerm(years=1)),
            ("1.0", LicenseTerm(years=1)),
            ("3 года", LicenseTerm(years=3)),
            ("5 лет", LicenseTerm(years=5)),
            ("на 2 года", LicenseTerm(years=2)),
            ("2028", LicenseTerm(valid_to=date(2028, 12, 31))),
            ("до 2028 г.", LicenseTerm(valid_to=date(2028, 12, 31))),
            ("31.08.2027", LicenseTerm(valid_to=date(2027, 8, 31))),
            ("до 31.08.2027", LicenseTerm(valid_to=date(2027, 8, 31))),
            ("Бессрочно", LicenseTerm(perpetual=True)),
            ("бессрочная", LicenseTerm(perpetual=True)),
        ],
    )
    def test_terms(self, raw: str, expected: LicenseTerm) -> None:
        assert parse_license_term(raw) == expected

    @pytest.mark.parametrize("raw", ["1,5", "0", "100", "полтора", "до весны"])
    def test_bad_terms(self, raw: str) -> None:
        with pytest.raises(ValueError):
            parse_license_term(raw)

    def test_license_end_is_inclusive_and_leap_safe(self) -> None:
        assert license_end(date(2026, 2, 1), 1) == date(2027, 1, 31)
        assert add_years(date(2028, 2, 29), 1) == date(2029, 2, 28)

    def test_merge_new_line(self) -> None:
        assert merge_license((None, None, None), date(2026, 2, 1), LicenseTerm(years=1)) == (
            date(2026, 2, 1),
            1,
            date(2027, 1, 31),
        )
        # срок без даты подписания — только годы
        assert merge_license((None, None, None), None, LicenseTerm(years=3)) == (None, 3, None)

    def test_merge_keeps_saved_values_on_empty_cells(self) -> None:
        saved = (date(2026, 2, 1), 1, date(2027, 1, 31))
        assert merge_license(saved, None, None) == saved

    def test_merge_resign_recomputes_end(self) -> None:
        saved = (date(2025, 2, 1), 1, date(2026, 1, 31))
        assert merge_license(saved, date(2026, 2, 1), None) == (
            date(2026, 2, 1),
            1,
            date(2027, 1, 31),
        )
        # явная старая дата окончания раньше нового подписания — устарела
        explicit = (date(2025, 2, 1), None, date(2025, 12, 31))
        assert merge_license(explicit, date(2026, 2, 1), None) == (date(2026, 2, 1), None, None)

    def test_merge_perpetual_clears_term(self) -> None:
        saved = (date(2026, 2, 1), 1, date(2027, 1, 31))
        assert merge_license(saved, None, LicenseTerm(perpetual=True)) == (
            date(2026, 2, 1),
            None,
            None,
        )

    def test_merge_rejects_end_before_signing(self) -> None:
        with pytest.raises(ValueError, match="раньше её подписания"):
            merge_license(
                (date(2026, 2, 1), None, None), None, LicenseTerm(valid_to=date(2025, 12, 31))
            )


class TestTransferStatus:
    @pytest.mark.parametrize(
        ("raw", "code"),
        [
            ("Передана", "transferred"),
            ("передано", "transferred"),
            ("В процессе передачи", "in_progress"),
            ("в работе", "in_progress"),
            ("Не передана", "not_started"),
            ("Отказ.", "cancelled"),
            ("Отменена", "cancelled"),
            ("transferred", "transferred"),
            ("  ПЕРЕДАЁТСЯ ", "in_progress"),
        ],
    )
    def test_known(self, raw: str, code: str) -> None:
        assert resolve_transfer_status(raw) == code

    def test_unknown(self) -> None:
        assert resolve_transfer_status("потеряна") is None


class TestKamName:
    def test_keys_from_full_name(self) -> None:
        assert person_name_keys("Петров Константин Андреевич") == {
            "петров константин андреевич",
            "петров константин",
            "константин петров",
            "петров к",
            "петров к а",
        }

    def test_yo_and_short_names(self) -> None:
        assert "федоров максим" in person_name_keys("Фёдоров Максим Ильич")
        assert person_name_keys("Петров") == set()  # одна фамилия — не ключ
        assert person_name_keys("") == set()


class TestContacts:
    def test_parse_rich_cell(self) -> None:
        contacts = parse_contacts(
            "Смирнова Анна Олеговна (проректор), +7 900 000-00-01, smirnova@example.edu;"
            " Орлов И.П. — декан\nБелова Ольга, зав. кафедрой; ; +7 900 000-00-02"
        )
        assert [(c.full_name, c.position, c.phone, c.email) for c in contacts] == [
            ("Смирнова Анна Олеговна", "проректор", "+79000000001", "smirnova@example.edu"),
            ("Орлов И.П.", "декан", None, None),
            ("Белова Ольга", "зав. кафедрой", None, None),
        ]

    def test_duplicates_and_empty(self) -> None:
        assert [c.full_name for c in parse_contacts("Орлов Игорь; орлов  игорь")] == ["Орлов Игорь"]
        assert parse_contacts("+7 900 000-00-02; ivan@example.edu") == []

    def test_mask_hides_phone_email_and_names(self) -> None:
        masked = mask_contacts("Смирнова Анна Олеговна (проректор), +7 900 000-00-01; Орлов Игорь")
        assert masked == "С***а А.О.; О***в И."
        assert "900" not in masked
        assert mask_contacts("+7 900 000-00-01") == "***"

    def test_row_masking_in_error_report(self) -> None:
        values = {"university_contacts": "Орлов Игорь Петрович, +7 900 000-00-02", "kam": "Петров"}
        assert mask_pii_values(values, KAM_CATALOG_FIELDS.pii_fields()) == {
            "university_contacts": "О***в И.П.",
            "kam": "Петров",
        }


class TestSmallHelpers:
    def test_directions_split(self) -> None:
        assert split_direction_names("DevOps, QA; devops\nАнализ данных") == [
            "DevOps",
            "QA",
            "Анализ данных",
        ]
        assert split_direction_names(None) == []

    def test_vendor_in_description(self) -> None:
        assert with_vendor(None, "ООО «Пример»") == "Вендор: ООО «Пример»"
        assert with_vendor("Вендор: Старый\nОписание", "Новый") == "Вендор: Новый\nОписание"
        assert with_vendor("Описание", "Новый") == "Вендор: Новый\nОписание"
        assert product_vendor("Вендор: ООО «Пример»\nОписание") == "ООО «Пример»"
        assert product_vendor("Описание") is None

    def test_apply_stats_details_only_when_present(self) -> None:
        assert ApplyStats(created=1).as_dict() == {"created": 1, "updated": 0, "skipped": 0}
        assert ApplyStats(details={"licenses_created": 2}).as_dict()["details"] == {
            "licenses_created": 2
        }


class TestCrookedFiles:
    """Кривые файлы: кодировка, лишние колонки, пустые строки, объединённые ячейки, дубли."""

    def test_fixture_merged_university_and_empty_row(self) -> None:
        table = parse_table(FIXTURE.read_bytes())
        assert len(table.columns) == 12
        # A2:A3 объединены — вуз есть и во второй строке лицензии
        assert table.rows[0][0] == table.rows[1][0] == "Северный технологический университет"
        assert not any(table.rows[2])  # пустая строка — её пропустит skip_empty_rows
        assert table.rows[3][5] == "2026-03-15"  # Excel-дата → ISO

    def test_fixture_maps_rows(self) -> None:
        table = parse_table(FIXTURE.read_bytes())
        entries = parse_mapping_entries(
            [
                {"source_column": s["source_column"], "target_field": s["target_field"]}
                for s in suggest_mapping(table.columns, KAM_CATALOG_FIELDS)
            ]
        )
        mapped, errors = map_row(table.rows[1], entries)
        assert errors == []
        assert mapped["university"] == "Северный технологический университет"
        assert mapped["contract_number"] == "ЛЦ-2026/001"
        assert mapped["license_term"] == "до 2028"
        assert mapped["it_direction"] == "DevOps, QA"

    def test_cp1251_csv_with_extra_columns_and_blank_lines(self) -> None:
        text = (
            "Вуз;Лишняя;ПО;Номер договора;Статус передачи;Ещё лишняя\r\n"
            "\r\n"
            "Северный технологический университет;x;Пример.Аналитика;ЛЦ-1;Передана;y\r\n"
            ";;;;;\r\n"
            "Южный гуманитарный институт;;Демо.Хранилище;ЛЦ-2;не передана;\r\n"
        )
        table = parse_table(text.encode("cp1251"))
        assert table.encoding == "cp1251"
        assert table.columns[0] == "Вуз"
        assert len(table.rows) == 2  # пустые строки CSV отброшены
        suggested = _suggested(table.columns)
        assert suggested == {
            "Вуз": "university",
            "ПО": "product",
            "Номер договора": "contract_number",
            "Статус передачи": "transfer_status",
        }

    def test_xlsx_merged_range_is_clipped(self) -> None:
        grid: list[list[object]] = [["a", None], [None, None]]
        _fill_merged(grid, [(0, 0, 1_048_575, 16_383)])
        assert grid == [["a", "a"], ["a", "a"]]

    def test_xlsx_merged_header_and_body(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Вуз", "Лицензия", None])
        sheet.append(["СТУ", "01.02.2026", "1"])
        sheet.append([None, "01.03.2026", "2"])
        sheet.merge_cells("B1:C1")
        sheet.merge_cells("A2:A3")
        buffer = io.BytesIO()
        workbook.save(buffer)
        table = parse_table(buffer.getvalue())
        assert table.columns == ["Вуз", "Лицензия", "Лицензия"]
        assert [row[0] for row in table.rows] == ["СТУ", "СТУ"]
